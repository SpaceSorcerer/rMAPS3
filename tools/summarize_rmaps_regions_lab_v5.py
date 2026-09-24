#!/usr/bin/env python3
"""SUPERSEDED - use tools/calibrate_ranksum_v2.py (the reported layer is the rank-sum statistic since 2026-09-20).

Kept because tools/rank_stability.py imports its readers, and to reproduce the 2026-09-17 Fisher
figures. Summarize rMAPS3 SE output with optional independent permutation refinement.

For a pre-specified statistic, each fresh stage-2 plus-one estimate is a valid
Monte Carlo p-value under label exchangeability. Selecting whether to report B1
or B2 using the stage-1 p-value is adaptive: the final mixture is not guaranteed
super-uniform, and the usual BH FDR guarantee is not established by refinement.
Stage 1 is unchanged from v4. With refinement disabled the scientific TSVs and
readout match v4 byte-for-byte; workbook provenance and command logs necessarily
retain the actual script identity, command, and measured timing.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import subprocess
import sys
import time
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font
from scipy.special import gammaln
from scipy.stats import false_discovery_control, fisher_exact, hypergeom

# The audited engine of this checkout (legacy/, rmaps_core/, cli.py unchanged since 3faead9).
ENGINE_ROOT = Path(__file__).resolve().parents[1]


REGIONS = (
    "UpstreamExon_3prime",
    "UpstreamExonIntron",
    "UpstreamIntron",
    "TargetExon_5prime",
    "TargetExon-3prime",
    "DownstreamIntron",
    "DownstreamExonIntron",
    "DownstreamExon_5prime",
)
POOL_TO_REGIONS = {
    "Upstream Intron": ("UpstreamExonIntron", "UpstreamIntron"),
    "Exon Body": ("TargetExon_5prime", "TargetExon-3prime"),
    "Downstream Intron": ("DownstreamIntron", "DownstreamExonIntron"),
    "Flanking Exon": ("UpstreamExon_3prime", "DownstreamExon_5prime"),
}
REGION_TO_POOL = {region: pool for pool, regions in POOL_TO_REGIONS.items() for region in regions}
PLOT_POOLS = {"Upstream Intron", "Exon Body", "Downstream Intron"}
DIRECTION_LABEL = {"up": "INCLUDED", "dn": "SKIPPED"}
DIRECTION_CODE = {"up": 0, "dn": 1}
EXON_REGIONS = {
    "UpstreamExon_3prime", "TargetExon_5prime", "TargetExon-3prime", "DownstreamExon_5prime"
}
ROOT_COLUMNS = {
    "UpstreamExon_3prime": "smallest_p_in_upstreamExon-3prime",
    "UpstreamExonIntron": "smallest_p_in_upstreamExonIntron",
    "UpstreamIntron": "smallest_p_in_upstreamIntron",
    "TargetExon_5prime": "smallest_p_in_targetExon-5prime",
    "TargetExon-3prime": "smallest_p_in_targetExon-3prime",
    "DownstreamIntron": "smallest_p_in_downstreamIntron",
    "DownstreamExonIntron": "smallest_p_in_downstreamExonIntron",
    "DownstreamExon_5prime": "smallest_p_in_downstreamExon-5prime",
}


class CalibrationResult(NamedTuple):
    observed_min_p: float
    calib_p: float
    perm_min_p: np.ndarray
    reason: str
    observed_p: np.ndarray


def _finite_float(value: str | float | None) -> float:
    if value is None or str(value).strip().upper() in {"", "NA", "NAN"}:
        return math.nan
    return float(value)


def window_counts(hits: np.ndarray, foreground: np.ndarray) -> dict[str, np.ndarray]:
    """Return per-window 2x2 ingredients, excluding every NA cell."""
    hits = np.asarray(hits, dtype=float)
    foreground = np.asarray(foreground, dtype=bool)
    eligible = ~np.isnan(hits)
    binary = np.nan_to_num(hits, nan=0.0).astype(np.int64)
    fg = foreground.astype(np.int64)
    bg = (~foreground).astype(np.int64)
    return {
        "fg_hit": fg @ binary,
        "fg_elig": fg @ eligible.astype(np.int64),
        "bg_hit": bg @ binary,
        "bg_elig": bg @ eligible.astype(np.int64),
    }


def matrix_counts(
    hits: np.ndarray, eligible: np.ndarray, foreground: np.ndarray
) -> dict[str, np.ndarray]:
    """Observed 2x2 ingredients from explicit sparse-schema eligibility."""
    binary = np.asarray(hits, dtype=np.float32)
    eligibility = np.asarray(eligible, dtype=np.float32)
    fg = np.asarray(foreground, dtype=np.float32)
    fg_hit = np.rint(fg @ binary).astype(np.int64)
    fg_elig = np.rint(fg @ eligibility).astype(np.int64)
    total_hit = np.rint(binary.sum(axis=0)).astype(np.int64)
    total_elig = np.rint(eligibility.sum(axis=0)).astype(np.int64)
    return {
        "fg_hit": fg_hit,
        "fg_elig": fg_elig,
        "bg_hit": total_hit - fg_hit,
        "bg_elig": total_elig - fg_elig,
    }


def fisher_pvalues(
    fg_hit: np.ndarray,
    fg_elig: np.ndarray,
    bg_hit: np.ndarray,
    bg_elig: np.ndarray,
    alternative: str,
) -> np.ndarray:
    """Vectorized Fisher p-values; invalid windows remain NA."""
    a = np.asarray(fg_hit, dtype=np.int64)
    n1 = np.asarray(fg_elig, dtype=np.int64)
    c = np.asarray(bg_hit, dtype=np.int64)
    n0 = np.asarray(bg_elig, dtype=np.int64)
    out = np.full(np.broadcast_shapes(a.shape, n1.shape, c.shape, n0.shape), np.nan)
    a, n1, c, n0 = np.broadcast_arrays(a, n1, c, n0)
    valid = (n1 > 0) & (n0 > 0) & (a >= 0) & (c >= 0) & (a <= n1) & (c <= n0)
    if not np.any(valid):
        return out
    if alternative == "greater":
        total = n1[valid] + n0[valid]
        successes = a[valid] + c[valid]
        out[valid] = hypergeom.sf(a[valid] - 1, total, successes, n1[valid])
    elif alternative == "two-sided":
        tuples = np.column_stack((a[valid], n1[valid] - a[valid], c[valid], n0[valid] - c[valid]))
        unique, inverse = np.unique(tuples, axis=0, return_inverse=True)
        vals = np.array([fisher_exact(row.reshape(2, 2), alternative="two-sided").pvalue for row in unique])
        out[valid] = vals[inverse]
    else:
        raise ValueError(f"Unsupported Fisher alternative: {alternative}")
    return out



def _greater_tail_tables(total, successes, draws):
    total, successes, draws = np.broadcast_arrays(
        np.asarray(total, dtype=np.int64), np.asarray(successes, dtype=np.int64),
        np.asarray(draws, dtype=np.int64),
    )
    total, successes, draws = total.ravel(), successes.ravel(), draws.ravel()
    upper = np.minimum(successes, draws)
    lower = np.maximum(0, draws - (total - successes))
    if total.size == 0:
        return []
    j = np.arange(int(upper.max()) + 1, dtype=np.float64)[None, :]
    N, K, n = total[:, None], successes[:, None], draws[:, None]
    logpmf = (gammaln(K + 1) - gammaln(j + 1) - gammaln(K - j + 1)
              + gammaln(N - K + 1) - gammaln(n - j + 1)
              - gammaln(N - K - n + j + 1) - gammaln(N + 1)
              + gammaln(n + 1) + gammaln(N - n + 1))
    support = (j >= lower[:, None]) & (j <= upper[:, None])
    logpmf[~support] = -np.inf
    logtail = np.logaddexp.accumulate(logpmf[:, ::-1], axis=1)[:, ::-1]
    # Subtract the same partition function from both tails to remove gamma
    # cancellation in normalization; retain the upper tail for tiny p-values.
    norm = logtail[:, :1].copy()
    logtail -= norm
    tail = np.exp(logtail)
    if j.shape[1] > 1:
        logcdf = np.logaddexp.accumulate(logpmf, axis=1) - norm
        high = logtail[:, 1:] > -np.log(2.0)
        complement = -np.expm1(logcdf[:, :-1])
        tail[:, 1:] = np.where(high, complement, tail[:, 1:])
    tail[j <= lower[:, None]] = 1.0
    # Binary support is Bernoulli: P(X >= 1) = E[X] = K*n/N.
    # This exact identity preserves ties such as 187/374 == 199/398.
    singleton = upper == 1
    if np.any(singleton):
        tail[singleton, 1] = (successes[singleton] * draws[singleton]) / total[singleton]
    np.clip(tail, 0.0, 1.0, out=tail)
    return [tail[i, :int(hi) + 1].copy() for i, hi in enumerate(upper)]


def greater_pvalues(fg_hit, fg_elig, bg_hit, bg_elig):
    """One-sided reference-facing helper, with bounded padded-table batches."""
    a, n1, c, n0 = np.broadcast_arrays(*[
        np.asarray(x, dtype=np.int64) for x in (fg_hit, fg_elig, bg_hit, bg_elig)
    ])
    out = np.full(a.shape, np.nan)
    valid = (n1 > 0) & (n0 > 0) & (a >= 0) & (c >= 0) & (a <= n1) & (c <= n0)
    indices = np.flatnonzero(valid)
    flat = [x.ravel() for x in (a, n1, c, n0)]
    for start in range(0, len(indices), 256):
        idx = indices[start:start + 256]
        aa, nn, cc, mm = [x[idx] for x in flat]
        tables = _greater_tail_tables(nn + mm, aa + cc, nn)
        out.ravel()[idx] = [table[k] for table, k in zip(tables, aa)]
    return out


def _cached_greater_pvalues(fg_hit, fg_elig, total_hit, total_elig, cache):
    bg_hit = total_hit[None, :] - fg_hit
    bg_elig = total_elig[None, :] - fg_elig
    valid = ((fg_elig > 0) & (bg_elig > 0) & (fg_hit >= 0)
             & (bg_hit >= 0) & (fg_hit <= fg_elig) & (bg_hit <= bg_elig))
    out = np.full(fg_hit.shape, np.nan)
    groups = []
    missing = []
    for w in range(fg_hit.shape[1]):
        rows = np.flatnonzero(valid[:, w])
        if not rows.size:
            continue
        values, inverse = np.unique(fg_elig[rows, w], return_inverse=True)
        groups.append((w, rows, values, inverse))
        missing.extend((w, int(n)) for n in values if (w, int(n)) not in cache)
    for start in range(0, len(missing), 256):
        keys = missing[start:start + 256]
        windows = np.array([w for w, _ in keys], dtype=np.int64)
        draws = np.array([n for _, n in keys], dtype=np.int64)
        tables = _greater_tail_tables(total_elig[windows], total_hit[windows], draws)
        cache.update(zip(keys, tables))
    for w, rows, values, inverse in groups:
        for i, n in enumerate(values):
            selected = rows[inverse == i]
            out[selected, w] = cache[w, int(n)][fg_hit[selected, w]]
    return out


def generate_permutations(labels: np.ndarray, permutations: int, rng: np.random.Generator) -> np.ndarray:
    labels = np.asarray(labels, dtype=np.uint8)
    out = np.empty((permutations, labels.size), dtype=np.uint8)
    for i in range(permutations):
        out[i] = rng.permutation(labels)
    return out


def _assignment_min_p(
    hits: np.ndarray,
    eligible: np.ndarray,
    assignments: np.ndarray,
    alternative: str,
    chunk_size: int,
) -> tuple[np.ndarray, np.ndarray]:
    """Calculate window p-values and minima for one or more label assignments."""
    if alternative != "greater":
        raise ValueError("Vectorized calibration requires one-sided Fisher alternative 'greater'")
    binary = np.asarray(hits, dtype=np.float32)
    eligibility = np.asarray(eligible, dtype=np.float32)
    assignments = np.asarray(assignments, dtype=np.uint8)
    if binary.shape != eligibility.shape or binary.ndim != 2:
        raise ValueError("hits and eligibility must be same-shaped two-dimensional matrices")
    if assignments.ndim != 2 or assignments.shape[1] != binary.shape[0]:
        raise ValueError("assignment columns must equal the number of events")
    total_hit = np.rint(binary.sum(axis=0)).astype(np.int64)
    total_elig = np.rint(eligibility.sum(axis=0)).astype(np.int64)
    minima = np.full(assignments.shape[0], np.nan)
    first_pvalues = np.full(binary.shape[1], np.nan)
    tail_cache = {}
    for start in range(0, assignments.shape[0], chunk_size):
        stop = min(start + chunk_size, assignments.shape[0])
        pmat = assignments[start:stop].astype(np.float32, copy=False)
        if pmat.shape[0] <= 16:
            # Small assignment batches can trigger disproportionate BLAS thread costs.
            fg_hit = np.rint(np.einsum("ij,jk->ik", pmat, binary, optimize=False)).astype(np.int64)
            fg_elig = np.rint(np.einsum("ij,jk->ik", pmat, eligibility, optimize=False)).astype(np.int64)
        else:
            fg_hit = np.rint(pmat @ binary).astype(np.int64)
            fg_elig = np.rint(pmat @ eligibility).astype(np.int64)
        bg_hit = total_hit[None, :] - fg_hit
        bg_elig = total_elig[None, :] - fg_elig
        pvalues = _cached_greater_pvalues(fg_hit, fg_elig, total_hit, total_elig, tail_cache)
        if start == 0:
            first_pvalues = pvalues[0].copy()
        valid_rows = np.any(np.isfinite(pvalues), axis=1)
        if np.any(valid_rows):
            minima[start:stop][valid_rows] = np.nanmin(pvalues[valid_rows], axis=1)
    return minima, first_pvalues


def calibrate_min_p(
    hits: np.ndarray,
    labels: np.ndarray,
    permutations: np.ndarray,
    alternative: str,
    chunk_size: int = 500,
    eligible: np.ndarray | None = None,
) -> CalibrationResult:
    raw_hits = np.asarray(hits)
    labels = np.asarray(labels, dtype=bool)
    if eligible is None:
        eligible = ~np.isnan(raw_hits.astype(float))
        hits = np.nan_to_num(raw_hits, nan=0.0).astype(bool)
    else:
        eligible = np.asarray(eligible, dtype=bool)
        hits = np.asarray(raw_hits, dtype=bool) & eligible
    observed_minima, observed_p = _assignment_min_p(
        hits, eligible, labels.astype(np.uint8)[None, :], alternative, 1
    )
    if not np.any(np.isfinite(observed_p)):
        return CalibrationResult(math.nan, math.nan, np.full(len(permutations), np.nan),
                                 "no window has eligible events in both sets", observed_p)
    observed_min = float(observed_minima[0])
    perm_min, _ = _assignment_min_p(
        hits, eligible, permutations, alternative, chunk_size
    )
    valid_perm = np.isfinite(perm_min)
    calibrated = (1 + int(np.sum(perm_min[valid_perm] <= observed_min))) / (1 + len(permutations))
    return CalibrationResult(observed_min, float(calibrated), perm_min, "", observed_p)


def _sample_assignment_counts(binary, eligibility, labels, number, rng,
                              total_hit=None, total_elig=None):
    """Uniform fixed-cardinality subsets; sum the smaller assigned event set."""
    labels = np.asarray(labels, dtype=bool)
    n_events = labels.size
    n_foreground = int(labels.sum())
    complement = n_foreground > n_events - n_foreground
    sample_size = n_events - n_foreground if complement else n_foreground
    if complement:
        if total_hit is None:
            total_hit = np.sum(binary, axis=0, dtype=np.int64)
        if total_elig is None:
            total_elig = np.sum(eligibility, axis=0, dtype=np.int64)
    fg_hit = np.empty((number, binary.shape[1]), dtype=np.int64)
    fg_elig = np.empty_like(fg_hit)
    for i in range(number):
        selected = rng.choice(n_events, sample_size, replace=False)
        hits_sum = np.sum(binary[selected], axis=0, dtype=np.int64)
        eligible_sum = np.sum(eligibility[selected], axis=0, dtype=np.int64)
        fg_hit[i] = total_hit - hits_sum if complement else hits_sum
        fg_elig[i] = total_elig - eligible_sum if complement else eligible_sum
    return fg_hit, fg_elig


def refine_calibration(
    hits: np.ndarray,
    labels: np.ndarray,
    permutations: int,
    seed: int,
    direction: str,
    statistic: str,
    chunk_size: int = 500,
    eligible: np.ndarray | None = None,
) -> CalibrationResult:
    """Fresh fixed-B Monte Carlo calibration using uniform assignment subsets.

The same stage-2 stream is reused across motifs for each statistic/direction.
Region indices are 0..7; pooled indices are 8..11 in declaration order.
    This standalone p-value is valid under exchangeability for a pre-specified
    statistic. Adaptive selection of the reported stage has no such guarantee.
    Uniform sampling without replacement of the smaller foreground/background
    set is distributionally identical to shuffling the complete binary labels.
    It uses a different realized RNG sequence from full label shuffles. Counts
    use exact int64 sums; auxiliary memory is chunk_size x n_windows plus one
    min(n_foreground, n_background) x n_windows gather, never B2 x n_events.
"""
    if permutations < 1 or chunk_size < 1:
        raise ValueError("permutations and chunk_size must be positive")
    statistics = list(REGIONS) + list(POOL_TO_REGIONS)
    if statistic not in statistics or direction not in DIRECTION_CODE:
        raise ValueError("Unknown statistic or direction for refinement seed stream")
    rng = np.random.default_rng(np.random.SeedSequence(
        [seed, 2, statistics.index(statistic), DIRECTION_CODE[direction]]
    ))
    raw_hits = np.asarray(hits)
    labels = np.asarray(labels, dtype=bool)
    if eligible is None:
        eligible = ~np.isnan(raw_hits.astype(float))
        hits = np.nan_to_num(raw_hits, nan=0.0).astype(bool)
    else:
        eligible = np.asarray(eligible, dtype=bool)
        hits = raw_hits.astype(bool) & eligible
    observed, observed_p = _assignment_min_p(
        hits, eligible, labels.astype(np.uint8)[None, :], "greater", 1
    )
    minima = np.full(permutations, np.nan)
    if not np.any(np.isfinite(observed_p)):
        return CalibrationResult(math.nan, math.nan, minima,
                                 "no window has eligible events in both sets", observed_p)
    binary = np.asarray(hits, dtype=bool)
    eligibility = np.asarray(eligible, dtype=bool)
    total_hit = np.sum(binary, axis=0, dtype=np.int64)
    total_elig = np.sum(eligibility, axis=0, dtype=np.int64)
    tail_cache = {}
    for start in range(0, permutations, chunk_size):
        stop = min(start + chunk_size, permutations)
        fg_hit, fg_elig = _sample_assignment_counts(
            binary, eligibility, labels, stop - start, rng, total_hit, total_elig
        )
        pvalues = _cached_greater_pvalues(fg_hit, fg_elig, total_hit, total_elig, tail_cache)
        valid = np.any(np.isfinite(pvalues), axis=1)
        minima[start:stop][valid] = np.nanmin(pvalues[valid], axis=1)
    observed_min = float(observed[0])
    p = (1 + int(np.sum(minima[np.isfinite(minima)] <= observed_min))) / (1 + permutations)
    return CalibrationResult(observed_min, float(p), minima, "", observed_p)


def select_refinement_tasks(records: list[dict], threshold: float, max_tasks: int):
    """Return selected and deferred record indices, ordered by p then input index."""
    if not 0 <= threshold <= 1 or max_tasks < 0:
        raise ValueError("threshold must be in [0, 1] and max_tasks nonnegative")
    qualifying = [i for i, record in enumerate(records)
                  if math.isfinite(record["calib_p"]) and record["calib_p"] <= threshold]
    qualifying.sort(key=lambda i: (records[i]["calib_p"], i))
    return qualifying[:max_tasks], qualifying[max_tasks:]


_STAGE1_PROGRESS = None


def _progress_motifs(motifs):
    for motif in motifs:
        yield motif
        if _STAGE1_PROGRESS is not None:
            _STAGE1_PROGRESS[0] += 1
            done, total, started = _STAGE1_PROGRESS
            if done % 100 == 0 or done == total:
                print(f"stage 1: {done}/{total} tasks, elapsed {time.perf_counter() - started:.1f}s",
                      file=sys.stderr, flush=True)


def bh_adjust(values: np.ndarray | list[float]) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    out = np.full(values.shape, np.nan)
    finite = np.isfinite(values)
    if np.any(finite):
        out[finite] = false_discovery_control(values[finite], method="bh")
    return out


def read_root_table(path: Path) -> dict[str, dict[str, float]]:
    with path.open(encoding="utf-8", newline="") as handle:
        rows = csv.DictReader(handle, delimiter="\t")
        return {
            row["RBP"]: {region: _finite_float(row.get(ROOT_COLUMNS[region])) for region in REGIONS}
            for row in rows
        }


def read_manifest(path: Path) -> dict:
    manifest = json.loads(path.read_text(encoding="utf-8"))
    schema = manifest.get("schema_version")
    if not isinstance(schema, int) or isinstance(schema, bool) or schema < 2:
        raise ValueError(
            f"Sparse positional input requires run_manifest schema_version 2 or newer; found {schema!r}"
        )
    return manifest


def load_sparse_hits(path: Path, engine_root: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Load schema-2 positional data only via the engine's public loader."""
    engine_text = str(Path(engine_root).resolve())
    if engine_text not in sys.path:
        sys.path.insert(0, engine_text)
    from rmaps_core.positional_io import load_hits

    hits, eligible, labels = load_hits(path)
    hits = np.asarray(hits, dtype=bool)
    eligible = np.asarray(eligible, dtype=bool)
    labels = np.asarray(labels)
    if hits.ndim != 2 or eligible.shape != hits.shape or labels.shape != (hits.shape[0],):
        raise ValueError(f"Invalid sparse positional shapes in {path}")
    if not np.all(np.isin(labels, [0, 1, 2])):
        raise ValueError(f"Invalid set_label value in {path}; expected only 0=up, 1=dn, 2=bg")
    return hits, eligible, labels


def expected_windows(manifest: dict, region: str) -> int:
    parameters = manifest.get("parameters", {})
    step = int(parameters.get("step", 1))
    if region in EXON_REGIONS:
        length = int(parameters.get("exon", 50))
        window = int(parameters.get("exon_window", parameters.get("window", 50)))
    else:
        length = int(parameters.get("intron", 250))
        window = int(parameters.get("window", 50))
    return len(range(0, length - window + 1, step))


def engine_git_revision(engine_root: Path) -> str:
    root = str(Path(engine_root).resolve())
    command = ["git", "-c", f"safe.directory={root.replace(chr(92), '/')}", "-C", root,
               "rev-parse", "HEAD"]
    completed = subprocess.run(command, check=True, capture_output=True, text=True)
    return completed.stdout.strip()


def _region_calibration_task(args):
    run_text, engine_text, motifs, region, direction, permutations, seed, chunk_size = args
    run = Path(run_text)
    engine_root = Path(engine_text)
    selected_label = DIRECTION_CODE[direction]
    template_labels = None
    template_eligible = None
    use = None
    perm_matrix = None
    output = []
    for motif in _progress_motifs(motifs):
        path = run / "positional" / f"{motif}.{region}.hits.npz"
        if not path.exists():
            output.append((motif, None, f"missing hit matrix: {path.name}"))
            continue
        hits, eligible, labels = load_sparse_hits(path, engine_root)
        if template_labels is None:
            template_labels = labels
            use = np.isin(labels, [selected_label, 2])
            binary_labels = labels[use] == selected_label
            template_eligible = eligible
            perm_matrix = generate_permutations(
                binary_labels, permutations, stable_rng(seed, direction, region)
            )
        elif not np.array_equal(labels, template_labels):
            raise ValueError(f"Event label mismatch across motifs for {direction} {region}: {motif}")
        elif not np.array_equal(eligible, template_eligible):
            raise ValueError(f"Eligibility mismatch across motifs for {direction} {region}: {motif}")
        selected_hits = hits[use]
        selected_eligible = eligible[use]
        result = calibrate_min_p(
            selected_hits, binary_labels, perm_matrix, "greater", chunk_size, selected_eligible
        )
        counts = matrix_counts(selected_hits, selected_eligible, binary_labels)
        output.append((motif, (hits.shape[1], counts, result), ""))
    if template_labels is None:
        raise FileNotFoundError(f"No sparse hit matrix exists for region {region}")
    return region, direction, output


def _pooled_calibration_task(args):
    run_text, engine_text, motifs, pool_name, member_regions, direction, permutations, seed, chunk_size = args
    run = Path(run_text)
    engine_root = Path(engine_text)
    selected_label = DIRECTION_CODE[direction]
    template_labels = None
    template_eligible_by_region = {}
    use = None
    binary_labels = None
    perm_matrix = None
    output = []
    for motif in _progress_motifs(motifs):
        loaded = []
        missing = None
        for region in member_regions:
            path = run / "positional" / f"{motif}.{region}.hits.npz"
            if not path.exists():
                missing = f"missing pooled hit matrix: {path.name}"
                break
            hits, eligible, labels = load_sparse_hits(path, engine_root)
            if template_labels is None:
                template_labels = labels
                use = np.isin(labels, [selected_label, 2])
                binary_labels = labels[use] == selected_label
                perm_matrix = generate_permutations(
                    binary_labels, permutations, stable_rng(seed, direction, f"pooled:{pool_name}")
                )
            elif not np.array_equal(labels, template_labels):
                raise ValueError(f"Event label mismatch in pooled calibration: {motif} {direction} {pool_name}")
            if region not in template_eligible_by_region:
                template_eligible_by_region[region] = eligible
            elif not np.array_equal(eligible, template_eligible_by_region[region]):
                raise ValueError(
                    f"Eligibility mismatch in pooled calibration: {motif} {direction} {pool_name} {region}"
                )
            loaded.append((hits[use], eligible[use]))
        if missing:
            output.append((motif, None, missing))
            continue
        results = [
            calibrate_min_p(h, binary_labels, perm_matrix, "greater", chunk_size, e)
            for h, e in loaded
        ]
        output.append((motif, results, ""))
    if template_labels is None:
        return pool_name, direction, [(m, None, "no motif has both pooled-region hit matrices") for m in motifs]
    return pool_name, direction, output


def read_positional_p(path: Path) -> dict[str, list[tuple[int, float, str]]]:
    result = {region: [] for region in REGIONS}
    if not path.exists():
        return result
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            if row["Region"] in result:
                result[row["Region"]].append(
                    (int(row["position"]), _finite_float(row["fisher.exact.pVal"]), row.get("reason", ""))
                )
    return result


def stable_rng(seed: int, direction: str, region: str) -> np.random.Generator:
    digest = hashlib.sha256(f"{seed}|{direction}|{region}".encode()).digest()
    derived = int.from_bytes(digest[:8], "little")
    return np.random.default_rng(derived)


def md5(path: Path) -> str:
    h = hashlib.md5()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _write_tsv(path: Path, rows: list[dict], columns: list[str] | None = None) -> None:
    if columns is None:
        columns = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _excel_value(value):
    if isinstance(value, (np.integer, np.floating)):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, float) and math.isinf(value):
        return "Inf" if value > 0 else "-Inf"
    return value


def write_workbook(path: Path, sheets: dict[str, list[dict]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        columns = list(rows[0]) if rows else ["status"]
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_excel_value(row.get(col)) for col in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            width = min(60, max(len(str(c.value)) if c.value is not None else 0 for c in col_cells) + 2)
            ws.column_dimensions[col_cells[0].column_letter].width = max(10, width)
    wb.save(path)


def flatten_json(value, prefix="") -> list[dict[str, str]]:
    rows = []
    if isinstance(value, dict):
        for key, child in value.items():
            rows.extend(flatten_json(child, f"{prefix}.{key}" if prefix else str(key)))
    elif isinstance(value, list):
        rows.append({"field": prefix, "value": json.dumps(value, separators=(",", ":"))})
    else:
        rows.append({"field": prefix, "value": str(value)})
    return rows


def _refine_records(run, engine_root, rows, pooled_records, permutations, seed,
                    refine_perms, threshold, max_tasks, chunk_size):
    started = time.perf_counter()
    records = []
    targets = []
    for row in rows:
        row.update(calib_perms_used=permutations, calib_stage=1)
        records.append({"kind": "region", "motif_key": row["motif_key"],
                        "direction": row["direction"], "statistic": row["region"],
                        "calib_p": row["calib_p"]})
        targets.append(row)
    for (motif, direction, pool_name), record in pooled_records.items():
        record.update(calib_perms_used_pooled=permutations, calib_stage_pooled=1)
        records.append({"kind": "pooled", "motif_key": motif, "direction": direction,
                        "statistic": pool_name, "calib_p": record["calib_p_pooled"]})
        targets.append(record)
    selected, deferred = select_refinement_tasks(records, threshold, max_tasks)
    print(f"stage 2: {len(selected)} selected, {len(deferred)} qualifying tasks deferred by guard",
          file=sys.stderr, flush=True)
    audit = []
    for done, index in enumerate(selected, 1):
        task = records[index]
        task_started = time.perf_counter()
        member_regions = ((task["statistic"],) if task["kind"] == "region"
                          else POOL_TO_REGIONS[task["statistic"]])
        matrices = []
        eligibilities = []
        labels = None
        for region in member_regions:
            path = run / "positional" / f"{task['motif_key']}.{region}.hits.npz"
            hits, eligible, current_labels = load_sparse_hits(path, engine_root)
            if labels is not None and not np.array_equal(labels, current_labels):
                raise ValueError(f"Event label mismatch during pooled refinement: {path}")
            labels = current_labels
            use = np.isin(labels, [DIRECTION_CODE[task["direction"]], 2])
            matrices.append(hits[use])
            eligibilities.append(eligible[use])
        result = refine_calibration(
            np.concatenate(matrices, axis=1), labels[use] == DIRECTION_CODE[task["direction"]],
            refine_perms, seed, task["direction"], task["statistic"], chunk_size,
            np.concatenate(eligibilities, axis=1),
        )
        target = targets[index]
        suffix = "" if task["kind"] == "region" else "_pooled"
        previous_observed = target[f"calib_observed_min_p{suffix}"]
        if not math.isclose(result.observed_min_p, previous_observed, rel_tol=1e-12, abs_tol=0):
            raise ValueError(f"Observed minimum changed during refinement: {task}")
        target[f"calib_p{suffix}"] = result.calib_p
        target[f"calib_perms_used{suffix}"] = refine_perms
        target[f"calib_stage{suffix}"] = 2
        audit.append({**task, "status": "refined", "final_calib_p": result.calib_p,
                      "calib_perms_used": refine_perms, "calib_stage": 2,
                      "wall_seconds": time.perf_counter() - task_started})
        if done % 100 == 0 or done == len(selected):
            print(f"stage 2: {done}/{len(selected)} tasks, elapsed {time.perf_counter() - started:.1f}s",
                  file=sys.stderr, flush=True)
    for index in deferred:
        audit.append({**records[index], "status": "deferred_by_guard",
                      "final_calib_p": records[index]["calib_p"],
                      "calib_perms_used": permutations, "calib_stage": 1, "wall_seconds": 0.0})
    for row in rows:
        row.update(pooled_records[(row["motif_key"], row["direction"], row["pooled_region"])])
    return {"qualified_tasks": len(selected) + len(deferred), "refined_tasks": len(selected),
            "deferred_tasks": len(deferred), "stage2_wall_seconds": time.perf_counter() - started,
            "refine_perms": refine_perms, "refine_threshold": threshold,
            "refine_max_tasks": max_tasks, "stage2_p_floor": 1 / (1 + refine_perms)}, audit


def summarize(
    run: Path,
    out: Path,
    arm: str,
    permutations: int,
    seed: int,
    workers: int,
    engine_root: Path = ENGINE_ROOT,
    chunk_size: int = 500,
    refine_perms: int = 100000,
    refine_threshold: float = 0.005,
    refine_max_tasks: int = 400,
) -> dict:
    if workers != 1:
        raise ValueError("v5 requires workers=1 on this workstation")
    if refine_perms < 0 or not 0 <= refine_threshold <= 1 or refine_max_tasks < 0:
        raise ValueError("Invalid refinement permutation count, threshold, or task guard")
    started = time.perf_counter()
    run = run.resolve()
    out = out.resolve()
    engine_root = engine_root.resolve()
    required = [run / "run_manifest.json", run / "pVal.up.vs.bg.RNAmap.txt",
                run / "pVal.dn.vs.bg.RNAmap.txt", run / "positional", run / "temp"]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing required rMAPS3 inputs: " + "; ".join(missing))
    manifest = read_manifest(run / "run_manifest.json")
    alternative = manifest.get("parameters", {}).get("fisher_alternative")
    if alternative != "greater":
        raise ValueError(f"Vectorized calibration requires fisher_alternative='greater'; found {alternative!r}")
    if manifest.get("parameters", {}).get("genome") != "hg38":
        raise ValueError("This summarizer run is scoped to human GRCh38/hg38")
    revision = engine_git_revision(engine_root)
    roots = {d: read_root_table(run / f"pVal.{d}.vs.bg.RNAmap.txt") for d in ("up", "dn")}
    motifs = sorted(set(roots["up"]) | set(roots["dn"]))
    if not motifs:
        raise ValueError("No motif rows found in root p-value tables")
    global _STAGE1_PROGRESS
    _STAGE1_PROGRESS = [0, len(motifs) * (len(REGIONS) + len(POOL_TO_REGIONS)) * 2, started]
    out.mkdir(parents=True, exist_ok=True)
    rows: list[dict] = []
    positions_rows: list[dict] = []
    perm_minima: dict[tuple[str, str, str], np.ndarray] = {}

    region_tasks = [
        (str(run), str(engine_root), motifs, region, direction, permutations, seed, chunk_size)
        for region in REGIONS for direction in ("up", "dn")
    ]
    if workers == 1:
        region_outputs = [_region_calibration_task(task) for task in region_tasks]
    else:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            region_outputs = list(pool.map(_region_calibration_task, region_tasks))
    region_lookup = {(region, direction): results for region, direction, results in region_outputs}

    for region in REGIONS:
        for direction in ("up", "dn"):
            results = region_lookup[(region, direction)]
            for motif, payload, matrix_reason in results:
                native_p = roots[direction].get(motif, {}).get(region, math.nan)
                pos_table = read_positional_p(run / "temp" / f"{motif}.pVal.{direction}.vs.bg.txt")
                p_lookup = {pos: p for pos, p, _ in pos_table[region]}
                reason_lookup = {pos: reason for pos, _, reason in pos_table[region]}
                finite_pos = [(pos, p) for pos, p, _ in pos_table[region] if math.isfinite(p)]
                argmin = min(finite_pos, key=lambda x: (x[1], x[0]))[0] if finite_pos else None
                native_reason = "" if math.isfinite(native_p) else (
                    next((reason for _, _, reason in pos_table[region] if reason), "native p unavailable")
                )
                if finite_pos and math.isfinite(native_p):
                    positional_min = min(p for _, p in finite_pos)
                    if not math.isclose(native_p, positional_min, rel_tol=1e-12, abs_tol=1e-15):
                        raise ValueError(f"Root/positional minimum mismatch: {motif} {direction} {region}")
                base = {
                    "arm": arm, "motif_key": motif, "RBP": motif.split(".", 1)[0],
                    "direction": direction, "direction_label": DIRECTION_LABEL[direction],
                    "region": region, "pooled_region": REGION_TO_POOL[region],
                    "plot": REGION_TO_POOL[region] in PLOT_POOLS,
                    "native_p": native_p, "native_argmin_position": argmin,
                    "native_reason": native_reason,
                }
                if payload is None:
                    base.update({"calib_observed_min_p": math.nan, "calib_p": math.nan,
                                 "calib_reason": matrix_reason, "fg_proportion": math.nan,
                                 "bg_proportion": math.nan, "enrichment_ratio": math.nan,
                                 "n_fg_hit": math.nan, "n_fg_elig": math.nan,
                                 "n_bg_hit": math.nan, "n_bg_elig": math.nan})
                    perm_minima[(motif, direction, region)] = np.full(permutations, np.nan)
                    rows.append(base)
                    continue
                n_windows, counts, result = payload
                if n_windows != expected_windows(manifest, region):
                    raise ValueError(
                        f"Unexpected window count for {motif} {region}: {n_windows}; "
                        f"expected {expected_windows(manifest, region)}"
                    )
                positions = np.asarray([pos for pos, _, _ in pos_table[region]], dtype=int)
                if positions.size != n_windows:
                    raise ValueError(
                        f"Positional p-value/window count mismatch for {motif} {direction} {region}: "
                        f"{positions.size} versus {n_windows}"
                    )
                perm_minima[(motif, direction, region)] = result.perm_min_p
                for j, position in enumerate(positions):
                    fge, bge = int(counts["fg_elig"][j]), int(counts["bg_elig"][j])
                    fgh, bgh = int(counts["fg_hit"][j]), int(counts["bg_hit"][j])
                    positions_rows.append({
                        "arm": arm, "motif_key": motif, "RBP": motif.split(".", 1)[0],
                        "direction": direction, "direction_label": DIRECTION_LABEL[direction],
                        "region": region, "pooled_region": REGION_TO_POOL[region], "position": int(position),
                        "native_p": p_lookup.get(int(position), math.nan),
                        "native_reason": reason_lookup.get(int(position), ""),
                        "fg_proportion": fgh / fge if fge else math.nan,
                        "bg_proportion": bgh / bge if bge else math.nan,
                        "n_fg_hit": fgh, "n_fg_elig": fge, "n_bg_hit": bgh, "n_bg_elig": bge,
                    })
                if argmin is not None and argmin in set(positions.tolist()):
                    j = int(np.where(positions == argmin)[0][0])
                    fge, bge = int(counts["fg_elig"][j]), int(counts["bg_elig"][j])
                    fgh, bgh = int(counts["fg_hit"][j]), int(counts["bg_hit"][j])
                    fgp = fgh / fge if fge else math.nan
                    bgp = bgh / bge if bge else math.nan
                    ratio = (fgp / bgp) if bgp > 0 else (math.inf if fgp > 0 else math.nan)
                else:
                    fgh = fge = bgh = bge = math.nan
                    fgp = bgp = ratio = math.nan
                base.update({"calib_observed_min_p": result.observed_min_p, "calib_p": result.calib_p,
                             "calib_reason": result.reason, "fg_proportion": fgp,
                             "bg_proportion": bgp, "enrichment_ratio": ratio,
                             "n_fg_hit": fgh, "n_fg_elig": fge,
                             "n_bg_hit": bgh, "n_bg_elig": bge})
                rows.append(base)

    by_key = {(r["motif_key"], r["direction"], r["region"]): r for r in rows}
    pooled_records = {}
    pooled_tasks = [
        (str(run), str(engine_root), motifs, pool_name, member_regions, direction,
         permutations, seed, chunk_size)
        for direction in ("up", "dn")
        for pool_name, member_regions in POOL_TO_REGIONS.items()
    ]
    if workers == 1:
        pooled_outputs = [_pooled_calibration_task(task) for task in pooled_tasks]
    else:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            pooled_outputs = list(pool.map(_pooled_calibration_task, pooled_tasks))
    pooled_lookup = {(pool_name, direction): results
                     for pool_name, direction, results in pooled_outputs}

    for direction in ("up", "dn"):
        for pool_name, member_regions in POOL_TO_REGIONS.items():
            pooled_results = pooled_lookup[(pool_name, direction)]
            for motif, region_results, pooled_reason in pooled_results:
                member_rows = [by_key[(motif, direction, region)] for region in member_regions]
                if region_results is None:
                    native_vals = np.asarray([r["native_p"] for r in member_rows], dtype=float)
                    pooled_native = float(np.nanmin(native_vals)) if np.any(np.isfinite(native_vals)) else math.nan
                    pooled_records[(motif, direction, pool_name)] = {
                        "calib_p_pooled": math.nan, "calib_observed_min_p_pooled": math.nan,
                        "native_p_pooled": pooled_native, "calib_pooled_reason": pooled_reason,
                    }
                    for row in member_rows:
                        row.update(pooled_records[(motif, direction, pool_name)])
                    continue
                observed_values = np.asarray([result.observed_min_p for result in region_results], dtype=float)
                observed = float(np.nanmin(observed_values)) if np.any(np.isfinite(observed_values)) else math.nan
                stack = np.vstack([result.perm_min_p for result in region_results])
                valid_cols = np.any(np.isfinite(stack), axis=0)
                pooled_perm = np.full(permutations, np.nan)
                if np.any(valid_cols):
                    pooled_perm[valid_cols] = np.nanmin(stack[:, valid_cols], axis=0)
                calib = ((1 + int(np.sum(pooled_perm[np.isfinite(pooled_perm)] <= observed))) / (1 + permutations)) \
                    if math.isfinite(observed) else math.nan
                native_vals = np.asarray([r["native_p"] for r in member_rows], dtype=float)
                pooled_native = float(np.nanmin(native_vals)) if np.any(np.isfinite(native_vals)) else math.nan
                pooled_records[(motif, direction, pool_name)] = {
                    "calib_p_pooled": calib, "calib_observed_min_p_pooled": observed,
                    "native_p_pooled": pooled_native,
                    "calib_pooled_reason": "" if math.isfinite(calib) else "no valid pooled observed minimum",
                }
                for row in member_rows:
                    row.update(pooled_records[(motif, direction, pool_name)])

    refinement = None
    refinement_audit = []
    if refine_perms:
        refinement, refinement_audit = _refine_records(
            run, engine_root, rows, pooled_records, permutations, seed,
            refine_perms, refine_threshold, refine_max_tasks, chunk_size,
        )
    native_q = bh_adjust([r["native_p"] for r in rows])
    for row, q in zip(rows, native_q):
        row["native_q"] = q
    bh_keys = [(m, d, p) for m in motifs for d in ("up", "dn") for p in POOL_TO_REGIONS if p in PLOT_POOLS]
    pooled_q = bh_adjust([pooled_records[key]["calib_p_pooled"] for key in bh_keys])
    q_lookup = {key: q for key, q in zip(bh_keys, pooled_q)}
    for row in rows:
        row["calib_q"] = q_lookup.get((row["motif_key"], row["direction"], row["pooled_region"]), math.nan)

    condensed = []
    rbps = sorted({m.split(".", 1)[0] for m in motifs})
    for rbp in rbps:
        rbp_motifs = [m for m in motifs if m.split(".", 1)[0] == rbp]
        for direction in ("up", "dn"):
            for pool_name in POOL_TO_REGIONS:
                candidates = []
                for motif in rbp_motifs:
                    rec = pooled_records[(motif, direction, pool_name)]
                    candidates.append((rec["calib_p_pooled"], rec["native_p_pooled"], motif))
                finite = [x for x in candidates if math.isfinite(x[0])]
                chosen = min(finite, key=lambda x: (x[0], x[1] if math.isfinite(x[1]) else math.inf, x[2])) if finite else (math.nan, math.nan, rbp_motifs[0])
                motif = chosen[2]
                member_rows = [by_key[(motif, direction, r)] for r in POOL_TO_REGIONS[pool_name]]
                effect_row = min(member_rows, key=lambda r: r["native_p"] if math.isfinite(r["native_p"]) else math.inf)
                condensed.append({
                    "arm": arm, "RBP": rbp, "direction": direction,
                    "direction_label": DIRECTION_LABEL[direction], "pooled_region": pool_name,
                    "plot": pool_name in PLOT_POOLS, "selected_motif_key": motif,
                    "selected_native_region": effect_row["region"],
                    "calib_p_pooled": pooled_records[(motif, direction, pool_name)]["calib_p_pooled"],
                    "calib_q": q_lookup.get((motif, direction, pool_name), math.nan),
                    "native_p_pooled": pooled_records[(motif, direction, pool_name)]["native_p_pooled"],
                    "native_q_at_selected_region": effect_row["native_q"],
                    "native_argmin_position": effect_row["native_argmin_position"],
                    "fg_proportion": effect_row["fg_proportion"], "bg_proportion": effect_row["bg_proportion"],
                    "enrichment_ratio": effect_row["enrichment_ratio"],
                    "n_fg_hit": effect_row["n_fg_hit"], "n_fg_elig": effect_row["n_fg_elig"],
                    "n_bg_hit": effect_row["n_bg_hit"], "n_bg_elig": effect_row["n_bg_elig"],
                    "n_motifs_total": len(rbp_motifs),
                    "n_motifs_calib_q_lt_0.05": sum(
                        math.isfinite(q_lookup.get((m, direction, pool_name), math.nan)) and q_lookup[(m, direction, pool_name)] < 0.05
                        for m in rbp_motifs
                    ),
                    "n_motifs_native_q_lt_0.05": sum(
                        any(math.isfinite(by_key[(m, direction, r)]["native_q"]) and by_key[(m, direction, r)]["native_q"] < 0.05
                            for r in POOL_TO_REGIONS[pool_name]) for m in rbp_motifs
                    ),
                })

    if refine_perms:
        for row in condensed:
            record = pooled_records[(row["selected_motif_key"], row["direction"], row["pooled_region"])]
            row.update(calib_perms_used=record["calib_perms_used_pooled"],
                       calib_stage=record["calib_stage_pooled"])
        final_p = [pooled_records[key]["calib_p_pooled"] for key in bh_keys]
        refinement.update(
            bh_family_size=len(bh_keys),
            smallest_calib_p=min((p for p in final_p if math.isfinite(p)), default=None),
            smallest_calib_q=min((float(q) for q in pooled_q if math.isfinite(q)), default=None),
            singleton_bh_q_floor=min(1.0, len(bh_keys) / (1 + refine_perms)),
        )
    elapsed_before_write = time.perf_counter() - started
    script_path = Path(__file__).resolve()
    provenance = flatten_json(manifest)
    provenance.extend([
        {"field": "summary.arm", "value": arm}, {"field": "summary.seed", "value": str(seed)},
        {"field": "summary.permutations", "value": str(permutations)},
        {"field": "summary.chunk_size", "value": str(chunk_size)},
        {"field": "summary.workers", "value": str(workers)},
        {"field": "summary.schema_version", "value": str(manifest["schema_version"])},
        {"field": "summary.engine_git_revision", "value": revision},
        {"field": "summary.engine_root", "value": str(engine_root)},
        {"field": "summary.wall_seconds_before_write", "value": f"{elapsed_before_write:.6f}"},
        {"field": "summary.script_path", "value": str(script_path)},
        {"field": "summary.script_md5", "value": md5(script_path)},
        {"field": "summary.python", "value": sys.version.replace("\n", " ")},
        {"field": "summary.platform", "value": platform.platform()},
    ])
    for package in ("numpy", "scipy", "openpyxl"):
        provenance.append({"field": f"summary.package.{package}", "value": importlib.metadata.version(package)})
    if refine_perms:
        provenance.extend(flatten_json(refinement, "summary.refinement"))
        provenance.extend([
            {"field": "summary.refinement.seed_stream", "value":
             "numpy SeedSequence([seed,2,statistic_index,direction_index]); regions 0..7 and pools 8..11 in declaration order; up=0,dn=1; shared across motifs; independent from v4 SHA256 stage-1 stream"},
            {"field": "summary.refinement.statistic_indices", "value":
             json.dumps(list(REGIONS) + list(POOL_TO_REGIONS))},
            {"field": "summary.refinement.assignment_sampler", "value":
             "For each replicate independently use Generator.choice(n_events,min(n_fg,n_bg),replace=False), choosing foreground when n_fg<=n_bg and background otherwise. Sum sampled hit and eligibility rows in int64; subtract from totals for background sampling. Uniform subsets are mathematically equivalent to permuting fixed-cardinality binary labels, but use different realized RNG draws than full-label shuffling."},
            {"field": "summary.refinement.validity", "value":
             "Each fresh fixed-B2 plus-one p-value is a valid Monte Carlo p-value for a pre-specified statistic under label exchangeability. Adaptive selection between stage-1 and stage-2 p-values is not guaranteed super-uniform; BH applied to this mixture has no established FDR guarantee from this procedure alone."},
            {"field": "summary.refinement.pooled_assignments", "value":
             "Member window matrices concatenated; one common permutation assignment across all pooled windows."},
            {"field": "summary.refinement.selection", "value":
             "Global threshold across regional and pooled tasks; smallest stage-1 p first; ties by original record order (regions then pools); deferred tasks retain stage-1 p."},
        ])
        _write_tsv(out / "refinement_tasks.tsv", refinement_audit,
                   ["kind", "motif_key", "direction", "statistic", "calib_p", "status",
                    "final_calib_p", "calib_perms_used", "calib_stage", "wall_seconds"])
        (out / "refinement_report.json").write_text(json.dumps(refinement, indent=2) + "\n", encoding="utf-8")

    _write_tsv(out / "per_motif_regions.tsv", rows)
    _write_tsv(out / "condensed_per_rbp.tsv", condensed)
    _write_tsv(out / "positions_long.tsv", positions_rows)
    workbook = out / f"{arm}_rmaps_summary.xlsx"
    write_workbook(workbook, {"condensed": condensed, "per_motif": rows,
                              "positions": positions_rows, "provenance": provenance})
    readout_lines = [
        f"# {arm} rMAPS3 region summary",
        f"Calibrated Westfall-Young min-P results used {permutations} label permutations with seed {seed}.",
    ]
    if refine_perms:
        readout_lines[1] = (f"Stage 1 used {permutations} label permutations with seed {seed}; "
                            f"{refinement['refined_tasks']} tasks replaced by independent {refine_perms}-permutation calibrations "
                            f"({refinement['deferred_tasks']} qualifying tasks deferred by guard).")
        readout_lines.append("Adaptive stage selection does not establish super-uniform final p-values or a BH FDR guarantee; see provenance.")
    for pool_name in ("Upstream Intron", "Exon Body", "Downstream Intron"):
        for direction in ("up", "dn"):
            subset = [r for r in condensed if r["pooled_region"] == pool_name and r["direction"] == direction and math.isfinite(r["calib_q"])]
            subset.sort(key=lambda r: (r["calib_q"], r["calib_p_pooled"], r["RBP"]))
            leaders = ", ".join(f"{r['RBP']} (q={r['calib_q']:.3g})" for r in subset[:3]) or "none available"
            readout_lines.append(f"- {pool_name}, {DIRECTION_LABEL[direction]}: {leaders}; n={len(subset)} RBPs tested.")
    readout_lines.append("Both native raw regional minima and calibrated layers are reported; the figure is driven by pooled calibrated q-values.")
    (out / "readout.md").write_text("\n".join(readout_lines) + "\n", encoding="utf-8")
    command = " ".join([str(Path(sys.executable).resolve()), str(script_path), "--run", str(run), "--out", str(out),
                        "--arm", arm, "--perms", str(permutations), "--seed", str(seed), "--workers", str(workers),
                        "--chunk-size", str(chunk_size), "--engine-root", str(engine_root)])
    command += (f" --refine-perms {refine_perms} --refine-threshold {refine_threshold}"
                f" --refine-max-tasks {refine_max_tasks}")
    (out / "command.log").write_text(command + "\n", encoding="utf-8")
    (out / "versions.txt").write_text(
        f"Python\t{platform.python_version()}\n" +
        "".join(f"{p}\t{importlib.metadata.version(p)}\n" for p in ("numpy", "scipy", "openpyxl")) +
        "scope\tHomo sapiens / GRCh38 (hg38) / GENCODE v49\n",
        encoding="utf-8",
    )
    wall = time.perf_counter() - started
    result = {"motifs": len(motifs), "rows": len(rows), "wall_seconds": wall,
              "workbook": str(workbook), "out": str(out)}
    if refinement is not None:
        result["refinement"] = refinement
    return result


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--arm", required=True)
    parser.add_argument("--perms", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=149)
    parser.add_argument("--workers", type=int, default=1, choices=[1])
    parser.add_argument("--refine-perms", type=int, default=100000)
    parser.add_argument("--refine-threshold", type=float, default=0.005)
    parser.add_argument("--refine-max-tasks", type=int, default=400)
    parser.add_argument("--chunk-size", type=int, default=500)
    parser.add_argument("--engine-root", type=Path, default=ENGINE_ROOT)
    args = parser.parse_args(argv)
    if args.perms < 1 or args.workers < 1 or args.chunk_size < 1:
        parser.error("--perms, --workers, and --chunk-size must be positive integers")
    if args.refine_perms < 0 or args.refine_max_tasks < 0 or not 0 <= args.refine_threshold <= 1:
        parser.error("--refine-perms/--refine-max-tasks must be nonnegative; --refine-threshold must be in [0,1]")
    return args


def main(argv=None) -> int:
    args = parse_args(argv)
    result = summarize(args.run, args.out, args.arm, args.perms, args.seed, args.workers,
                       args.engine_root, args.chunk_size, args.refine_perms,
                       args.refine_threshold, args.refine_max_tasks)
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

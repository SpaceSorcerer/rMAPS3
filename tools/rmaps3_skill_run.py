#!/usr/bin/env python3
"""One CLI for a faithful rMAPS3 SE motif map: --mode quick (the tool's own layer) or --mode full.

quick  runs the engine exactly as the lab launchers do, converts the countDist temporaries into
       per-motif npz archives, proves those archives reproduce the engine's own root tables,
       deletes only the verified temporaries, and writes quick_summary.xlsx, the main-layer
       region lollipops (released rank-sum raw p only; --no-figures skips them) + index.html.
full   adds the lab layers: the reportable calibrated supplement (calibrate_ranksum_v2.py: target-exon
       cluster permutation, RBP-level min-P, unique-k-mer family, row-unit sensitivity columns),
       foreground-bootstrap rank stability and the version 4 region lollipops for both layers with
       their rank workbook.

Nothing here recomputes or rewrites an engine output. Every path default is a lab convenience and
is overridable by a flag.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import rmaps_countdist_io as io  # noqa: E402
from rmaps3_lab_run import COORD_HEADER, normalize_coordinates  # noqa: E402

# Site paths (genome root, released-engine checkout, GTF, exclusion lists) are arguments with no
# default: the lab passes them from the rmaps3-quick / rmaps3-full skills. Only repo files default.
DEFAULT_ALIAS = ROOT / "data" / "rbp_alias_hgnc_2026-09-17.tsv"


def figure_version() -> str:
    """FIG_VERSION of tools/build_region_lollipops_v4.py, read from its source so labels never drift."""
    text = (ROOT / "tools" / "build_region_lollipops_v4.py").read_text(encoding="utf-8")
    match = re.search(r"^FIG_VERSION = '([^']+)'", text, re.M)
    if not match:
        raise ValueError("FIG_VERSION not found in tools/build_region_lollipops_v4.py")
    return match.group(1)

STAT_CAVEATS = {
    "fisher": ("Fisher counts MOTIF HITS, not exons: repeated hits in one exon enter the 2x2 table "
               "more than once, so the margins are not event counts. Use for the tool's own layer, "
               "never as an event-level test."),
    "mannwhitney": ("The rank-sum p is scipy's asymptotic normal approximation with tie and "
                    "continuity correction. With almost every eligible exon carrying no hit it is "
                    "severely anti-conservative: report the RBP ORDER, not the p as a p."),
}
GATE_RULES = ("A", "B", "Beffect")
RULE_B_SUFFIXES = ("B", "Beffect")
BLAS_VARS = ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS",
             "NUMEXPR_NUM_THREADS", "VECLIB_MAXIMUM_THREADS")
FILTER_ALIASES = {
    "fdr": "fdr",
    "min_abs_dpsi": "abs_dpsi",
    "abs_dpsi": "abs_dpsi",
    "min_jc_per_sample": "min_sample_coverage",
    "min_sample_coverage": "min_sample_coverage",
    "bg_fdr_min": "background_fdr",
    "background_fdr": "background_fdr",
    "base_mean_floor": "base_mean_floor",
    "expr_unknown": "expr_unknown",
    "treatment_is": "treatment_is",
    "rule": "rule",
}


# ------------------------------------------------------------------ small helpers
def md5(path: Path) -> str:
    digest = hashlib.md5()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def versions() -> dict:
    out = {"python": sys.version.split()[0], "executable": sys.executable,
           "platform": platform.platform()}
    for package in ("numpy", "scipy", "openpyxl", "matplotlib", "pandas"):
        try:
            out[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            out[package] = "not installed"
    return out


def git_revision(repo: Path) -> str:
    proc = subprocess.run(["git", "-C", str(repo), "rev-parse", "HEAD"],
                          capture_output=True, text=True)
    return proc.stdout.strip() if proc.returncode == 0 else "unavailable"


def load_spec(path: Path) -> dict:
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() in {".yaml", ".yml"}:
        try:
            import yaml
        except ImportError as exc:
            raise ValueError("YAML filter spec needs PyYAML; use .json instead") from exc
        spec = yaml.safe_load(text)
    elif path.suffix.lower() == ".json":
        spec = json.loads(text)
    else:
        raise ValueError("Filter spec must be .json, .yaml or .yml")
    if not isinstance(spec, dict):
        raise ValueError("Filter spec must be an object")
    return spec


def fresh_directory(path: Path) -> None:
    if path.exists() and (not path.is_dir() or any(path.iterdir())):
        raise ValueError(f"Refusing non-empty output directory: {path}")
    path.mkdir(parents=True, exist_ok=True)


class Logger:
    def __init__(self, path: Path):
        self.path = path
        self.handle = open(path, "a", encoding="utf-8")

    def __call__(self, message: str) -> None:
        stamped = time.strftime("%Y-%m-%dT%H:%M:%S") + " " + message
        print(stamped, flush=True)
        self.handle.write(stamped + "\n")
        self.handle.flush()

    def close(self) -> None:
        self.handle.close()


def run_step(name: str, command: list, log: Logger, env: dict, cwd: Path, out: Path) -> float:
    log(f"STEP {name}: {subprocess.list2cmdline(command)}")
    started = time.perf_counter()
    stdout_path = out / "logs" / f"{name}.stdout.log"
    stderr_path = out / "logs" / f"{name}.stderr.log"
    stdout_path.parent.mkdir(parents=True, exist_ok=True)
    with open(stdout_path, "w", encoding="utf-8") as so, open(stderr_path, "w", encoding="utf-8") as se:
        proc = subprocess.run(command, cwd=str(cwd), env=env, stdout=so, stderr=se)
    wall = time.perf_counter() - started
    log(f"STEP {name}: exit={proc.returncode} wall_s={wall:.1f}")
    if proc.returncode != 0:
        raise RuntimeError(f"{name} failed (exit {proc.returncode}); see {stderr_path}")
    return wall


# ------------------------------------------------------------------ inputs
def build_event_sets(args, out: Path, log: Logger, env: dict) -> Path:
    """Pre-split an rMATS SE table with the filter spec, through the tested rule-A builder."""
    spec = load_spec(Path(args.filter))
    gates, unknown = {}, []
    expr_table = None
    for key, value in spec.items():
        if key in {"expr_table", "deseq2"}:
            expr_table = value
            continue
        if key not in FILTER_ALIASES:
            unknown.append(key)
            continue
        gates[FILTER_ALIASES[key]] = value
    if unknown:
        raise ValueError(f"Unknown filter keys: {sorted(unknown)}; allowed: "
                         f"{sorted(set(FILTER_ALIASES) | {'expr_table'})}")
    if gates.get("base_mean_floor") is not None and expr_table is None:
        raise ValueError("base_mean_floor needs expr_table (a DESeq2 table with gene_id, baseMean)")
    if gates.get("rule", "A") != "A":
        raise ValueError("the filter spec asks for rule {}; the portable builder implements rule A only. "
                         "Rule B sets are frozen concordant files: supply pre-split inputs".format(gates["rule"]))
    config = {"arm": args.arm,
              "inputs": {"rmats_se": str(Path(args.rmats_se).resolve())},
              "gates": gates}
    if expr_table is not None:
        config["inputs"]["deseq2"] = str(Path(expr_table).resolve())
    config_path = out / "event_set_config.json"
    config_path.write_text(json.dumps(config, indent=2) + "\n", encoding="utf-8")
    dest = out / "event_sets"
    run_step("build_event_sets",
             [sys.executable, str(ROOT / "tools" / "build_event_sets.py"),
              "--config", str(config_path), "--out", str(dest)], log, env, ROOT, out)
    return dest


def prepare_inputs(args, out: Path, log: Logger, env: dict):
    """Resolve up/dn/bg, check every contig against the FASTA index, and report set sizes."""
    if args.rmats_se:
        source_dir = build_event_sets(args, out, log, env)
        sources = {key: source_dir / f"{key}.coord.txt" for key in ("up", "dn", "bg")}
    else:
        sources = {"up": Path(args.up), "dn": Path(args.dn), "bg": Path(args.bg)}
        for key, path in sources.items():
            if not path.is_file():
                raise ValueError(f"Missing pre-split {key} file: {path}")
    genome_root = Path(args.genome_root)
    fasta = genome_root / args.genome / f"{args.genome}.fa"
    fai = Path(str(fasta) + ".fai")
    if not fasta.is_file():
        raise ValueError(f"Genome FASTA absent: {fasta}")
    if not fai.is_file():
        raise ValueError(f"FASTA index absent, so the contigs cannot be checked: {fai}")
    mapped = out / "inputs"
    mapped.mkdir(parents=True, exist_ok=True)
    paths, counts = normalize_coordinates(sources, fai, mapped)
    for key in ("up", "dn", "bg"):
        renamed = sum(v["count"] for v in counts[key]["renamed"].values())
        log(f"events {key}={counts[key]['n_events']} contigs_renamed={renamed}")
    summary = {key: counts[key]["n_events"] for key in counts}
    record = {"n_up": summary["up"], "n_dn": summary["dn"], "n_bg": summary["bg"],
              "n_expr_unknown_in_fg": None, "n_expr_unknown_in_bg": None,
              "source": {k: str(v) for k, v in sources.items()}}
    if args.gate_counts:
        gate = json.loads(Path(args.gate_counts).read_text(encoding="utf-8-sig"))
        for key, direction in (("n_up", "up"), ("n_dn", "dn"), ("n_bg", "bg")):
            if int(gate[key]) != summary[direction]:
                raise ValueError(f"--gate-counts {key}={gate[key]} disagrees with the "
                                 f"{direction} input ({summary[direction]} events)")
        record.update({k: int(gate[k]) for k in ("n_expr_unknown_in_fg", "n_expr_unknown_in_bg")})
        record["gate_counts_source"] = str(Path(args.gate_counts).resolve())
        log(f"gate record {args.gate_counts}: event counts match the inputs")
    (out / "event_counts.json").write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")
    if args.rmats_se:
        built = json.loads((out / "event_sets" / "counts.json").read_text())
        (out / "event_counts.json").write_text(json.dumps(built, indent=2) + "\n", encoding="utf-8")
    return paths, counts, fasta, fai


# ------------------------------------------------------------------ engine
def engine_root_of(args) -> Path:
    """The released engine must be named explicitly; the audited engine defaults to this checkout."""
    if args.engine_root:
        return Path(args.engine_root)
    if args.engine == "released":
        raise ValueError("--engine released needs --engine-root <checkout of the released engine, "
                         "e.g. a worktree of tag upstream-base-2026-09-16>")
    return ROOT


def run_engine(args, paths, out: Path, log: Logger, env: dict):
    engine_root = engine_root_of(args)
    if not (engine_root / "cli.py").is_file():
        raise ValueError(f"No cli.py under the {args.engine} engine root: {engine_root}")
    known = Path(args.known_motifs) if args.known_motifs else engine_root / "data" / "knownMotifs.human.mouse.txt"
    if args.additional_motifs and args.additional_motifs.upper() == "NA":
        additional = None
    else:
        additional = (Path(args.additional_motifs) if args.additional_motifs
                      else engine_root / "data" / "ESRP.like.motif.txt")
    for path in filter(None, (known, additional)):
        if not path.is_file():
            raise ValueError(f"Motif table absent: {path}")
    engine_out = out / "engine" / args.arm
    command = [sys.executable, str(engine_root / "cli.py"), "motif-map", "se",
               "--known-motifs", str(known), "--motifs", str(additional) if additional else "NA",
               "--fasta-root", str(Path(args.genome_root)), "--genome", args.genome,
               "--rMATS", "NA", "--miso", "NA",
               "--up", str(paths["up"]), "--down", str(paths["dn"]),
               "--background", str(paths["bg"]), "--output", str(engine_out),
               "--stat-method", args.stat_method,
               "--intron", str(args.intron), "--exon", str(args.exon),
               "--window", str(args.window), "--step", str(args.step), "--keep-temp"]
    if args.engine == "audited":
        command += ["--workers", str(args.workers)]
        if args.stat_method == "fisher":
            command += ["--fisher-alternative", "greater"]
    wall = run_step("motif_map", command, log, env, engine_root, out)
    roots = {d: engine_out / f"pVal.{d}.vs.bg.RNAmap.txt" for d in ("up", "dn")}
    for direction, path in roots.items():
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"Engine exit was zero but the {direction} root table is missing: {path}")
    # The released-run log shape the lab layers assert on before they will read an arm.
    (out / "engine" / f"{args.arm}_command.log").write_text(
        "set={} engine={} ({} @ {}) stat={} start={}\n".format(
            args.arm, args.engine, engine_root, git_revision(engine_root), args.stat_method,
            datetime.now(timezone.utc).isoformat())
        + "".join("{} *{}\n".format(md5(paths[k]), paths[k]) for k in ("up", "dn", "bg"))
        + "exit=0 wall_s={:.0f}\n".format(wall), encoding="utf-8")
    return engine_out, wall, engine_root, known, additional


def convert_and_verify(args, engine_out: Path, out: Path, log: Logger, env: dict):
    """countDist temporaries -> npz archives -> proof they reproduce the engine's root tables."""
    temp_dir = engine_out / "temp"
    result = {"converted": False, "verified": False, "deleted_files": 0, "deleted_bytes": 0,
              "reason": ""}
    if args.engine != "released":
        result["reason"] = ("the count-archive reader targets the released engine's countDist "
                            "schema; an audited-engine run keeps its temporaries and its own "
                            "positional/*.hits.npz instead")
        log("WARN " + result["reason"])
        return None, result
    if not temp_dir.is_dir() or not list(temp_dir.glob("*.countDist.bg.txt")):
        result["reason"] = f"no countDist temporaries under {temp_dir}; nothing to convert"
        log("WARN " + result["reason"])
        return None, result
    counts_root = out / "counts"
    import countdist_to_npz
    rc = countdist_to_npz.main(["--arm", args.arm, "--released-root", str(engine_out.parent),
                               "--out-root", str(counts_root)])
    if rc != 0:
        raise RuntimeError("countDist -> npz conversion failed")
    result["converted"] = True
    log(f"archives written to {counts_root / args.arm}")
    if args.stat_method != "mannwhitney":
        result["reason"] = ("verification recomputes the rank-sum statistic, so it cannot check a "
                            f"--stat-method {args.stat_method} run; temporaries kept")
        log("WARN " + result["reason"])
        return counts_root, result
    import verify_ranksum_archives
    rc = verify_ranksum_archives.main(["--arm", args.arm, "--released-root", str(engine_out.parent),
                                       "--counts-root", str(counts_root)])
    if rc != 0:
        raise RuntimeError(f"archive verification FAILED; see {counts_root / args.arm / 'VERIFY.md'}")
    result["verified"] = True
    log("archives reproduce the engine root tables (see VERIFY.md)")
    if args.keep_temp:
        result["reason"] = "--keep-temp given; verified temporaries kept"
        return counts_root, result
    deleted, total = delete_verified_temporaries(
        counts_root / args.arm / verify_ranksum_archives.MANIFEST_NAME, temp_dir,
        out / "logs" / "temp_deletion.log")
    result["deleted_files"], result["deleted_bytes"] = deleted, total
    result["reason"] = f"deleted after verification; inventory in {out / 'logs' / 'temp_deletion.log'}"
    log(f"deleted {deleted} verified temporaries ({total / 1e9:.2f} GB), exactly the verifier's list")
    return counts_root, result


def delete_verified_temporaries(manifest: Path, temp_dir: Path, inventory: Path):
    """Delete exactly the files the verifier listed, each still byte-identical to its listed md5.

    Nothing outside the list is touched, and nothing is deleted when the list is absent, names a file
    outside temp_dir, or any listed file changed since verification: every check runs before the
    first deletion."""
    import verify_ranksum_archives
    if not manifest.is_file():
        raise RuntimeError(f"verification passed but wrote no deletion list: {manifest}; nothing deleted")
    listed = verify_ranksum_archives.read_manifest(manifest)
    root = temp_dir.resolve()
    for path, size, digest, _ in listed:
        resolved = path.resolve()
        if resolved.parent != root:
            raise RuntimeError(f"deletion list names a file outside {root}: {path}; nothing deleted")
        if not resolved.is_file() or resolved.stat().st_size != size or md5(resolved) != digest:
            raise RuntimeError(f"{path} changed since verification; nothing deleted")
    deleted, total = 0, 0
    with open(inventory, "w", encoding="utf-8") as handle:
        handle.write("path\tbytes\tmd5\tkind\n")
        for path, size, digest, kind in listed:
            handle.write(f"{path}\t{size}\t{digest}\t{kind}\n")
            path.resolve().unlink()
            deleted += 1
            total += size
    return deleted, total


# ------------------------------------------------------------------ quick summary
def load_alias(path: Path) -> dict:
    alias = {}
    with open(path, "r", encoding="utf-8") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        if header[:2] != ["table_name", "hgnc_symbol"]:
            raise ValueError(f"Unexpected alias header in {path}")
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            alias[fields[0]] = fields[1]
    return alias


def motif_scores(counts_root, arm: str, motif: str):
    """Per-sub-region argmin window and the engine's own counts there, from the npz archive."""
    import numpy as np
    if counts_root is None:
        return None
    path = Path(counts_root) / arm / f"{motif}.counts.npz"
    if not path.is_file():
        return None
    out = {}
    with np.load(path, allow_pickle=False) as data:
        region_index = np.asarray(data["region_of_position"])
        position = np.asarray(data["position"])
        bg = io.group_csr(data, "bg")
        bg_carrying = np.asarray((bg > 0).sum(axis=0)).ravel()
        bg_total = np.asarray(bg.sum(axis=0)).ravel()
        n_bg = bg.shape[0]
        for direction in ("up", "dn"):
            fg = io.group_csr(data, direction)
            fg_carrying = np.asarray((fg > 0).sum(axis=0)).ravel()
            fg_total = np.asarray(fg.sum(axis=0)).ravel()
            n_fg = fg.shape[0]
            kmax = int(max(bg.data.max(initial=0), fg.data.max(initial=0)))
            _, _, _, z, _ = io.rank_statistics(io.csr_histograms(fg, kmax),
                                               io.csr_histograms(bg, kmax))
            p = io.p_from_z(z)
            for region_id, region in enumerate(io.REGIONS):
                slots = np.flatnonzero(region_index == region_id)
                best = int(slots[int(np.argmin(p[slots]))])
                fg_mean = float(fg_total[best]) / n_fg
                bg_mean = float(bg_total[best]) / n_bg
                fg_prop = float(fg_carrying[best]) / n_fg
                bg_prop = float(bg_carrying[best]) / n_bg
                out[(direction, region)] = {
                    "position": int(position[best]),
                    "fg_mean_count": fg_mean, "bg_mean_count": bg_mean,
                    "count_ratio": (fg_mean / bg_mean) if bg_mean > 0 else None,
                    "fg_proportion": fg_prop, "bg_proportion": bg_prop,
                    "enrichment_ratio": (fg_prop / bg_prop) if bg_prop > 0 else None,
                    "n_fg_carrying": int(fg_carrying[best]), "n_bg_carrying": int(bg_carrying[best]),
                }
    return out


def read_roots(engine_out: Path):
    roots = {d: io.read_root_table(engine_out / f"pVal.{d}.vs.bg.RNAmap.txt")
             for d in ("up", "dn")}
    motifs = sorted(roots["up"])
    if sorted(roots["dn"]) != motifs:
        raise ValueError("The two root tables do not carry the same motifs")
    return roots, motifs


def quick_tables(args, engine_out: Path, counts_root, counts: dict, alias: dict, scores=None):
    """One row per motif per direction, plus the per-RBP best-motif panel ranking."""
    roots, motifs = read_roots(engine_out)
    if scores is None:
        scores = {m: motif_scores(counts_root, args.arm, m) for m in motifs}
    n_events = {"up": counts["up"]["n_events"], "dn": counts["dn"]["n_events"],
                "bg": counts["bg"]["n_events"]}
    per_motif = {"up": [], "dn": []}
    for direction in ("up", "dn"):
        for motif in motifs:
            table_name = motif.split(".", 1)[0]
            row = {"arm": args.arm, "direction": direction,
                   "direction_label": io.DIRECTION_LABEL[direction],
                   "motif_key": motif, "rbp_table_name": table_name,
                   "RBP": alias.get(table_name, table_name),
                   "n_changed_events": n_events[direction], "n_background_events": n_events["bg"]}
            for region in io.REGIONS:
                row[f"p_{region}"] = roots[direction][motif][region]
            for pool, members in io.POOL_TO_REGIONS.items():
                best = min(members, key=lambda r: (roots[direction][motif][r], r))
                row[f"{pool} | min_p"] = roots[direction][motif][best]
                row[f"{pool} | sub_region"] = best
                score = (scores[motif] or {}).get((direction, best))
                for field in ("position", "fg_mean_count", "bg_mean_count", "count_ratio",
                              "fg_proportion", "bg_proportion", "enrichment_ratio"):
                    row[f"{pool} | {field}"] = None if score is None else score[field]
            per_motif[direction].append(row)
    rbp_rows = []
    for direction in ("up", "dn"):
        for pool in io.POOL_TO_REGIONS:
            best_by_rbp = {}
            for row in per_motif[direction]:
                key = row["RBP"]
                current = best_by_rbp.get(key)
                if current is None or (row[f"{pool} | min_p"], row["motif_key"]) < (
                        current[f"{pool} | min_p"], current["motif_key"]):
                    best_by_rbp[key] = row
            ordered = sorted(best_by_rbp.values(),
                             key=lambda r: (r[f"{pool} | min_p"], r["motif_key"]))
            for rank, row in enumerate(ordered, 1):
                rbp_rows.append({
                    "arm": args.arm, "direction_label": io.DIRECTION_LABEL[direction],
                    "pooled_region": pool, "rank": rank, "RBP": row["RBP"],
                    "rbp_table_name": row["rbp_table_name"],
                    "selected_motif_key": row["motif_key"],
                    "selected_sub_region": row[f"{pool} | sub_region"],
                    "min_p": row[f"{pool} | min_p"],
                    "argmin_position": row[f"{pool} | position"],
                    "fg_mean_count": row[f"{pool} | fg_mean_count"],
                    "bg_mean_count": row[f"{pool} | bg_mean_count"],
                    "count_ratio": row[f"{pool} | count_ratio"],
                    "enrichment_ratio": row[f"{pool} | enrichment_ratio"],
                    "n_changed_events": row["n_changed_events"],
                    "n_background_events": row["n_background_events"],
                    "plotted_pool": pool in io.PLOT_POOLS,
                })
    return per_motif, rbp_rows, motifs


def write_workbook(path: Path, sheets: dict) -> None:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, (columns, rows) in sheets.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([excel_value(row.get(c)) for c in columns])
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = (f"A1:{get_column_letter(len(columns))}{len(rows) + 1}")
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                44, max(12, len(str(column)) + 2))
    book.save(path)


def excel_value(value):
    if value is None:
        return "NA"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and not math.isfinite(value):
        return "NA"
    return value


def readme_rows(args, engine_out: Path, engine_root: Path, counts, conversion, motifs,
                known: Path, additional: Path) -> list:
    rows = [
        ("Arm", args.arm),
        ("Mode", args.mode),
        ("Species / assembly", f"{args.species} / {args.genome}; SE (skipped exon) events only"),
        ("Genome FASTA", str(Path(args.genome_root) / args.genome / f"{args.genome}.fa")),
        ("Engine", f"{args.engine} at {engine_root} @ {git_revision(engine_root)}"),
        ("Statistic", args.stat_method),
        ("Statistic caveat", STAT_CAVEATS[args.stat_method]),
        ("Geometry", f"window {args.window} / step {args.step} / intron {args.intron} / exon {args.exon}"),
        ("Motif tables", "{} + {}; {} motif keys scored".format(
            known.name, additional.name if additional else "no additional list",
            len(motifs))),
        ("n events", "changed included {n_up}, changed skipped {n_dn}, background {n_bg}".format(
            n_up=counts["up"]["n_events"], n_dn=counts["dn"]["n_events"],
            n_bg=counts["bg"]["n_events"])),
        ("Counted unit", "One row is one rMATS SE event, not one distinct target exon; an exon "
                         "paired with different flanking exons contributes more than one row."),
        ("Root tables", "pVal.{up,dn}.vs.bg.RNAmap.txt hold the smallest p over the 50-nt windows "
                        "of each sub-region. They are read here, never recomputed."),
        ("Pooled regions", "Upstream Intron = min(UpstreamExonIntron, UpstreamIntron); Exon Body = "
                           "min(TargetExon_5prime, TargetExon-3prime); Downstream Intron = "
                           "min(DownstreamIntron, DownstreamExonIntron); Flanking Exon kept as a "
                           "column and not plotted."),
        ("Motif score columns", "The engine's own countDist numbers at the window that produced "
                                "that sub-region's minimum p: fg/bg_mean_count are hits per event "
                                "per window, count_ratio is their quotient, fg/bg_proportion are "
                                "the fractions of events carrying the motif at all. NA when the "
                                "count archives were not built."),
        ("Multiple testing", "NONE is applied here. The tool reports raw p; this workbook reports "
                             "what the tool reported."),
        ("Archive verification", conversion.get("reason") or
         ("archives reproduce the root tables" if conversion.get("verified") else "not run")),
        ("Direction", "up = the exon is MORE included in the treatment group; dn = LESS included"),
    ]
    if args.gate_note:
        rows.append(("Gate", args.gate_note))
    return [{"field": k, "description": v} for k, v in rows]


# ------------------------------------------------------------------ positive control
def positive_control(rbp_rows, symbol: str, panels) -> list:
    out = []
    for label, pool in panels:
        ordered = [r for r in rbp_rows
                   if r["direction_label"] == label and r["pooled_region"] == pool]
        ordered.sort(key=lambda r: r["rank"])
        top = ordered[0] if ordered else None
        hit = next((r for r in ordered if r["RBP"] == symbol), None)
        out.append({"symbol": symbol, "panel": f"{label} x {pool}",
                    "top_rbp": None if top is None else top["RBP"],
                    "top_p": None if top is None else top["min_p"],
                    "symbol_rank": None if hit is None else hit["rank"],
                    "symbol_p": None if hit is None else hit["min_p"],
                    "symbol_motif": None if hit is None else hit["selected_motif_key"],
                    "pass": bool(hit is not None and hit["rank"] == 1)})
    return out


# ------------------------------------------------------------------ quick-mode figures
MAIN_LAYER = "released_ranksum_rawP"
FIGURE_KINDS = ("byRBP", "byMotif")
FIGURE_VARIANTS = ("main", "noSpliceosome_noBroad")
SCORE_TABLE_COLUMNS = [
    "arm", "motif_key", "RBP", "rbp_table_name", "direction", "direction_label", "region",
    "pooled_region", "plot", "native_ranksum_p", "argmin_window_position",
    "n_fg_carrying", "n_bg_carrying", "fg_proportion", "bg_proportion", "enrichment_ratio",
    "fg_mean_count", "bg_mean_count", "count_ratio",
]


def figures_wanted(args) -> bool:
    return args.mode == "quick" and not args.no_figures


def figure_skip_reason(args, conversion) -> str:
    if args.engine != "released" or args.stat_method != "mannwhitney":
        return ("the main-layer figures draw the released engine's rank-sum p; this run is "
                f"--engine {args.engine} --stat-method {args.stat_method}")
    if not conversion.get("verified"):
        return "dot size needs verified count archives; " + (conversion.get("reason") or "not verified")
    return ""


def figure_stem(arm: str, kind: str, variant: str) -> str:
    return f"{arm}_SE_{kind}_{MAIN_LAYER}" + ("" if variant == "main" else "_" + variant)


def cell(value) -> str:
    if value is None:
        return "NA"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return repr(value) if math.isfinite(value) else "NA"
    return str(value)


def write_score_table(path: Path, arm: str, roots: dict, motifs: list, scores: dict,
                      alias: dict) -> int:
    """The tool's own per-sub-region motif scores, in the column names the figure builder reads.

    Released-layer columns only: no calibrated p, no q. The builder sizes dots by count_ratio and
    skips its calibrated layer because no refinement_report.json sits beside this table.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with open(path, "w", encoding="utf-8", newline="\n") as handle:
        handle.write("\t".join(SCORE_TABLE_COLUMNS) + "\n")
        for motif in motifs:
            table_name = motif.split(".", 1)[0]
            for direction in ("up", "dn"):
                for region in io.REGIONS:
                    score = (scores.get(motif) or {}).get((direction, region))
                    if score is None:
                        raise ValueError(f"No motif score for {motif} {direction} {region}")
                    pool = io.REGION_TO_POOL[region]
                    row = {"arm": arm, "motif_key": motif, "RBP": alias.get(table_name, table_name),
                           "rbp_table_name": table_name, "direction": direction,
                           "direction_label": io.DIRECTION_LABEL[direction], "region": region,
                           "pooled_region": pool, "plot": pool in io.PLOT_POOLS,
                           "native_ranksum_p": roots[direction][motif][region],
                           "argmin_window_position": score["position"],
                           **{k: score[k] for k in SCORE_TABLE_COLUMNS[11:]}}
                    handle.write("\t".join(cell(row[c]) for c in SCORE_TABLE_COLUMNS) + "\n")
                    n += 1
    return n


def figure_positive_control(selections: dict, arm: str, symbol: str, panels) -> list:
    """Rank-1 check on every drawn main-layer panel: both kinds, both exclusion variants."""
    out = []
    for variant in FIGURE_VARIANTS:
        for kind in FIGURE_KINDS:
            drawn_panels = selections[(variant, kind)]
            for label, pool in panels:
                drawn = drawn_panels.get((label, pool), [])
                first = drawn[0] if drawn else None
                rank = next((i for i, r in enumerate(drawn, 1) if r["_label"] == symbol), None)
                out.append({"arm": arm, "layer": MAIN_LAYER, "variant": variant, "kind": kind,
                            "panel": f"{label} x {pool}", "symbol": symbol,
                            "first": first["_label"] if first else None,
                            "first_p": first["_p"] if first else None,
                            "first_motif": first["motif_key"] if first else None,
                            "symbol_rank_in_drawn_top_n": rank,
                            "pass": bool(first is not None and first["_label"] == symbol)})
    return out


def build_quick_figures(args, out: Path, engine_out: Path, roots: dict, motifs: list,
                        scores: dict, alias: dict, known: Path, additional, log) -> dict:
    """Main-layer region lollipops (released rank-sum, raw p) drawn by the v4 builder's own
    functions: naming, released-table reader, motif scores, shared y-scale, panel selection and
    draw_figure with its layout audit. The builder's main() is not used because its
    rank-comparison step needs a second layer or a v3.1 archive; nothing in it is modified."""
    import matplotlib.pyplot as plt
    import build_region_lollipops_v4 as lol
    arm = args.arm
    figures = out / "figures"
    score_root = figures / "motif_scores"
    score_table = score_root / arm / "per_motif_regions.tsv"
    n_rows = write_score_table(score_table, arm, roots, motifs, scores, alias)
    (score_root / arm / "command.log").write_text(
        f"source=count archives {out / 'counts' / arm} (verified against the root tables)\n"
        f"writer={Path(__file__).resolve()} rows={n_rows} layer=released motif scores only\n"
        "exit=0\n", encoding="utf-8")
    log(f"wrote {score_table} ({n_rows} rows; released-layer motif scores, no calibrated columns)")
    plt.rcParams.update({"font.family": "Arial", "svg.fonttype": "none",
                         "svg.hashsalt": "rmaps-v4", "pdf.fonttype": 42})
    if args.arm_label:
        lol.LABELS[arm] = args.arm_label
    lol.LABELS.setdefault(arm, arm.replace("_", " "))
    mappings, naming_rows, naming_stats = lol.load_naming(
        known, additional or lol.ESRP_TABLE, args.gtf, args.alias_table)
    aliases = naming_stats["alias_mappings"]
    lists = {"spliceosome_census": lol.load_exclusion_list(args.spliceosome_list),
             "broad_binders": lol.load_exclusion_list(args.broad_binders_list)}
    lists = {k: {aliases.get(s, s) for s in v} for k, v in lists.items()}
    counts, _ = lol.gate_counts(None, arm, None, out / "event_counts.json")
    engine_root = engine_root_of(args)
    entries, _, n_motifs = lol.released_entries(arm, engine_out.parent, mappings,
                                                git_revision(engine_root)[:7], args.stat_method)
    exclusion_rows, dropped = lol.exclusion_audit(arm, entries, lists)
    lookup, size_key = lol.load_motif_scores(arm, score_root)
    lol.attach_scores(entries, lookup, arm)
    scale = lol.y_scale(entries, MAIN_LAYER)
    power = None
    if arm in set(args.underpowered_arms):
        power = {"label": f"underpowered: {counts['n_up']} included and {counts['n_dn']} skipped events",
                 "adequate": False}
    selections, audits, layouts, built = {}, [], [], []
    for variant in FIGURE_VARIANTS:
        chosen = entries if variant == "main" else [e for e in entries
                                                    if e["_exclusion_symbol"] not in dropped]
        for kind in FIGURE_KINDS:
            panels = lol.select_panels(chosen, MAIN_LAYER, kind, args.top_n)
            selections[(variant, kind)] = panels
            for (direction, region), rows in panels.items():
                for rank, r in enumerate(rows, 1):
                    audits.append({"arm": arm, "layer": MAIN_LAYER, "variant": variant, "type": kind,
                                   "direction": direction, "pooled_region": region, "rank": rank,
                                   "label": r["_label"], "figure_tick_label": r["_tick"],
                                   "source_rbp": r["RBP"], "motif_key": r["motif_key"],
                                   "selected_sub_region": r["region"], "p": r["_p"],
                                   "count_ratio": r["_count_ratio"],
                                   "k_motifs_p_lt_0.05": r["_k"], "n_motifs_in_group": r["_n"],
                                   "hgnc_symbol": r["_hgnc_symbol"],
                                   "naming_action": r["_naming_action"]})
            stem = figures / figure_stem(arm, kind, variant)
            report = lol.draw_figure(arm, MAIN_LAYER, kind, panels, counts, None, scale, power,
                                     size_key, stem, args, excluded=variant != "main")
            layouts.append(report)
            built += [Path(str(stem) + ".png"), Path(str(stem) + ".svg")]
    for path in built:
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"Figure missing after drawing: {path}")
    lol.write_tsv(figures / "selection_audit.tsv", list(audits[0]), audits)
    lol.write_tsv(figures / "exclusion_audit.tsv", list(exclusion_rows[0]), exclusion_rows)
    lol.write_tsv(figures / "naming_audit.tsv",
                  ["table_name", "hgnc_symbol", "source", "evidence", "ambiguity_note", "n_motifs",
                   "merged_into"], naming_rows)
    (figures / "layout_report.json").write_text(json.dumps(
        {"figure_version": lol.FIG_VERSION, "layer": MAIN_LAYER, "n_motifs": n_motifs,
         "y_scale": scale, "size_key": size_key, "dropped_in_noSpliceosome_noBroad": sorted(dropped),
         "figures": layouts}, indent=2, default=str) + "\n", encoding="utf-8")
    result = {"figures": built, "score_table": score_table, "controls": [],
              "audits": [figures / n for n in ("selection_audit.tsv", "exclusion_audit.tsv",
                                               "naming_audit.tsv", "layout_report.json")]}
    if args.positive_control:
        panels = [tuple(part.split(":", 1)) for part in args.positive_control_panels.split(",")]
        controls = figure_positive_control(selections, arm, args.positive_control, panels)
        lol.write_tsv(figures / "positive_control_audit.tsv", list(controls[0]), controls)
        result["controls"] = controls
        for row in controls:
            log("figure positive control {} {} {} {}: first={} -> {}".format(
                row["symbol"], row["variant"], row["kind"], row["panel"], row["first"],
                "PASS" if row["pass"] else "FAIL"))
    log(f"figures: {len(built)} files in {figures}; y limit {scale['ymax']:g}, "
        f"{sum(r['n_stems_truncated'] for r in layouts)} truncated stems, layout audit passed")
    return result


# ------------------------------------------------------------------ index page
def write_index(args, out: Path, engine_out: Path, counts, controls, conversion, extras,
                figures=None) -> Path:
    import html as html_mod

    def esc(value):
        return html_mod.escape(str(value))

    def rel(path):
        try:
            return os.path.relpath(Path(path), out).replace("\\", "/")
        except ValueError:
            return Path(path).resolve().as_uri()

    parts = [
        '<!doctype html><html lang="en"><meta charset="utf-8">',
        f"<title>rMAPS3 {esc(args.arm)} — {esc(args.mode)} run</title>",
        "<style>body{font-family:Segoe UI,Arial,sans-serif;margin:2rem;max-width:1100px;color:#222}"
        "table{border-collapse:collapse;margin:1rem 0}td,th{border:1px solid #CCC;padding:4px 10px;"
        "text-align:left}th{background:#F2F2F2}code{background:#F6F6F6;padding:1px 4px}"
        ".pass{color:#046A38;font-weight:bold}.fail{color:#D55E00;font-weight:bold}"
        "img{max-width:100%;border:1px solid #DDD}</style>",
        f"<h1>rMAPS3 SE motif map — {esc(args.arm)}</h1>",
        f"<p>Mode <b>{esc(args.mode)}</b>; engine <b>{esc(args.engine)}</b>; statistic "
        f"<b>{esc(args.stat_method)}</b>; {esc(args.species)} / {esc(args.genome)}; SE events only.</p>",
        f"<p><b>{esc(STAT_CAVEATS[args.stat_method])}</b></p>",
        "<h2>Event sets</h2><table><tr><th>set</th><th>n events</th><th>file</th></tr>",
    ]
    for key, label in (("up", "changed, more included"), ("dn", "changed, more skipped"),
                       ("bg", "background")):
        parts.append(f"<tr><td>{esc(label)}</td><td>{counts[key]['n_events']}</td>"
                     f"<td><code>{esc(counts[key]['output']['path'])}</code></td></tr>")
    parts.append("</table>")
    parts.append("<h2>The tool's own tables</h2><ul>")
    for direction in ("up", "dn"):
        path = engine_out / f"pVal.{direction}.vs.bg.RNAmap.txt"
        parts.append(f'<li><a href="{rel(path)}">pVal.{direction}.vs.bg.RNAmap.txt</a> — '
                     "smallest p per motif per sub-region, raw, as the tool reports it</li>")
    parts.append(f'<li><a href="quick_summary.xlsx">quick_summary.xlsx</a> — README, one sheet per '
                 "direction, and the per-RBP panel ranking</li></ul>")
    if controls:
        parts.append("<h2>Positive control</h2><table><tr><th>panel</th><th>top RBP</th>"
                     "<th>control rank</th><th>verdict</th></tr>")
        for row in controls:
            klass = "pass" if row["pass"] else "fail"
            parts.append(f"<tr><td>{esc(row['panel'])}</td><td>{esc(row['top_rbp'])}</td>"
                         f"<td>{esc(row['symbol_rank'])}</td>"
                         f"<td class=\"{klass}\">{'PASS' if row['pass'] else 'FAIL'}</td></tr>")
        parts.append("</table>")
    if figures is not None and figures.get("figures"):
        parts.append(
            "<h2>Main-layer region lollipops</h2>"
            "<p><b>What they show:</b> the authors' released rMAPS3 rank-sum p, raw, exactly as the "
            "root tables report it, reduced to the smallest p per pooled region. Stems are "
            "−log10 p, dot colour bins the same raw p, dot size is the tool's own motif-score "
            "ratio (changed ÷ background hits per event per 50-nt window). Included panels rise, "
            "skipped panels hang down, all six share one y-scale, and the footer gives n events.</p>"
            "<p><b>What they do not show:</b> a calibrated p, a q-value or any multiple-testing "
            "adjustment. The p is anti-conservative; <b>use them for the RBP ORDER only</b>. "
            "The calibrated supplement, the p and q to report, is built by <code>--mode full</code>.</p>"
            '<div style="display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px">')
        for path in figures["figures"]:
            if path.suffix != ".png":
                continue
            svg = path.with_suffix(".svg")
            parts.append(f'<div><a href="{rel(path)}"><img loading="lazy" src="{rel(path)}" '
                         f'alt="{esc(path.name)}"></a><br>{esc(path.stem)} · '
                         f'<a href="{rel(svg)}">editable SVG</a></div>')
        parts.append("</div>")
        links = [("selection audit: every drawn row, its motif, sub-region, p and ratio",
                  out / "figures" / "selection_audit.tsv"),
                 ("layout, y-scale, dot-size and truncation record", out / "figures" / "layout_report.json"),
                 ("exclusions applied in the _noSpliceosome_noBroad variant",
                  out / "figures" / "exclusion_audit.tsv"),
                 ("HGNC naming audit", out / "figures" / "naming_audit.tsv"),
                 ("motif-score table used for dot size", figures["score_table"])]
        if figures.get("controls"):
            links.insert(0, ("figure positive-control audit",
                             out / "figures" / "positive_control_audit.tsv"))
        parts.append("<ul>" + "".join(f'<li><a href="{rel(t)}">{esc(l)}</a></li>'
                                      for l, t in links) + "</ul>")
    elif figures is not None and figures.get("skipped"):
        parts.append(f"<h2>Main-layer region lollipops</h2><p>Not built: {esc(figures['skipped'])}</p>")
    maps_dir = engine_out / "maps"
    if maps_dir.is_dir():
        pngs = sorted(maps_dir.glob("*.png"))
        parts.append(f"<h2>The authors' own per-motif RNA maps</h2><p>Unmodified engine output: "
                     f'<a href="{rel(maps_dir)}">{esc(maps_dir)}</a> ({len(pngs)} PNG).</p>')
        if args.positive_control:
            for png in pngs:
                if png.stem.split(".", 1)[-1].split("-", 1)[0].upper() == args.positive_control.upper():
                    parts.append(f'<p><a href="{rel(png)}"><img loading="lazy" src="{rel(png)}" '
                                 f'alt="{esc(png.name)}"></a><br>{esc(png.name)} — the authors\' '
                                 "map for the positive control, unmodified</p>")
                    break
    for heading, items in extras:
        parts.append(f"<h2>{esc(heading)}</h2><ul>")
        for label, target in items:
            parts.append(f'<li><a href="{rel(target)}">{esc(label)}</a></li>'
                         if target else f"<li>{esc(label)}</li>")
        parts.append("</ul>")
    parts.append("<h2>Provenance</h2><ul>"
                 '<li><a href="command.log">command.log</a></li>'
                 '<li><a href="versions.txt">versions.txt</a></li>'
                 '<li><a href="md5.txt">md5.txt</a></li>'
                 '<li><a href="run_manifest.json">run_manifest.json</a></li></ul>')
    if conversion.get("reason"):
        parts.append(f"<p>Count archives: {esc(conversion['reason'])}</p>")
    parts.append("</html>")
    path = out / "index.html"
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")
    return path


# ------------------------------------------------------------------ full-mode layers
def target_exon_line(report_path: Path) -> str:
    """'n events (rMATS SE rows) over N target exons' for the three sets, from the calibration report."""
    report = json.loads(Path(report_path).read_text(encoding="utf-8"))
    ev, tx = report["events"], report["target_exons"]
    return ("n events (rMATS SE rows) over N target exons: included {} over {}; skipped {} over {}; "
            "background {} over {}; target exons in both changed sets (kept in both, as the tool does) {}"
            .format(ev["up"], tx["up"], ev["dn"], tx["dn"], ev["bg"], tx["bg"],
                    report["target_exons_in_both_changed_foregrounds"]))


def run_full_layers(args, out: Path, engine_out: Path, counts_root, log: Logger, env: dict):
    """Row-unit sensitivity, the reportable calibration v2.1, optional rank stability, v4 figures."""
    produced = []
    if counts_root is None:
        raise RuntimeError("--mode full needs the count archives; the countDist temporaries were "
                           "absent, so the calibration layer cannot be built")
    if args.stat_method != "mannwhitney":
        raise RuntimeError("--mode full calibrates the rank-sum statistic; rerun with "
                           "--stat-method mannwhitney or use --mode quick")
    common = ["--arm", args.arm, "--counts-root", str(counts_root),
              "--released-root", str(engine_out.parent), "--alias-table", str(args.alias_table),
              "--permutations", str(args.permutations), "--refine-perms", str(args.refine_perms),
              "--refine-threshold", str(args.refine_threshold), "--seed", str(args.seed)]
    rowunit_root = out / "summary_rowunit"
    run_step("calibrate_ranksum_rowunit_sensitivity",
             [sys.executable, str(ROOT / "tools" / "calibrate_ranksum.py"), *common,
              "--out-root", str(rowunit_root), "--permutation-unit", "row"], log, env, ROOT, out)
    summary_root = out / "summary"
    command = [sys.executable, str(ROOT / "tools" / "calibrate_ranksum_v2.py"), *common,
               "--out-root", str(summary_root), "--rowunit-root", str(rowunit_root)]
    if args.arm in set(args.underpowered_arms):
        command += ["--underpowered-arms", args.arm]
    run_step("calibrate_ranksum_v2", command, log, env, ROOT, out)
    report_path = summary_root / args.arm / "refinement_report.json"
    counts_line = target_exon_line(report_path)
    log(counts_line)
    produced.append(("Calibrated supplement: the p and q to report (target-exon cluster permutation, "
                     "RBP-level min-P)",
                     [(counts_line, None),
                      ("calibration workbook", summary_root / args.arm / f"{args.arm}_calibrated_ranksum_v2.xlsx"),
                      ("RBP-level table (min-P primary; max-z, mean-z sensitivity)",
                       summary_root / args.arm / "rbp_level.tsv"),
                      ("per-motif regions TSV", summary_root / args.arm / "per_motif_regions.tsv"),
                      ("readout", summary_root / args.arm / "readout.md"),
                      ("row-unit sensitivity (not reportable: duplicate target exons make the row an "
                       "invalid permutation unit)", rowunit_root / args.arm / "readout.md")]))

    stability_dir = None
    stability_run = Path(args.rank_stability_run) if args.rank_stability_run else (
        engine_out if args.engine == "audited" else None)
    if stability_run is not None and (stability_run / "positional").is_dir():
        stability_dir = out / "stability"
        stability_dir.mkdir(parents=True, exist_ok=True)
        run_step("rank_stability",
                 [sys.executable, str(ROOT / "tools" / "rank_stability.py"),
                  "--run", str(stability_run), "--arm", args.arm, "--out", str(stability_dir),
                  "--boot", str(args.bootstraps), "--seed", str(args.seed),
                  "--aliases", str(args.alias_table)], log, env, ROOT, out)
        produced.append(("Rank stability (foreground bootstrap)",
                         [("stability workbook", stability_dir / f"{args.arm}_rank_stability.xlsx")]))
    else:
        log("SKIP rank_stability: it reads positional/*.hits.npz, which only the audited engine "
            "writes; pass --rank-stability-run <audited engine dir> to include it")

    figures = out / "figures"
    command = [sys.executable, str(ROOT / "tools" / "build_region_lollipops_v4.py"),
               "--arms", args.arm, "--out-root", str(figures),
               "--released-root", str(engine_out.parent),
               "--calibrated-root", str(summary_root),
               "--counts-json", f"{args.arm}={out / 'event_counts.json'}",
               "--alias-table", str(args.alias_table), "--gtf", str(args.gtf),
               "--spliceosome-list", str(args.spliceosome_list),
               "--broad-binders-list", str(args.broad_binders_list),
               "--author-maps-root", str(engine_out.parent),
               "--released-commit", git_revision(engine_root_of(args))[:7],
               "--released-stat-method", args.stat_method,
               "--top-n", str(args.top_n)]
    if args.arm in set(args.underpowered_arms):
        command += ["--underpowered-arms", args.arm]
    if args.method_comparison:
        command += ["--method-comparison", str(args.method_comparison)]
    for flag, value in (("--gate-text", args.gate_text), ("--gate-record", args.gate_record),
                        ("--direction-text", args.direction_text), ("--tail-text", args.tail_text)):
        if value:
            command += [flag, str(value)]
    run_step("figures_v" + figure_version(), command, log, env, ROOT, out)
    produced.append(("Region lollipops, figure version " + figure_version(),
                     [("figure index", figures / "index.html"),
                      ("rank workbook", figures / args.arm / f"{args.arm}_rank_comparison.xlsx")]))
    return produced


# ------------------------------------------------------------------ main
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--mode", choices=("quick", "full"), required=True)
    p.add_argument("--arm", required=True, help="filename-safe run name")
    p.add_argument("--out", required=True, help="output root; must be absent or empty")
    p.add_argument("--rmats-se", help="SE.MATS.JC.txt to pre-split with --filter")
    p.add_argument("--filter", help="json/yaml gate spec: fdr, min_abs_dpsi, min_jc_per_sample, "
                                    "bg_fdr_min, expr_table, base_mean_floor")
    p.add_argument("--up")
    p.add_argument("--dn")
    p.add_argument("--bg")
    p.add_argument("--gate-rule", choices=GATE_RULES, default=None,
                   help="event-set rule the inputs were built under; must equal the arm suffix. Raw "
                        "--rmats-se input is split by the portable rule-A builder, so it takes rule A only")
    p.add_argument("--genome-root", default=None,
                   help="required: directory holding <genome>/<genome>.fa and its .fai")
    p.add_argument("--genome", default="hg38", help="FASTA build directory name, e.g. hg38 or mm10")
    p.add_argument("--species", default="Homo sapiens")
    p.add_argument("--engine", choices=("released", "audited"), default="released")
    p.add_argument("--engine-root", default=None,
                   help=f"required for --engine released; audited default {ROOT}")
    p.add_argument("--stat-method", choices=("fisher", "mannwhitney"), default="mannwhitney")
    p.add_argument("--known-motifs", default=None)
    p.add_argument("--additional-motifs", default=None)
    p.add_argument("--window", type=int, default=50)
    p.add_argument("--step", type=int, default=1)
    p.add_argument("--intron", type=int, default=250)
    p.add_argument("--exon", type=int, default=50)
    p.add_argument("--workers", type=int, default=1, help="audited engine only; 1 avoids the "
                                                          "Windows DuplicateHandle spawn failure")
    p.add_argument("--blas-threads", type=int, default=4)
    p.add_argument("--alias-table", default=str(DEFAULT_ALIAS))
    p.add_argument("--positive-control", default=None,
                   help="HGNC symbol expected to rank first in the control panels")
    p.add_argument("--positive-control-panels", default="INCLUDED:Upstream Intron,SKIPPED:Downstream Intron")
    p.add_argument("--positive-control-advisory", action="store_true",
                   help="record a failed positive control instead of exiting non-zero")
    p.add_argument("--keep-temp", action="store_true",
                   help="keep the countDist temporaries even after they verify")
    p.add_argument("--gate-note", default=None, help="free text describing the gate, for the README")
    p.add_argument("--gate-counts", default=None,
                   help="the gate record counts.json for pre-split inputs (n_up, n_dn, n_bg, "
                        "n_expr_unknown_in_fg/bg); its event counts must match the inputs")
    p.add_argument("--no-figures", action="store_true",
                   help="quick mode: skip the main-layer region lollipops")
    p.add_argument("--arm-label", default=None, help="figure title for an arm the builder does not name")
    # full mode
    p.add_argument("--calib-unit", choices=("cluster",), default="cluster",
                   help="permutation unit of the reportable calibration: the target-exon cluster; the "
                        "row unit is always run beside it as a sensitivity column")
    p.add_argument("--permutations", type=int, default=2000)
    p.add_argument("--refine-perms", type=int, default=100000)
    p.add_argument("--refine-threshold", type=float, default=0.005)
    p.add_argument("--seed", type=int, default=149)
    p.add_argument("--bootstraps", type=int, default=300)
    p.add_argument("--rank-stability-run", default=None,
                   help="audited engine directory carrying positional/*.hits.npz")
    p.add_argument("--gtf", default=None,
                   help="GENCODE GTF for figure naming; required whenever figures are drawn")
    p.add_argument("--spliceosome-list", default=None,
                   help="core-spliceosome list for the _noSpliceosome_noBroad variant; required with figures")
    p.add_argument("--broad-binders-list", default=None,
                   help="broad-binder list for the _noSpliceosome_noBroad variant; required with figures")
    p.add_argument("--method-comparison", default=None)
    p.add_argument("--underpowered-arms", nargs="*", default=[])
    p.add_argument("--top-n", type=int, default=10)
    p.add_argument("--gate-text", default=None,
                   help="full mode: figure footer gate text (builder --gate-text; default = the dissertation "
                        "reli_v121 gate sentence)")
    p.add_argument("--gate-record", default=None,
                   help="full mode: JSON with gate_text / direction_text / tail_text / tail_text_supplement")
    p.add_argument("--direction-text", default=None, help="full mode: figure legend direction line")
    p.add_argument("--tail-text", default=None, help="full mode: figure bottom tail line")
    return p


def validate(args) -> None:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", args.arm) or ".." in args.arm:
        raise ValueError("--arm must be a filename-safe name")
    given = [key for key in ("up", "dn", "bg") if getattr(args, key)]
    if given and len(given) != 3:
        raise ValueError("--up, --dn and --bg must be given together")
    presplit = len(given) == 3
    from_rmats = bool(args.rmats_se)
    if presplit == from_rmats:
        raise ValueError("Give either --up/--dn/--bg or --rmats-se with --filter, not both")
    if from_rmats and not args.filter:
        raise ValueError("--rmats-se needs --filter")
    if args.mode == "full" and args.stat_method != "mannwhitney":
        raise ValueError("--mode full calibrates the rank-sum statistic; it needs "
                         "--stat-method mannwhitney, or use --mode quick")
    if args.mode == "full" and args.engine != "released":
        raise ValueError("--mode full builds the authors' layer plus its calibration, so it needs "
                         "--engine released")
    draws = args.mode == "full" or (figures_wanted(args) and args.engine == "released"
                                    and args.stat_method == "mannwhitney")
    if draws and "_" not in args.arm:
        raise ValueError("the figure builder reads the gate rule from the text after the arm's last "
                         "'_' (e.g. QKI_KO_B -> rule B); name the arm <NAME>_<RULE> or pass --no-figures")
    if draws and presplit and not args.gate_counts:
        raise ValueError("the figure footer prints the gate's expression-unknown event counts, which "
                         "pre-split files do not carry; pass --gate-counts <counts.json> (the gate "
                         "record <event_sets>/<ARM>/<gate>_rule<R>/counts.json) or --no-figures")
    if args.gate_counts and not presplit:
        raise ValueError("--gate-counts is for pre-split inputs; --rmats-se writes its own gate record")
    check_gate_rule(args, from_rmats)
    for key in ("window", "step", "intron", "exon", "workers", "blas_threads", "permutations",
                "refine_perms", "bootstraps", "top_n"):
        if getattr(args, key) < 1:
            raise ValueError(f"--{key.replace('_', '-')} must be positive")
    if args.genome.lower().startswith("mm") or "mouse" in args.species.lower():
        print("*" * 78, file=sys.stderr)
        print(f"MOUSE REFERENCE IN PLAY: genome={args.genome} species={args.species}. "
              "Do not merge or compare this run with a human arm.", file=sys.stderr)
        print("*" * 78, file=sys.stderr)


def arm_rule(arm: str):
    """The rule the figure builder reads from the arm name: the text after the last '_', or None."""
    return arm.rsplit("_", 1)[1] if "_" in arm else None


def check_gate_rule(args, from_rmats: bool) -> None:
    """The effective gate rule must agree with the arm name, so no set is labelled with a rule it
    was not built under. Raw rMATS input only ever goes through the portable rule-A builder."""
    suffix = arm_rule(args.arm)
    if from_rmats and suffix in RULE_B_SUFFIXES:
        raise ValueError(f"arm {args.arm} names rule {suffix}, but --rmats-se input is split by the portable "
                         "rule-A builder. Rule B sets are frozen concordant files: supply pre-split inputs "
                         "(--up/--dn/--bg with --gate-counts) built under rule B")
    if from_rmats and args.gate_rule not in (None, "A"):
        raise ValueError(f"--gate-rule {args.gate_rule} with --rmats-se: the portable builder implements rule A "
                         "only. Rule B sets are frozen concordant files: supply pre-split inputs")
    if args.gate_rule is not None and suffix != args.gate_rule:
        raise ValueError(f"--gate-rule {args.gate_rule} disagrees with the arm name {args.arm} (rule "
                         f"{suffix}); name the arm <NAME>_{args.gate_rule}")


def require_site_paths(args) -> None:
    """Site paths have no default; refuse before any output exists when a needed one is missing."""
    missing = []
    if not args.genome_root:
        missing.append("--genome-root")
    if args.engine == "released" and not args.engine_root:
        missing.append("--engine-root (checkout of the released engine)")
    draws = args.mode == "full" or (figures_wanted(args) and args.engine == "released"
                                    and args.stat_method == "mannwhitney")
    if draws:
        missing += [flag for flag, value in (("--gtf", args.gtf),
                                              ("--spliceosome-list", args.spliceosome_list),
                                              ("--broad-binders-list", args.broad_binders_list))
                    if not value]
    if missing:
        raise ValueError("missing site path(s): " + ", ".join(missing)
                         + (" (figures need the naming inputs; --no-figures skips them)" if draws else ""))


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        validate(args)
        require_site_paths(args)
    except ValueError as exc:
        parser.exit(2, f"ERROR: {exc}\n")
    out = Path(args.out).resolve()
    fresh_directory(out)
    (out / "logs").mkdir(exist_ok=True)
    log = Logger(out / "command.log")
    started = time.perf_counter()
    manifest = {"schema_version": 1, "status": "running", "mode": args.mode, "arm": args.arm,
                "started_utc": datetime.now(timezone.utc).isoformat(),
                "command": subprocess.list2cmdline(sys.orig_argv),
                "args": {k: (str(v) if isinstance(v, Path) else v) for k, v in vars(args).items()},
                "wrapper_md5": md5(Path(__file__).resolve()),
                "effective_gate_rule": "A" if args.rmats_se else (args.gate_rule or arm_rule(args.arm)),
                "fork_revision": git_revision(ROOT), "versions": versions(), "steps": []}
    (out / "versions.txt").write_text(
        "\n".join(f"{k}\t{v}" for k, v in manifest["versions"].items())
        + f"\nfork_revision\t{manifest['fork_revision']}"
        + f"\nwrapper\t{Path(__file__).resolve()}"
        + f"\nwrapper_md5\t{manifest['wrapper_md5']}"
        + f"\nseed\t{args.seed}"
        + f"\ngenerated\t{time.strftime('%Y-%m-%dT%H:%M:%S')}\n", encoding="utf-8")
    log("command: " + manifest["command"])
    env = os.environ.copy()
    for key in BLAS_VARS:
        env[key] = str(args.blas_threads)
    env["RMAPS_FORCE_MOTIF_FALLBACK"] = "1"
    env["PYTHONPATH"] = os.pathsep.join(filter(None, [str(ROOT / "tools"), env.get("PYTHONPATH", "")]))
    exit_code = 0
    try:
        paths, counts, fasta, fai = prepare_inputs(args, out, log, env)
        engine_out, wall, engine_root, known, additional = run_engine(args, paths, out, log, env)
        counts_root, conversion = convert_and_verify(args, engine_out, out, log, env)
        alias = load_alias(Path(args.alias_table))
        roots, root_motifs = read_roots(engine_out)
        scores = {m: motif_scores(counts_root, args.arm, m) for m in root_motifs}
        per_motif, rbp_rows, motifs = quick_tables(args, engine_out, counts_root, counts, alias,
                                                   scores=scores)
        sheets = {
            "README": (["field", "description"],
                       readme_rows(args, engine_out, engine_root, counts, conversion, motifs,
                                   known, additional)),
        }
        for direction in ("up", "dn"):
            columns = list(per_motif[direction][0])
            sheets[io.DIRECTION_LABEL[direction]] = (columns, per_motif[direction])
        sheets["RBP_best_motif"] = (list(rbp_rows[0]), rbp_rows)
        controls = []
        if args.positive_control:
            panels = [tuple(part.split(":", 1)) for part in args.positive_control_panels.split(",")]
            controls = positive_control(rbp_rows, args.positive_control, panels)
            sheets["positive_control"] = (list(controls[0]), controls)
            for row in controls:
                log("positive control {} {}: top={} rank={} -> {}".format(
                    row["symbol"], row["panel"], row["top_rbp"], row["symbol_rank"],
                    "PASS" if row["pass"] else "FAIL"))
        write_workbook(out / "quick_summary.xlsx", sheets)
        log(f"wrote {out / 'quick_summary.xlsx'}")
        extras = []
        if args.mode == "full":
            extras = run_full_layers(args, out, engine_out, counts_root, log, env)
        figures = None
        if figures_wanted(args):
            reason = figure_skip_reason(args, conversion)
            if reason:
                log("SKIP figures: " + reason)
                figures = {"skipped": reason}
            else:
                figures = build_quick_figures(args, out, engine_out, roots, root_motifs, scores,
                                              alias, known, additional, log)
        write_index(args, out, engine_out, counts, controls, conversion, extras, figures)
        figure_controls = (figures or {}).get("controls", [])
        manifest.update(status="complete", engine_wall_seconds=wall,
                        n_motifs=len(motifs), conversion=conversion,
                        positive_control=controls, figure_positive_control=figure_controls,
                        figures=[str(p) for p in (figures or {}).get("figures", [])],
                        figures_skipped=(figures or {}).get("skipped"),
                        event_counts={k: counts[k]["n_events"] for k in counts})
        checks = controls + figure_controls
        if checks and not all(row["pass"] for row in checks) and not args.positive_control_advisory:
            manifest["status"] = "complete_with_failed_positive_control"
            exit_code = 3
    except Exception as exc:  # recorded, then re-raised through the exit code
        manifest.update(status="failed", error=f"{type(exc).__name__}: {exc}")
        log(f"FAILED {type(exc).__name__}: {exc}")
        exit_code = 1
    finally:
        manifest["finished_utc"] = datetime.now(timezone.utc).isoformat()
        manifest["wall_seconds"] = time.perf_counter() - started
        (out / "run_manifest.json").write_text(json.dumps(manifest, indent=2, default=str) + "\n",
                                               encoding="utf-8")
        with open(out / "md5.txt", "w", encoding="utf-8") as handle:
            handle.write("md5\tbytes\tpath\n")
            for path in sorted(out.rglob("*")):
                if path.is_file() and path.name != "md5.txt":
                    handle.write(f"{md5(path)}\t{path.stat().st_size}\t{path}\n")
        log(f"exit={exit_code} wall_s={manifest['wall_seconds']:.1f}")
        log.close()
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())

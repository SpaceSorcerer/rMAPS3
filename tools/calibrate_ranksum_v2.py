"""Calibration v2.1 of the rMAPS3 rank-sum supplement: the reportable p and q (2026-09-22).

1. Permutation unit = target-exon cluster (chr, strand, exonStart, exonEnd): every rMATS SE row
   is kept, labels move per target exon, and the observed foreground's cluster-size composition
   is held fixed. A row-unit result from tools/calibrate_ranksum.py (--permutation-unit row) can
   be carried beside it as *_rowunit sensitivity columns (--rowunit-root); it is never recomputed.
2. RBP-level statistics inside the permutation, per RBP, direction and pooled region:
   min-P over the RBP's unique motifs (PRIMARY: each motif's pooled regional-max z becomes its own
   tail count in its permutation distribution, the minimum is calibrated against its permutation
   distribution); max-z and mean-z over the same motifs as sensitivity columns.
3. Motif keys that share one k-mer are one motif-level test, mapped back to every carrier key.
4. Target exons present in both changed foregrounds are kept, as the tool does, and counted.
5. Two-stage scheme: --permutations (2,000) for every test, --refine-perms (100,000) for tests
   with stage-1 p <= --refine-threshold (0.005); stage streams are independent; seed 149.

Human / GRCh38 (hg38) / GENCODE v49 unless the caller's inputs say otherwise; SE events only.
"""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import math
import os
import platform
import random
import sys
import time
from pathlib import Path

import numpy as np
import scipy

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calibrate_ranksum as calib  # noqa: E402
import rmaps_calib_v2_lib as lib  # noqa: E402
import rmaps_countdist_io as io  # noqa: E402

POOLS = tuple(io.POOL_TO_REGIONS)
PERMUTATION_UNIT = "target-exon cluster (chr, strand, exonStart, exonEnd), size-matched"
ADEQUATE_LABEL = "adequate"
IDENTITY_FIELDS = {"motif"}


def power_label(arm, underpowered, n_inc, n_skip, t_inc, t_skip):
    """Power labelling is a NAMED-ARM decision passed by the caller, never a numeric rule."""
    if arm in set(underpowered):
        return ("underpowered (changed included/skipped = {}/{} events over {}/{} target exons) "
                "- landscape only, not for RBP ranking".format(n_inc, n_skip, t_inc, t_skip))
    return ADEQUATE_LABEL


def verify_duplicate_archives(counts_dir: Path, groups):
    """Carrier keys of one k-mer must hold bit-identical data (all fields but the key name)."""
    audit = []
    for kmer, keys in groups.items():
        if len(keys) < 2:
            continue
        with np.load(counts_dir / (keys[0] + ".counts.npz"), allow_pickle=False) as ref:
            reference = {name: np.asarray(ref[name]) for name in ref.files}
        for key in keys[1:]:
            with np.load(counts_dir / (key + ".counts.npz"), allow_pickle=False) as other:
                names = set(other.files)
                if names != set(reference):
                    raise ValueError("archive fields differ between " + keys[0] + " and " + key)
                for name in names - IDENTITY_FIELDS:
                    if not np.array_equal(np.asarray(other[name]), reference[name]):
                        raise ValueError("archive {} differs between {} and {}".format(
                            name, keys[0], key))
        audit.append({"kmer": kmer, "carrier_keys": ",".join(keys), "n_carriers": len(keys),
                      "archives_identical": True})
    return audit


def native_region_summary(model, roots, carrier_keys, direction, emit_positions, arm, alias):
    """Native per-region minimum, its window, the effect there; checked against the root table."""
    native_p = io.p_from_z(model.observed_z)
    argmin, effect, region_p = {}, {}, {}
    for region_id, region in enumerate(io.REGIONS):
        window_ids = np.flatnonzero(model.region_index == region_id)
        local = native_p[window_ids]
        best = int(window_ids[int(np.argmin(local))])
        argmin[region] = best
        effect[region] = calib.effect_at(model, best)
        region_p[region] = float(local.min())
        for key in carrier_keys:
            released = roots[direction][key][region]
            if released > 0 and region_p[region] > 0:
                if abs(math.log10(region_p[region]) - math.log10(released)) > 1e-6:
                    raise ValueError("native p differs from the released table: {} {} {}"
                                     .format(key, direction, region))
            if emit_positions is not None:
                rbp = alias.get(lib.table_name_of(key), lib.table_name_of(key))
                prefix = "\t".join([arm, key, rbp, direction, io.DIRECTION_LABEL[direction],
                                    region, io.REGION_TO_POOL[region]])
                for slot in window_ids:
                    emit_positions.write(prefix + "\t" + "\t".join([
                        str(int(model.position[slot])), repr(float(native_p[slot])),
                        str(int(model.fg_carrying[slot])), str(int(model.bg_carrying[slot])),
                        str(int(model.fg_total[slot])), str(int(model.bg_total[slot])),
                    ]) + "\n")
    for pool, members in io.POOL_TO_REGIONS.items():
        best_region = min(members, key=lambda r: region_p[r])
        argmin[pool] = argmin[best_region]
        effect[pool] = effect[best_region]
        region_p[pool] = region_p[best_region]
        argmin["native_region_" + pool] = best_region
    return argmin, effect, region_p


def load_rowunit(args, motifs, rbps):
    """The row-unit sensitivity tables, or empty maps when --rowunit-root was not given."""
    if not args.rowunit_root:
        return {}, {}, None
    base = Path(args.rowunit_root) / args.arm
    a_motif = {(r["motif_key"], r["direction"], r["region"]): r
               for r in lib.read_tsv(base / "per_motif_regions.tsv")}
    a_cond = {(r["RBP"], r["direction"], r["pooled_region"]): r
              for r in lib.read_tsv(base / "condensed_per_rbp.tsv")}
    if len(a_motif) != len(motifs) * 2 * len(io.REGIONS) or len(a_cond) != len(rbps) * 2 * len(POOLS):
        raise ValueError("row-unit tables in {} do not match the motif / RBP axes".format(base))
    units = {r.get("permutation_unit") for r in a_motif.values()}
    if units - {"row", None}:
        raise ValueError("--rowunit-root holds a non-row permutation unit: {}".format(sorted(units)))
    return a_motif, a_cond, base


def run_arm(args, emit):
    started_arm = time.perf_counter()
    arm = args.arm
    counts_dir = Path(args.counts_root) / arm
    out_dir = Path(args.out_root) / arm
    out_dir.mkdir(parents=True, exist_ok=True)
    alias = calib.load_alias(Path(args.alias_table))
    stage1_perms, stage2_perms = args.permutations, args.refine_perms
    threshold, chunk, seed = args.refine_threshold, args.chunk_size, args.seed

    ids, motifs = lib.exon_ids_for_counts_dir(counts_dir, verify_all_motifs=True)
    keys = {g: lib.target_keys(ids[g]) for g in io.GROUPS}
    n_rows = {g: int(ids[g].shape[0]) for g in io.GROUPS}
    n_target = {g: len(set(keys[g].tolist())) for g in io.GROUPS}
    both_way_target = len(set(keys["up"].tolist()) & set(keys["dn"].tolist()))
    both_way_rows = len(set(ids["up"].tolist()) & set(ids["dn"].tolist()))
    emit("{}: events (rMATS SE rows) / target exons: up {}/{} dn {}/{} bg {}/{}; "
         "target exons in both changed foregrounds {} (event rows in both {})".format(
             arm, n_rows["up"], n_target["up"], n_rows["dn"], n_target["dn"],
             n_rows["bg"], n_target["bg"], both_way_target, both_way_rows))

    groups = lib.unique_motif_groups(motifs)
    kmers = list(groups)
    dup_audit = verify_duplicate_archives(counts_dir, groups)
    emit("{} motif keys -> {} unique k-mers; {} shared k-mers verified bit-identical "
         "({} redundant keys)".format(len(motifs), len(kmers), len(dup_audit),
                                     len(motifs) - len(kmers)))
    rbp_kmers = lib.rbp_unique_kmers(motifs, alias)
    rbps = sorted(rbp_kmers)
    a_motif, a_cond, rowunit_dir = load_rowunit(args, motifs, rbps)

    drawers = {d: lib.ClusterDrawer(keys[d], keys["bg"]) for d in ("up", "dn")}
    for d in ("up", "dn"):
        emit("  {}: {} rows in {} clusters; pooled pool {} clusters; sizes {}".format(
            d, drawers[d].n1, drawers[d].n_clusters_fg, drawers[d].n_clusters_pooled,
            drawers[d].need))
    roots = {d: io.read_root_table(Path(args.released_root) / arm / ("pVal." + d + ".vs.bg.RNAmap.txt"))
             for d in ("up", "dn")}

    positions = open(out_dir / "positions_long.tsv", "w", newline="\n", encoding="utf-8")
    positions.write("\t".join(calib.POSITIONS_COLUMNS) + "\n")

    stage1, null1, meta = {}, {}, {}
    selection1 = {d: drawers[d].draw(lib.stage_rng(seed, 1, d), stage1_perms) for d in ("up", "dn")}
    started = time.perf_counter()
    for index, kmer in enumerate(kmers, 1):
        representative = groups[kmer][0]
        for d in ("up", "dn"):
            model = calib.MotifModel(counts_dir / (representative + ".counts.npz"), d)
            masks = model.window_masks()
            result, maxima = lib.calibrate_keep(model, masks, selection1[d], chunk)
            stage1[(kmer, d)] = result
            null1[(kmer, d)] = {pool: maxima[pool] for pool in POOLS}
            argmin, effect, region_p = native_region_summary(
                model, roots, groups[kmer], d, positions, arm, alias)
            meta[(kmer, d)] = {"n1": model.n1, "n0": model.n0, "argmin": argmin,
                               "effect": effect, "native": region_p}
            del model
        if index % 25 == 0 or index == len(kmers):
            emit("  stage 1 [{}]: {}/{} unique motifs, {:.1f}s".format(
                arm, index, len(kmers), time.perf_counter() - started))
    positions.close()
    stage1_seconds = time.perf_counter() - started

    rbp_obs, rbp_p1 = {}, {}
    for rbp in rbps:
        for d in ("up", "dn"):
            for pool in POOLS:
                observed = np.asarray([stage1[(k, d)][pool]["observed_max_z"] for k in rbp_kmers[rbp]])
                null = np.vstack([null1[(k, d)][pool] for k in rbp_kmers[rbp]])
                o_max, n_max, o_mean, n_mean, arg = lib.rbp_combine(observed, null)
                rbp_obs[(rbp, d, pool)] = {"max": o_max, "mean": o_mean,
                                           "argmax_kmer": rbp_kmers[rbp][arg] if arg is not None else None}
                p_minp, arg_minp = lib.rbp_minp(observed, null)
                rbp_p1[(rbp, d, pool)] = {"max": lib.permutation_p(o_max, n_max),
                                          "mean": lib.permutation_p(o_mean, n_mean),
                                          "minp": p_minp,
                                          "minp_kmer": rbp_kmers[rbp][arg_minp]
                                          if arg_minp is not None else None}
    del null1

    def promoted_motif(k, d):
        return any(math.isfinite(v["calibrated_p"]) and v["calibrated_p"] <= threshold
                   for v in stage1[(k, d)].values())

    motif_promoted = {(k, d) for k in kmers for d in ("up", "dn") if promoted_motif(k, d)}
    rbp_promoted = set()
    for rbp in rbps:
        for d in ("up", "dn"):
            own = any(math.isfinite(rbp_p1[(rbp, d, pool)][s]) and rbp_p1[(rbp, d, pool)][s] <= threshold
                      for pool in POOLS for s in ("max", "mean", "minp"))
            via_motif = any((k, d) in motif_promoted for k in rbp_kmers[rbp])
            if own or via_motif:
                rbp_promoted.add((rbp, d))
    rbp_needed = {(k, d) for rbp, d in rbp_promoted for k in rbp_kmers[rbp]}
    needed = sorted(motif_promoted | rbp_needed)
    emit("stage 2 [{}]: {} of {} unique motif-direction pairs promoted at motif level; "
         "{} of {} RBP-directions promoted; {} pairs to run".format(
             arm, len(motif_promoted), len(stage1), len(rbp_promoted), len(rbps) * 2, len(needed)))
    if args.stage1_only:
        emit("stage 2 SKIPPED by --stage1-only")
        motif_promoted, rbp_promoted, needed = set(), set(), []

    final = {k: {n: dict(v) for n, v in r.items()} for k, r in stage1.items()}
    motif_stage = {k: 1 for k in stage1}
    null2 = {}
    stage2_started = time.perf_counter()
    for d in ("up", "dn"):
        pairs = [k for k in needed if k[1] == d]
        if not pairs:
            continue
        draw_started = time.perf_counter()
        selection2 = drawers[d].draw(lib.stage_rng(seed, 2, d), stage2_perms)
        emit("  stage 2 [{} {}]: {} x {} cluster draws in {:.1f}s".format(
            arm, d, selection2.shape[0], selection2.shape[1], time.perf_counter() - draw_started))
        for done, (kmer, _) in enumerate(pairs, 1):
            model = calib.MotifModel(counts_dir / (groups[kmer][0] + ".counts.npz"), d)
            masks = model.window_masks()
            result, maxima = lib.calibrate_keep(model, masks, selection2, chunk)
            for name in masks:
                if math.isfinite(result[name]["observed_max_z"]) and abs(
                        result[name]["observed_max_z"] - stage1[(kmer, d)][name]["observed_max_z"]) > 1e-9:
                    raise ValueError("observed statistic changed during refinement")
            if (kmer, d) in motif_promoted:
                final[(kmer, d)] = result
                motif_stage[(kmer, d)] = 2
            if (kmer, d) in rbp_needed:
                null2[(kmer, d)] = {pool: maxima[pool] for pool in POOLS}
            del model
            if done % 10 == 0 or done == len(pairs):
                emit("  stage 2 [{} {}]: {}/{} pairs, {:.1f}s".format(
                    arm, d, done, len(pairs), time.perf_counter() - stage2_started))
        del selection2
    stage2_seconds = time.perf_counter() - stage2_started

    rbp_final = {}
    for rbp in rbps:
        for d in ("up", "dn"):
            for pool in POOLS:
                entry = {"p_max": rbp_p1[(rbp, d, pool)]["max"], "p_mean": rbp_p1[(rbp, d, pool)]["mean"],
                         "p_minp": rbp_p1[(rbp, d, pool)]["minp"],
                         "minp_kmer": rbp_p1[(rbp, d, pool)]["minp_kmer"],
                         "perms": stage1_perms, "stage": 1}
                if (rbp, d) in rbp_promoted:
                    observed = np.asarray([stage1[(k, d)][pool]["observed_max_z"] for k in rbp_kmers[rbp]])
                    null = np.vstack([null2[(k, d)][pool] for k in rbp_kmers[rbp]])
                    o_max, n_max, o_mean, n_mean, _ = lib.rbp_combine(observed, null)
                    p_minp, arg_minp = lib.rbp_minp(observed, null)
                    entry = {"p_max": lib.permutation_p(o_max, n_max),
                             "p_mean": lib.permutation_p(o_mean, n_mean), "p_minp": p_minp,
                             "minp_kmer": rbp_kmers[rbp][arg_minp] if arg_minp is not None else None,
                             "perms": stage2_perms, "stage": 2}
                rbp_final[(rbp, d, pool)] = entry
    del null2

    motif_family = [(k, d, p) for k in kmers for d in ("up", "dn") for p in io.PLOT_POOLS]
    motif_q_values, motif_divisor = calib.bh_adjust(
        [final[(k, d)][p]["calibrated_p"] for k, d, p in motif_family], return_divisor=True)
    motif_q = dict(zip(motif_family, (float(q) for q in motif_q_values)))
    rbp_family = [(r, d, p) for r in rbps for d in ("up", "dn") for p in io.PLOT_POOLS]
    rbp_q, rbp_divisor = {}, {}
    for s in ("max", "mean", "minp"):
        values, rbp_divisor[s] = calib.bh_adjust([rbp_final[k]["p_" + s] for k in rbp_family], return_divisor=True)
        rbp_q[s] = dict(zip(rbp_family, (float(q) for q in values)))
    emit("BH divisors (testable cells actually used): motif level {} of {} = {} unique motifs x 2 x 3 ({} "
         "untestable); RBP level min-P {} of {} = {} RBPs x 2 x 3 ({} untestable)".format(
             motif_divisor, len(motif_family), len(kmers), len(motif_family) - motif_divisor,
             rbp_divisor["minp"], len(rbp_family), len(rbps), len(rbp_family) - rbp_divisor["minp"]))

    counts = {"n_changed_included": n_rows["up"], "n_changed_skipped": n_rows["dn"],
              "n_changed_included_target_exons": n_target["up"],
              "n_changed_skipped_target_exons": n_target["dn"]}
    label = power_label(arm, args.underpowered_arms, n_rows["up"], n_rows["dn"],
                        n_target["up"], n_target["dn"])
    common = dict(counts, power_label=label, permutation_unit=PERMUTATION_UNIT)

    per_motif = []
    for key in motifs:
        kmer = lib.kmer_of(key)
        rbp = alias.get(lib.table_name_of(key), lib.table_name_of(key))
        carriers = groups[kmer]
        for d in ("up", "dn"):
            m = meta[(kmer, d)]
            for region in io.REGIONS:
                pool = io.REGION_TO_POOL[region]
                stat, pooled = final[(kmer, d)][region], final[(kmer, d)][pool]
                a_row = a_motif.get((key, d, region), {})
                row = {
                    "arm": arm, "motif_key": key, "RBP": rbp, "rbp_table_name": lib.table_name_of(key),
                    "direction": d, "direction_label": io.DIRECTION_LABEL[d], "region": region,
                    "pooled_region": pool, "plot": pool in io.PLOT_POOLS,
                    "native_ranksum_p": roots[d][key][region],
                    "native_argmin_position": int(m["argmin"][region]),
                    "calib_observed_min_p": stat["observed_min_p"], "calibrated_p": stat["calibrated_p"],
                    "calib_perms_used": stat["permutations"], "calib_stage": motif_stage[(kmer, d)],
                    "calib_reason": stat["reason"],
                    "native_ranksum_p_pooled": min(roots[d][key][r] for r in io.POOL_TO_REGIONS[pool]),
                    "calib_observed_min_p_pooled": pooled["observed_min_p"],
                    "calibrated_p_pooled": pooled["calibrated_p"],
                    "calib_perms_used_pooled": pooled["permutations"],
                    "calib_stage_pooled": motif_stage[(kmer, d)], "calib_pooled_reason": pooled["reason"],
                    "n_fg_exons": m["n1"], "n_bg_exons": m["n0"],
                    "n_fg_target_exons": n_target[d], "n_bg_target_exons": n_target["bg"],
                    "calibrated_q": motif_q.get((kmer, d, pool), math.nan),
                    "untestable": not math.isfinite(pooled["calibrated_p"]),
                    "unique_motif_id": kmer, "unique_motif_representative_key": carriers[0],
                    "n_keys_sharing_motif": len(carriers),
                    "rbps_sharing_motif": ",".join(sorted({alias.get(lib.table_name_of(c), lib.table_name_of(c))
                                                           for c in carriers})),
                    "calibrated_p_rowunit": lib.as_float(a_row.get("calibrated_p")),
                    "calibrated_p_pooled_rowunit": lib.as_float(a_row.get("calibrated_p_pooled")),
                    "calibrated_q_rowunit": lib.as_float(a_row.get("calibrated_q")),
                }
                row.update(m["effect"][region])
                row.update(common)
                per_motif.append(row)
    native_q = calib.bh_adjust([r["native_ranksum_p"] for r in per_motif])
    for row, q in zip(per_motif, native_q):
        row["native_q"] = float(q)
    by_key = {(r["motif_key"], r["direction"], r["region"]): r for r in per_motif}

    condensed, rbp_rows = [], []
    for rbp in rbps:
        rbp_keys = [k for k in motifs if alias.get(lib.table_name_of(k), lib.table_name_of(k)) == rbp]
        table_names = sorted({lib.table_name_of(k) for k in rbp_keys})
        for d in ("up", "dn"):
            for pool in POOLS:
                candidates = [(final[(lib.kmer_of(k), d)][pool]["calibrated_p"],
                               meta[(lib.kmer_of(k), d)]["native"][pool], k) for k in rbp_keys]
                finite = [c for c in candidates if math.isfinite(c[0])]
                motif = min(finite)[2] if finite else rbp_keys[0]
                kmer = lib.kmer_of(motif)
                region = meta[(kmer, d)]["argmin"]["native_region_" + pool]
                effect_row = by_key[(motif, d, region)]
                a_row = a_cond.get((rbp, d, pool), {})
                obs = rbp_obs[(rbp, d, pool)]
                fin = rbp_final[(rbp, d, pool)]
                argmax_key = (next(k for k in rbp_keys if lib.kmer_of(k) == obs["argmax_kmer"])
                              if obs["argmax_kmer"] else "NA")
                record = {
                    "arm": arm, "RBP": rbp, "rbp_table_names": ",".join(table_names),
                    "direction": d, "direction_label": io.DIRECTION_LABEL[d],
                    "pooled_region": pool, "plot": pool in io.PLOT_POOLS,
                    "selected_motif_key": motif, "selected_native_region": region,
                    "native_ranksum_p": effect_row["native_ranksum_p_pooled"],
                    "native_q_at_selected_region": effect_row["native_q"],
                    "calibrated_p": final[(kmer, d)][pool]["calibrated_p"],
                    "calibrated_q": motif_q.get((kmer, d, pool), math.nan),
                    "untestable": not math.isfinite(final[(kmer, d)][pool]["calibrated_p"]),
                    "calib_perms_used": final[(kmer, d)][pool]["permutations"],
                    "calib_stage": motif_stage[(kmer, d)],
                    "native_argmin_position": effect_row["native_argmin_position"],
                    "n_motifs_total": len(rbp_keys), "n_unique_motifs": len(rbp_kmers[rbp]),
                    "n_motifs_calib_q_lt_0.05": sum(
                        1 for k in rbp_kmers[rbp]
                        if math.isfinite(motif_q.get((k, d, pool), math.nan)) and motif_q[(k, d, pool)] < 0.05),
                    "n_motifs_native_q_lt_0.05": sum(
                        1 for k in rbp_keys
                        if any(math.isfinite(by_key[(k, d, r)]["native_q"]) and by_key[(k, d, r)]["native_q"] < 0.05
                               for r in io.POOL_TO_REGIONS[pool])),
                    "rbp_observed_max_z": obs["max"], "rbp_max_z_motif_key": argmax_key,
                    "rbp_calibrated_p_maxz": fin["p_max"],
                    "rbp_calibrated_q_maxz": rbp_q["max"].get((rbp, d, pool), math.nan),
                    "rbp_calibrated_p_minp": fin["p_minp"],
                    "rbp_calibrated_q_minp": rbp_q["minp"].get((rbp, d, pool), math.nan),
                    "rbp_untestable": not math.isfinite(fin["p_minp"]),
                    "rbp_minp_motif_key": (next(k for k in rbp_keys if lib.kmer_of(k) == fin["minp_kmer"])
                                           if fin["minp_kmer"] else "NA"),
                    "rbp_observed_mean_z": obs["mean"], "rbp_calibrated_p_meanz": fin["p_mean"],
                    "rbp_calibrated_q_meanz": rbp_q["mean"].get((rbp, d, pool), math.nan),
                    "rbp_calib_perms_used": fin["perms"], "rbp_calib_stage": fin["stage"],
                    "selected_motif_key_rowunit": a_row.get("selected_motif_key", "NA"),
                    "calibrated_p_rowunit": lib.as_float(a_row.get("calibrated_p")),
                    "calibrated_q_rowunit": lib.as_float(a_row.get("calibrated_q")),
                    "n_fg_target_exons": n_target[d], "n_bg_target_exons": n_target["bg"],
                }
                for column in ("n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
                               "fg_proportion", "bg_proportion", "enrichment_ratio",
                               "fg_mean_count", "bg_mean_count", "count_ratio"):
                    record[column] = effect_row[column]
                record.update(common)
                condensed.append(record)

    for d in ("up", "dn"):
        for pool in io.PLOT_POOLS:
            panel = [r for r in condensed if r["direction"] == d and r["pooled_region"] == pool]
            for r in panel:
                r["_neg_max"] = -r["rbp_observed_max_z"] if math.isfinite(r["rbp_observed_max_z"]) else math.inf
                r["_neg_mean"] = -r["rbp_observed_mean_z"] if math.isfinite(r["rbp_observed_mean_z"]) else math.inf
                r["_rowunit_native"] = lib.as_float(a_cond.get((r["RBP"], d, pool), {}).get("native_ranksum_p"))
            ranks = {
                "rank_motif_level": lib.panel_rank(panel, "calibrated_q", "calibrated_p", "native_ranksum_p"),
                "rank_rbp_maxz": lib.panel_rank(panel, "rbp_calibrated_q_maxz", "rbp_calibrated_p_maxz", "_neg_max"),
                "rank_rbp_meanz": lib.panel_rank(panel, "rbp_calibrated_q_meanz", "rbp_calibrated_p_meanz",
                                                 "_neg_mean"),
                "rank_rbp_minp": lib.panel_rank(panel, "rbp_calibrated_q_minp", "rbp_calibrated_p_minp", "_neg_max"),
            }
            if a_cond:
                ranks["rank_motif_level_rowunit"] = lib.panel_rank(
                    panel, "calibrated_q_rowunit", "calibrated_p_rowunit", "_rowunit_native")
            for r in panel:
                for name, table in ranks.items():
                    r[name] = table[r["RBP"]]
                for tmp in ("_neg_max", "_neg_mean", "_rowunit_native"):
                    r.pop(tmp)

    for r in condensed:
        rbp_rows.append({c: r.get(c) for c in RBP_LEVEL_COLUMNS if c in r} | {
            "unique_motifs": ",".join(rbp_kmers[r["RBP"]])})

    report = {
        "arm": arm, "seed": seed, "permutation_unit": PERMUTATION_UNIT,
        "stage1_permutations": stage1_perms, "stage2_permutations": stage2_perms,
        "refine_threshold": threshold, "stage1_only": bool(args.stage1_only),
        "n_motif_keys": len(motifs), "n_unique_motifs": len(kmers),
        "n_redundant_keys": len(motifs) - len(kmers),
        "redundant_motif_tests_removed": (len(motifs) - len(kmers)) * 2 * len(io.PLOT_POOLS),
        "motif_family_size": len(motif_family), "rbp_family_size": len(rbp_family), "n_rbps": len(rbps),
        "motif_bh_divisor": motif_divisor,
        "rbp_bh_divisor": {"minp": rbp_divisor["minp"], "maxz": rbp_divisor["max"], "meanz": rbp_divisor["mean"]},
        "motif_untestable_cells": len(motif_family) - motif_divisor,
        "rbp_untestable_cells": {"minp": len(rbp_family) - rbp_divisor["minp"],
                                 "maxz": len(rbp_family) - rbp_divisor["max"],
                                 "meanz": len(rbp_family) - rbp_divisor["mean"]},
        "unique_pairs_total": len(stage1), "unique_pairs_promoted_motif_level": len(motif_promoted),
        "rbp_directions_promoted": len(rbp_promoted), "stage2_pairs_run": len(needed),
        "stage1_wall_seconds": stage1_seconds, "stage2_wall_seconds": stage2_seconds,
        "cluster_sizes": {d: drawers[d].need for d in drawers},
        "events": n_rows, "target_exons": n_target,
        "target_exons_in_both_changed_foregrounds": both_way_target,
        "event_rows_in_both_changed_foregrounds": both_way_rows,
        "stage2_p_floor": 1.0 / (1.0 + stage2_perms),
        "singleton_bh_q_floor_motif": min(1.0, motif_divisor / (1.0 + stage2_perms)),
        "singleton_bh_q_floor_rbp": min(1.0, rbp_divisor["minp"] / (1.0 + stage2_perms)),
        "max_calibrated_p_among_q_lt_0.05": {
            "motif": max([final[k[:2]][k[2]]["calibrated_p"] for k in motif_family if motif_q[k] < 0.05],
                         default=None),
            "rbp_maxz": max([rbp_final[k]["p_max"] for k in rbp_family if rbp_q["max"][k] < 0.05], default=None),
            "rbp_meanz": max([rbp_final[k]["p_mean"] for k in rbp_family if rbp_q["mean"][k] < 0.05],
                             default=None),
            "rbp_minp": max([rbp_final[k]["p_minp"] for k in rbp_family if rbp_q["minp"][k] < 0.05],
                            default=None)},
        "rowunit_source": str(rowunit_dir) if rowunit_dir else None,
        "duplicate_kmers": dup_audit,
        "arm_wall_seconds_before_writing": time.perf_counter() - started_arm,
    }
    return out_dir, per_motif, condensed, rbp_rows, report, motifs


PER_MOTIF_COLUMNS = [
    "arm", "power_label", "n_changed_included", "n_changed_skipped",
    "n_changed_included_target_exons", "n_changed_skipped_target_exons", "permutation_unit",
    "motif_key", "RBP", "rbp_table_name", "direction", "direction_label", "region", "pooled_region", "plot",
    "native_ranksum_p", "native_argmin_position",
    "calib_observed_min_p", "calibrated_p", "calib_perms_used", "calib_stage", "calib_reason",
    "native_ranksum_p_pooled", "calib_observed_min_p_pooled", "calibrated_p_pooled",
    "calib_perms_used_pooled", "calib_stage_pooled", "calib_pooled_reason",
    "n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
    "fg_proportion", "bg_proportion", "enrichment_ratio", "fg_mean_count", "bg_mean_count", "count_ratio",
    "native_q", "calibrated_q", "untestable", "n_fg_target_exons", "n_bg_target_exons", "unique_motif_id",
    "unique_motif_representative_key", "n_keys_sharing_motif", "rbps_sharing_motif",
    "calibrated_p_rowunit", "calibrated_p_pooled_rowunit", "calibrated_q_rowunit",
]
CONDENSED_COLUMNS = [
    "arm", "power_label", "n_changed_included", "n_changed_skipped",
    "n_changed_included_target_exons", "n_changed_skipped_target_exons", "permutation_unit",
    "RBP", "rbp_table_names", "direction", "direction_label", "pooled_region", "plot",
    "selected_motif_key", "selected_native_region", "native_ranksum_p", "native_q_at_selected_region",
    "calibrated_p", "calibrated_q", "untestable", "calib_perms_used", "calib_stage", "native_argmin_position",
    "n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
    "fg_proportion", "bg_proportion", "enrichment_ratio", "fg_mean_count", "bg_mean_count", "count_ratio",
    "n_motifs_total", "n_motifs_calib_q_lt_0.05", "n_motifs_native_q_lt_0.05",
    "n_unique_motifs", "n_fg_target_exons", "n_bg_target_exons",
    "rbp_calibrated_p_minp", "rbp_calibrated_q_minp", "rbp_untestable", "rbp_minp_motif_key",
    "rbp_observed_max_z", "rbp_max_z_motif_key", "rbp_calibrated_p_maxz", "rbp_calibrated_q_maxz",
    "rbp_observed_mean_z", "rbp_calibrated_p_meanz", "rbp_calibrated_q_meanz",
    "rbp_calib_perms_used", "rbp_calib_stage",
    "selected_motif_key_rowunit", "calibrated_p_rowunit", "calibrated_q_rowunit",
    "rank_motif_level", "rank_rbp_minp", "rank_rbp_maxz", "rank_rbp_meanz", "rank_motif_level_rowunit",
]
RBP_LEVEL_COLUMNS = [
    "arm", "power_label", "n_changed_included", "n_changed_skipped",
    "n_changed_included_target_exons", "n_changed_skipped_target_exons",
    "RBP", "rbp_table_names", "direction", "direction_label", "pooled_region", "plot",
    "n_unique_motifs", "unique_motifs",
    "rbp_calibrated_p_minp", "rbp_calibrated_q_minp", "rbp_untestable", "rbp_minp_motif_key", "rank_rbp_minp",
    "rbp_observed_max_z", "rbp_max_z_motif_key", "rbp_calibrated_p_maxz", "rbp_calibrated_q_maxz", "rank_rbp_maxz",
    "rbp_observed_mean_z", "rbp_calibrated_p_meanz", "rbp_calibrated_q_meanz", "rank_rbp_meanz",
    "rbp_calib_perms_used", "rbp_calib_stage",
    "selected_motif_key", "calibrated_p", "calibrated_q", "rank_motif_level",
    "selected_motif_key_rowunit", "calibrated_p_rowunit", "calibrated_q_rowunit",
    "rank_motif_level_rowunit", "permutation_unit",
]

NULL_DESCRIPTION = (
    "Reportable p and q. The statistic is the released engine's own one-sided rank-sum on per-event "
    "motif hit counts, reduced to the smallest p over the windows of a region. Its null distribution "
    "comes from permuting the changed/background labels over TARGET EXONS (chr, strand, exonStart, "
    "exonEnd): rMATS rows that share a target exon move together and are never deduplicated, and every "
    "draw keeps the observed foreground's cluster-size composition. The minimum over windows is taken "
    "inside every permutation (Westfall-Young min-P), so the p is corrected for the window scan. The "
    "null asks whether changed target exons are an exchangeable subset of the tested universe given "
    "their motif counts; exon covariates that track motif content, such as length, enter it.")


def validity_sentence(report):
    """The one inferential statement the README, LESSONS.md, the workbook, the readout and the legend share."""
    return ("Every calibrated p is a valid permutation p: an unpromoted test keeps its stage-1 p (B = {:,}, "
            "resolution 1/{:,}) and a promoted test takes its stage-2 p (B = {:,}, resolution 1/{:,}); "
            "calib_perms_used gives each p's B. BH over valid p-values controls the FDR under PRDS-type "
            "dependence, and mixed resolutions do not break that.".format(
                report["stage1_permutations"], report["stage1_permutations"] + 1,
                report["stage2_permutations"], report["stage2_permutations"] + 1))


def readme_rows(report, alias_table, rowunit_root):
    shared = "; ".join("{} ({})".format(a["kmer"], a["carrier_keys"]) for a in report["duplicate_kmers"])
    return [
        ("Layer", "Released rMAPS3 engine, --stat-method mannwhitney: one-sided Mann-Whitney U on per-exon "
                  "motif hit counts, changed > background, tie-corrected normal approximation with continuity "
                  "correction. The released statistic and its inputs are unchanged; only its p is calibrated."),
        ("Permutation null", NULL_DESCRIPTION),
        ("n_changed_included / n_changed_skipped", "Events (rMATS SE rows) in the changed INCLUDED / SKIPPED "
                                                  "set, as the released engine counted them."),
        ("n_changed_*_target_exons", "Distinct target exons behind those events: the independent units the "
                                     "calibration resolves. Counts read 'n events (rMATS SE rows) over N "
                                     "target exons'."),
        ("n_fg_exons / n_bg_exons", "Events (rMATS SE rows) in the foreground / background of this direction "
                                    "(v1 column names kept); n_fg_target_exons / n_bg_target_exons are the "
                                    "distinct target exons."),
        ("power_label", "'adequate' or 'underpowered (...) - landscape only, not for RBP ranking'; set per "
                        "named arm by --underpowered-arms, never by a numeric threshold."),
        ("Both-way target exons", "Target exons present in both changed sets are kept in both, as the "
                                  "released tool does; refinement_report.json counts them "
                                  "(target_exons_in_both_changed_foregrounds)."),
        ("native_ranksum_p", "Released engine value: smallest p over the windows of the sub-region "
                             "(per_motif) or pooled region (condensed). Normal approximation; order only."),
        ("calibrated_p", "Motif level: permutation p of the released regional minimum under the target-exon "
                         "cluster permutation."),
        ("calibrated_q", "Motif level: Benjamini-Hochberg with divisor {} = the testable cells of unique motifs "
                         "x 2 directions x 3 plotted pooled regions (nominal family {} = {} x 2 x 3; {} untestable "
                         "cells carry untestable = TRUE and no q). Keys carrying one k-mer are one test, copied to "
                         "every carrier.".format(report["motif_bh_divisor"], report["motif_family_size"],
                                                 report["n_unique_motifs"],
                                                 report["motif_family_size"] - report["motif_bh_divisor"])),
        ("untestable / rbp_untestable", "TRUE when the calibrated p is NA because no window of the pooled region "
                                        "has usable variance (for example no motif hit in any event). Such a cell "
                                        "is not counted in the BH divisor."),
        ("unique_motif_id", "The k-mer of the motif key. {} keys carry {} unique k-mers; shared k-mers, "
                            "archives verified bit-identical: {}.".format(
                                report["n_motif_keys"], report["n_unique_motifs"], shared or "none")),
        ("selected_motif_key", "Condensed sheet: the RBP's motif with the smallest motif-level calibrated p "
                               "(ties by native p, then key); its q is motif-level. The RBP-level columns "
                               "correct for choosing the best of the RBP's motifs."),
        ("rbp_calibrated_p_minp / _q_minp", "PRIMARY RBP-level statistic: Westfall-Young min-P over the RBP's "
                                            "unique motifs. In the observed labels and every permutation each "
                                            "motif's pooled regional-max z becomes its own tail count in that "
                                            "motif's permutation distribution (same draws); the minimum over "
                                            "motifs is calibrated against its permutation distribution. Motifs "
                                            "weigh equally whatever the spread of their null. With one motif it "
                                            "equals the motif p. BH divisor {} = the testable cells of RBPs x 2 "
                                            "x 3 (nominal {}).".format(report["rbp_bh_divisor"]["minp"],
                                                                       report["rbp_family_size"])),
        ("rbp_*_maxz / rbp_*_meanz", "SENSITIVITY: max and mean over the RBP's motifs of the same per-motif z, "
                                     "each calibrated against its own permutation distribution; BH divisor "
                                     "{} (max-z) and {} (mean-z) of nominal {}. Max-z penalises an RBP whose "
                                     "second motif has a heavier-tailed null; min-P does not.".format(
                                         report["rbp_bh_divisor"]["maxz"], report["rbp_bh_divisor"]["meanz"],
                                         report["rbp_family_size"])),
        ("RBP grouping", "Alias table {}; table names without an alias entry stay their own RBP.".format(
            alias_table)),
        ("*_rowunit", ("SENSITIVITY: the row-unit result (each rMATS row an independent permutation unit), read "
                       "from {} and never recomputed. Duplicate target exons make it an invalid permutation "
                       "unit; do not report it.".format(rowunit_root)) if rowunit_root else
                      "Not computed for this run (no --rowunit-root); the columns are NA."),
        ("rank_*", "Rank among RBPs within a plotted panel by q, then p, then the observed statistic, then "
                   "RBP name."),
        ("Stages", "Stage 1 = {} permutations for every test; stage 2 = {} for a motif-direction pair with any "
                   "stage-1 p <= {}, and for an RBP-direction when any RBP-level statistic did or any of its "
                   "motifs was promoted. Seed {}; independent stage streams.".format(
                       report["stage1_permutations"], report["stage2_permutations"],
                       report["refine_threshold"], report["seed"])),
        ("Validity", validity_sentence(report)),
        ("calib_perms_used / rbp_calib_perms_used", "The B behind each calibrated p: {:,} (stage 1, resolution "
                                                    "1/{:,}) or {:,} (stage 2, resolution 1/{:,}). calib_stage "
                                                    "says which.".format(
                                                        report["stage1_permutations"],
                                                        report["stage1_permutations"] + 1,
                                                        report["stage2_permutations"],
                                                        report["stage2_permutations"] + 1)),
        ("Floors", "Stage-2 floor 1/({} + 1). A p at the floor means no permutation reached the observed "
                   "statistic, not that the true p equals the floor.".format(report["stage2_permutations"])),
        ("Families", "No correction across arms; each arm is its own family at each level."),
        ("Effect columns", "At the native argmin window of the sub-region with the smallest released p inside "
                           "the pooled region, for the selected motif; count_ratio = fg_mean_count / "
                           "bg_mean_count = the tool's motif-score ratio."),
        ("positions", "Released per-window native p and carrying / hit counts per motif key, direction and "
                      "window."),
        ("Species / assembly", "Human / GRCh38 (hg38) / GENCODE v49 unless the run's inputs say otherwise; "
                               "SE events only."),
    ]


def write_workbook_fast(path: Path, sheets):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook(write_only=True)
    bold = Font(bold=True)
    for name, (columns, rows) in sheets.items():
        sheet = book.create_sheet(name[:31])
        sheet.freeze_panes = "A2"
        header = []
        for c in columns:
            cell = WriteOnlyCell(sheet, value=c)
            cell.font = bold
            header.append(cell)
        sheet.append(header)
        n = 0
        for row in rows:
            sheet.append([calib.excel_value(row.get(c)) for c in columns])
            n += 1
        if n:
            sheet.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(columns)), n + 1)
    book.save(path)


def iter_positions(path: Path):
    with open(path, "r", encoding="utf-8") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        for line in handle:
            row = dict(zip(header, line.rstrip("\n").split("\t")))
            for c in ("position", "n_fg_carrying", "n_bg_carrying", "fg_total_hits", "bg_total_hits"):
                row[c] = int(row[c])
            row["native_ranksum_p"] = float(row["native_ranksum_p"])
            yield row


def fmt(value):
    return "none" if value is None else "{:.3g}".format(value)


def readout_lines(arm, condensed, report):
    first = condensed[0]
    lines = [
        "# {} - calibrated rank-sum supplement v2.1 (target-exon unit, RBP-level min-P): the reportable p "
        "and q".format(arm), "",
        "Changed sets: **{} INCLUDED events (rMATS SE rows) over {} target exons; {} SKIPPED events over {} "
        "target exons**. Background {} events over {} target exons. Target exons in both changed sets, kept "
        "in both as the tool does: {}. Power label: **{}**.".format(
            first["n_changed_included"], first["n_changed_included_target_exons"],
            first["n_changed_skipped"], first["n_changed_skipped_target_exons"],
            report["events"]["bg"], report["target_exons"]["bg"],
            report["target_exons_in_both_changed_foregrounds"], first["power_label"]), "",
    ]
    if first["power_label"] != ADEQUATE_LABEL:
        lines += ["> Underpowered arm: the calibration is complete, but with this few changed target exons the "
                  "calibrated p takes few distinct values. Read as landscape; do not rank RBPs from it.", ""]
    lines += [
        NULL_DESCRIPTION, "",
        "- **Motif level**: permutation p per unique motif and pooled region; BH divisor {} testable cells of "
        "the {} = {} unique motifs x 2 directions x 3 pooled regions ({} untestable).".format(
            report["motif_bh_divisor"], report["motif_family_size"], report["n_unique_motifs"],
            report["motif_untestable_cells"]),
        "- **RBP level, min-P (primary)**: smallest per-motif permutation p per RBP, the minimum taken inside "
        "every permutation; BH divisor {} of {} ({} untestable).".format(
            report["rbp_bh_divisor"]["minp"], report["rbp_family_size"], report["rbp_untestable_cells"]["minp"]),
        "- **RBP level, max-z and mean-z (sensitivity)**: BH divisor {} and {} of {}.".format(
            report["rbp_bh_divisor"]["maxz"], report["rbp_bh_divisor"]["meanz"], report["rbp_family_size"]),
        "- **Row unit (sensitivity, not reportable)**: {}.".format(
            "*_rowunit columns from " + report["rowunit_source"] if report["rowunit_source"] else "not computed"),
        "",
        "Stage 1: {} permutations, seed {}. Stage 2: {} permutations for {} unique motif-direction pairs and {} "
        "RBP-directions ({} pairs run). {} The largest calibrated p among q<0.05 calls is {} (motif), {} (RBP "
        "min-P), {} (RBP max-z), {} (RBP mean-z).".format(
            report["stage1_permutations"], report["seed"], report["stage2_permutations"],
            report["unique_pairs_promoted_motif_level"], report["rbp_directions_promoted"],
            report["stage2_pairs_run"], validity_sentence(report),
            fmt(report["max_calibrated_p_among_q_lt_0.05"]["motif"]),
            fmt(report["max_calibrated_p_among_q_lt_0.05"]["rbp_minp"]),
            fmt(report["max_calibrated_p_among_q_lt_0.05"]["rbp_maxz"]),
            fmt(report["max_calibrated_p_among_q_lt_0.05"]["rbp_meanz"])),
        "", "Top three RBPs per plotted panel (q):", "",
        "| panel | RBP level, min-P | RBP level, max-z | RBP level, mean-z | motif level (motif) |",
        "|---|---|---|---|---|",
    ]
    for pool in io.PLOT_POOLS:
        for d in ("up", "dn"):
            panel = [r for r in condensed if r["pooled_region"] == pool and r["direction"] == d]
            cells = []
            for rank_key, q_key, with_motif in (("rank_rbp_minp", "rbp_calibrated_q_minp", False),
                                                ("rank_rbp_maxz", "rbp_calibrated_q_maxz", False),
                                                ("rank_rbp_meanz", "rbp_calibrated_q_meanz", False),
                                                ("rank_motif_level", "calibrated_q", True)):
                top = sorted(panel, key=lambda r: r[rank_key])[:3]
                cells.append(", ".join("{}{} {:.3g}".format(
                    r["RBP"], " (" + r["selected_motif_key"] + ")" if with_motif else "", r[q_key]) for r in top))
            lines.append("| {}, {} | {} |".format(pool, io.DIRECTION_LABEL[d], " | ".join(cells)))
    lines += ["", "RBP x panel calls at q<0.05 over the six plotted panels: motif level {}, RBP min-P {}, RBP "
              "max-z {}, RBP mean-z {}, row unit {}.".format(
                  *(sum(1 for r in condensed if r["plot"] and math.isfinite(lib.as_float(r.get(k)))
                        and lib.as_float(r.get(k)) < 0.05)
                    for k in ("calibrated_q", "rbp_calibrated_q_minp", "rbp_calibrated_q_maxz",
                              "rbp_calibrated_q_meanz", "calibrated_q_rowunit"))),
              "", "The main figure remains the released engine output (order only)."]
    return lines


def write_arm(args, out_dir, per_motif, condensed, rbp_rows, report, motifs, started):
    calib.write_tsv(out_dir / "per_motif_regions.tsv", per_motif, PER_MOTIF_COLUMNS)
    calib.write_tsv(out_dir / "condensed_per_rbp.tsv", condensed, CONDENSED_COLUMNS)
    calib.write_tsv(out_dir / "rbp_level.tsv", rbp_rows, RBP_LEVEL_COLUMNS)
    (out_dir / "readout.md").write_text("\n".join(readout_lines(args.arm, condensed, report)) + "\n",
                                        encoding="utf-8")
    write_workbook_fast(out_dir / (args.arm + "_calibrated_ranksum_v2.xlsx"), {
        "README": (["field", "description"],
                   [{"field": k, "description": v}
                    for k, v in readme_rows(report, args.alias_table, report["rowunit_source"])]),
        "condensed": (CONDENSED_COLUMNS, condensed),
        "per_motif": (PER_MOTIF_COLUMNS, per_motif),
        "rbp_level": (RBP_LEVEL_COLUMNS, rbp_rows),
        "positions": (calib.POSITIONS_COLUMNS, iter_positions(out_dir / "positions_long.tsv")),
    })
    counts_dir = Path(args.counts_root) / args.arm
    released = Path(args.released_root) / args.arm
    required = [Path(args.alias_table), released / "pVal.up.vs.bg.RNAmap.txt",
                released / "pVal.dn.vs.bg.RNAmap.txt"] + [counts_dir / (m + ".counts.npz") for m in motifs]
    optional = [counts_dir / "conversion_manifest.tsv", counts_dir / "VERIFY.md"]
    if report["rowunit_source"]:
        required += [Path(report["rowunit_source"]) / n for n in ("per_motif_regions.tsv", "condensed_per_rbp.tsv")]
    with open(out_dir / "input_md5s.tsv", "w", newline="\n", encoding="utf-8") as handle:
        handle.write("path\tmd5\n")
        for path in required + [p for p in optional if p.exists()]:
            if not path.exists():
                raise FileNotFoundError("input missing: " + str(path))
            handle.write("{}\t{}\n".format(path, io.md5(path)))
    with open(out_dir / "versions.txt", "w", encoding="utf-8") as handle:
        handle.write("python\t{}\n".format(sys.version.split()[0]))
        handle.write("numpy\t{}\n".format(np.__version__))
        handle.write("scipy\t{}\n".format(scipy.__version__))
        handle.write("openpyxl\t{}\n".format(importlib.metadata.version("openpyxl")))
        handle.write("platform\t{}\n".format(platform.platform()))
        for module in (Path(__file__).resolve(), Path(lib.__file__).resolve(),
                       Path(calib.__file__).resolve(), Path(io.__file__).resolve()):
            handle.write("{}\t{}\n".format(module, io.md5(module)))
        handle.write("seed\t{}\nstage1_permutations\t{}\nstage2_permutations\t{}\nrefine_threshold\t{}\n".format(
            args.seed, args.permutations, args.refine_perms, args.refine_threshold))
        handle.write("generated\t{}\n".format(time.strftime("%Y-%m-%dT%H:%M:%S")))
        for var in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
            handle.write("{}\t{}\n".format(var, os.environ.get(var, "unset")))
    report["arm_wall_seconds_total"] = time.perf_counter() - started
    (out_dir / "refinement_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")


def build_parser():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--arm", required=True)
    p.add_argument("--counts-root", required=True, help="root holding <arm>/<motif>.counts.npz")
    p.add_argument("--released-root", required=True,
                   help="root holding <arm>/pVal.{up,dn}.vs.bg.RNAmap.txt from the released engine")
    p.add_argument("--out-root", required=True, help="writes <out-root>/<arm>/")
    p.add_argument("--alias-table", required=True, help="table_name<TAB>hgnc_symbol RBP grouping")
    p.add_argument("--rowunit-root", default=None,
                   help="optional: <root>/<arm>/ from calibrate_ranksum.py --permutation-unit row, carried as "
                        "*_rowunit sensitivity columns")
    p.add_argument("--underpowered-arms", nargs="*", default=[])
    p.add_argument("--permutations", type=int, default=2000)
    p.add_argument("--refine-perms", type=int, default=100000)
    p.add_argument("--refine-threshold", type=float, default=0.005)
    p.add_argument("--chunk-size", type=int, default=500)
    p.add_argument("--seed", type=int, default=149)
    p.add_argument("--stage1-only", action="store_true")
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    random.seed(args.seed)
    np.random.seed(args.seed)
    out_dir = Path(args.out_root) / args.arm
    out_dir.mkdir(parents=True, exist_ok=True)
    log = open(out_dir / "command.log", "a", encoding="utf-8")

    def emit(message):
        print(message, flush=True)
        log.write(message + "\n")
        log.flush()

    started = time.perf_counter()
    emit("start={} cmd={}".format(time.strftime("%Y-%m-%dT%H:%M:%S"),
                                  " ".join([sys.executable] + (sys.argv if argv is None else [__file__] + argv))))
    emit("threads OMP={} OPENBLAS={} MKL={}".format(*(os.environ.get(v, "unset") for v in (
        "OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"))))
    out_dir, per_motif, condensed, rbp_rows, report, motifs = run_arm(args, emit)
    write_arm(args, out_dir, per_motif, condensed, rbp_rows, report, motifs, started)
    emit("wrote per_motif_regions.tsv ({}), condensed_per_rbp.tsv ({}), rbp_level.tsv ({}), xlsx, readout.md"
         .format(len(per_motif), len(condensed), len(rbp_rows)))
    emit("wall_seconds={:.1f} stage1={:.1f} stage2={:.1f} exit=0".format(
        time.perf_counter() - started, report["stage1_wall_seconds"], report["stage2_wall_seconds"]))
    log.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

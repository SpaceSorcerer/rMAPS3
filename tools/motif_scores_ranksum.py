"""SUPERSEDED - use tools/calibrate_ranksum_v2.py (fg_mean_count, bg_mean_count, count_ratio columns).

Motif scores at the regional-minimum window of the released rMAPS3 rank-sum layer, added on
2026-09-21 to the v1 ROW-UNIT summaries of tools/calibrate_ranksum.py. Calibration v2 carries
the same quantity natively (count_ratio = motif_score_ratio wherever both are defined) and the
v4 figures take dot size from it; this tool is kept to reproduce the 2026-09-21 v1 workbooks,
their exact score verification table and their links to the released per-motif map images.

Scope decision 2026-09-21: carry only quantities the released tool itself outputs. The motif
score of a group at a window is the countDist row sum for that window divided by the number of
exons in the group. Nothing is added: no U, no z, no AUC, no fold enrichment, no carrying counts.

This score is the curve the released per-motif map draws, verified 2026-09-21 against the
released engine legacy/motifMapSE_MP.py at commit b9a9dce (tag upstream-base-2026-09-16):
ccc() at lines 1020-1036 sums the per-exon counts at each window, plotMotifs() at lines
1088-1105 divides that sum by uNum, dNum or bNum, the group exon counts set at lines 1463-1476,
and initMotif() at lines 486-494 reads the one-motif temp file written at lines 1235-1238, so
the len(motifs) factor in the denominator is 1 for every per-motif map. The map also applies a
min(1.0, ...) clamp that no score in this panel approaches.

The audited engine on this branch (audits F6, R4 and F15) rewrote the countDist writer to six
columns and changed the plotted value to a proportion of eligible events, so it does not
describe a released-engine run; read archives made from released runs only.

The released engine is neither run nor modified. Inputs are the verified count archives, the
v1 summary tables and the released run directory, all read only; every product is a new file
or a v2 copy. All roots are required arguments.
"""

from __future__ import annotations

import argparse
import hashlib
import math
import os
import platform
import random
import re
import sys
import time
from pathlib import Path

import numpy as np
import scipy

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rmaps_countdist_io as io  # noqa: E402

SEED = 149

SCORE_COLUMNS = [
    "n_changed", "n_background",
    "motif_score_changed", "motif_score_background", "motif_score_ratio",
    "motif_score_flag",
]

MOTIF_SCORE_COLUMNS = [
    "arm", "power_label", "n_changed_included", "n_changed_skipped",
    "motif_key", "RBP", "rbp_table_name", "direction", "direction_label",
    "region_level", "region", "pooled_region", "source_sub_region", "plot",
    "argmin_window", "argmin_window_position", "native_ranksum_p",
] + SCORE_COLUMNS + ["authors_map_png"]

VERIFY_COLUMNS = [
    "arm", "motif_key", "direction", "region", "argmin_window",
    "sum_hits_changed", "n_changed", "motif_score_changed",
    "sum_hits_background", "n_background", "motif_score_background",
    "score_equals_sum_over_n",
]

NEW_SUB_COLUMNS = [
    "argmin_window", "argmin_window_position",
    "motif_score_changed", "motif_score_background", "motif_score_ratio",
    "motif_score_flag",
]
NEW_POOLED_COLUMNS = ["pooled_source_sub_region"] + [c + "_pooled" for c in NEW_SUB_COLUMNS]
RETIRED_HEADING = "\n## Effect sizes\n"


def md5(path: Path) -> str:
    digest = hashlib.md5()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tsv(path: Path):
    with open(path, "r", encoding="utf-8") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        return [dict(zip(header, line.rstrip("\n").split("\t"))) for line in handle], header


def format_cell(value):
    if value is None:
        return "NA"
    if isinstance(value, float):
        return "NA" if not math.isfinite(value) else repr(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def write_tsv(path: Path, rows, columns):
    with open(path, "w", newline="\n", encoding="utf-8") as handle:
        handle.write("\t".join(columns) + "\n")
        for row in rows:
            handle.write("\t".join(format_cell(row.get(c)) for c in columns) + "\n")


def motif_scores(fg, bg):
    """Per-window motif score of each group and their ratio.

    The score is the countDist row sum over the exons of the group divided by the group
    size. fg and bg are exons x windows count matrices, dense or sparse.
    """
    sum_fg = np.asarray(fg.sum(axis=0), dtype=np.float64).ravel()
    sum_bg = np.asarray(bg.sum(axis=0), dtype=np.float64).ravel()
    n1, n0 = int(fg.shape[0]), int(bg.shape[0])
    score_fg = sum_fg / n1
    score_bg = sum_bg / n0
    with np.errstate(invalid="ignore", divide="ignore"):
        ratio = np.where(sum_bg > 0, score_fg / np.where(score_bg > 0, score_bg, np.nan), np.nan)
    return {
        "sum_hits_changed": sum_fg, "sum_hits_background": sum_bg,
        "n_changed": np.full(sum_fg.shape, n1, dtype=np.int64),
        "n_background": np.full(sum_fg.shape, n0, dtype=np.int64),
        "motif_score_changed": score_fg, "motif_score_background": score_bg,
        "motif_score_ratio": ratio,
    }


def score_flag(scores, window) -> str:
    return "background_zero_hits" if float(scores["sum_hits_background"][window]) == 0.0 else ""


def scores_at(scores, window):
    out = {}
    for name in SCORE_COLUMNS:
        if name == "motif_score_flag":
            out[name] = score_flag(scores, window)
            continue
        value = np.asarray(scores[name])[window]
        out[name] = int(value) if np.issubdtype(value.dtype, np.integer) else float(value)
    return out


def map_image_index(maps_dir: Path):
    """Exact on-disk per-motif map images of the released run, keyed by normalised name."""
    index = {}
    for path in sorted(maps_dir.glob("SE.*.png")):
        index[path.name[len("SE."):-len(".png")]] = path
    return index


def map_key(motif_key: str) -> str:
    """The released engine names a map after tableName-motif with bracket runs collapsed."""
    table_name, _, motif = motif_key.partition(".")
    return re.sub(r"[\[\]]+", "_", table_name + "-" + motif)


def resolve_map_images(maps_dir: Path, motifs):
    index = map_image_index(maps_dir)
    resolved, missing = {}, []
    for motif in motifs:
        path = index.get(map_key(motif))
        if path is None:
            missing.append(motif)
        else:
            resolved[motif] = str(path)
    if missing:
        raise SystemExit("no released map image for {} motifs, first {}".format(
            len(missing), missing[:3]))
    return resolved


def run_arm(arm: str, counts_root: Path, summary_root: Path, released_root: Path, emit):
    out_dir = summary_root / arm
    per_motif, per_motif_header = read_tsv(out_dir / "per_motif_regions.tsv")
    condensed, condensed_header = read_tsv(out_dir / "condensed_per_rbp.tsv")
    by_key = {(r["motif_key"], r["direction"], r["region"]): r for r in per_motif}
    motifs = sorted({r["motif_key"] for r in per_motif})
    images = resolve_map_images(released_root / arm / "maps", motifs)

    score_rows, verify_rows = [], []
    computed = {}
    for motif in motifs:
        with np.load(counts_root / arm / (motif + ".counts.npz"), allow_pickle=False) as data:
            region_index = np.asarray(data["region_of_position"])
            position = np.asarray(data["position"])
            bg = io.group_csr(data, "bg")
            groups = {d: io.group_csr(data, d) for d in ("up", "dn")}
        for direction in ("up", "dn"):
            scores = motif_scores(groups[direction], bg)
            template = by_key[(motif, direction, io.REGIONS[0])]
            # The 2026-09-20 lab v1 schema carries these three; tools/calibrate_ranksum.py in this
            # fork does not (power labels moved to v2), so fall back to the archive group sizes.
            base = {
                "arm": arm, "power_label": template.get("power_label", "NA"),
                "n_changed_included": int(template.get("n_changed_included",
                                                       groups["up"].shape[0])),
                "n_changed_skipped": int(template.get("n_changed_skipped",
                                                      groups["dn"].shape[0])),
                "motif_key": motif, "RBP": template["RBP"],
                "rbp_table_name": template["rbp_table_name"],
                "direction": direction, "direction_label": io.DIRECTION_LABEL[direction],
                "authors_map_png": images[motif],
            }
            per_region = {}
            for region in io.REGIONS:
                source = by_key[(motif, direction, region)]
                window = int(source["native_argmin_position"])
                if int(region_index[window]) != io.REGIONS.index(region):
                    raise SystemExit("argmin window {} is not inside {} for {} {}".format(
                        window, region, motif, direction))
                sum_fg = float(scores["sum_hits_changed"][window])
                sum_bg = float(scores["sum_hits_background"][window])
                n1 = int(scores["n_changed"][window])
                n0 = int(scores["n_background"][window])
                verify_rows.append({
                    "arm": arm, "motif_key": motif, "direction": direction, "region": region,
                    "argmin_window": window,
                    "sum_hits_changed": int(sum_fg), "n_changed": n1,
                    "motif_score_changed": float(scores["motif_score_changed"][window]),
                    "sum_hits_background": int(sum_bg), "n_background": n0,
                    "motif_score_background": float(scores["motif_score_background"][window]),
                    "score_equals_sum_over_n": (
                        float(scores["motif_score_changed"][window]) == sum_fg / n1
                        and float(scores["motif_score_background"][window]) == sum_bg / n0),
                })
                record = dict(base)
                record.update({
                    "region_level": "sub_region", "region": region,
                    "pooled_region": io.REGION_TO_POOL[region],
                    "source_sub_region": region,
                    "plot": io.REGION_TO_POOL[region] in io.PLOT_POOLS,
                    "argmin_window": window,
                    "argmin_window_position": int(position[window]),
                    "native_ranksum_p": float(source["native_ranksum_p"]),
                })
                record.update(scores_at(scores, window))
                score_rows.append(record)
                per_region[region] = record
            for pool, members in io.POOL_TO_REGIONS.items():
                best = min(members, key=lambda r: float(
                    by_key[(motif, direction, r)]["native_ranksum_p"]))
                record = dict(per_region[best])
                record.update({
                    "region_level": "pooled", "region": pool, "pooled_region": pool,
                    "source_sub_region": best, "plot": pool in io.PLOT_POOLS,
                })
                score_rows.append(record)
                per_region[pool] = record
            computed[(motif, direction)] = per_region
    emit("  {}: {} motifs x 2 directions, {} score rows, {} verification rows, "
         "{} released map images matched".format(
             arm, len(motifs), len(score_rows), len(verify_rows), len(images)))

    per_motif_v2 = []
    for row in per_motif:
        key = (row["motif_key"], row["direction"])
        sub = computed[key][row["region"]]
        pool = computed[key][row["pooled_region"]]
        new = dict(row)
        for name in NEW_SUB_COLUMNS:
            new[name] = sub[name]
            new[name + "_pooled"] = pool[name]
        new["pooled_source_sub_region"] = pool["source_sub_region"]
        new["authors_map_png"] = sub["authors_map_png"]
        per_motif_v2.append(new)

    condensed_v2 = []
    for row in condensed:
        key = (row["selected_motif_key"], row["direction"])
        sub = computed[key][row["selected_native_region"]]
        new = dict(row)
        for name in NEW_SUB_COLUMNS:
            new[name] = sub[name]
        new["authors_map_png"] = sub["authors_map_png"]
        condensed_v2.append(new)

    v2_per_motif_columns = (per_motif_header + NEW_SUB_COLUMNS + NEW_POOLED_COLUMNS
                            + ["authors_map_png"])
    v2_condensed_columns = condensed_header + NEW_SUB_COLUMNS + ["authors_map_png"]
    write_tsv(out_dir / "motif_scores_per_motif.tsv", score_rows, MOTIF_SCORE_COLUMNS)
    write_tsv(out_dir / "motif_scores_verification.tsv", verify_rows, VERIFY_COLUMNS)
    write_tsv(out_dir / "per_motif_regions_v2.tsv", per_motif_v2, v2_per_motif_columns)
    write_tsv(out_dir / "condensed_per_rbp_v2.tsv", condensed_v2, v2_condensed_columns)
    return (score_rows, verify_rows, per_motif_v2, condensed_v2,
            v2_per_motif_columns, v2_condensed_columns)


README_ROWS = [
    ("Scope", "Human / GRCh38 (hg38) / GENCODE v49; SE (skipped exon) events only. Motif scores "
              "for the released rMAPS3 rank-sum layer and its permutation calibration."),
    ("What is new in v2", "Every column of the v1 tables is carried over unchanged. v2 adds the "
                          "argmin window identifiers, the motif scores of the two groups and "
                          "their ratio, and the path of the released per-motif map image. No p or "
                          "q value was recomputed, altered or rewritten, and no statistic beyond "
                          "what the released tool itself outputs was added."),
    ("Where the scores are measured", "At the single window that produced the reported regional "
                                      "minimum p, the same window the released engine reports for "
                                      "that sub-region. Pooled-region values are those of the "
                                      "member sub-region carrying the regional minimum, named in "
                                      "source_sub_region."),
    ("motif_score_changed", "The countDist row sum of the changed exons at that window divided by "
                            "n_changed: motif hits per exon. Identical to the v1 column "
                            "fg_mean_count."),
    ("Relation to the map image", "motif_score = mean per-exon hit count per window = the curve "
                                  "drawn by the released tool's per-motif map (verified in "
                                  "released motifMapSE_MP.py, commit b9a9dce: ccc at lines "
                                  "1020-1036 sums the per-exon counts and plotMotifs at lines "
                                  "1088-1105 divides by the group exon count, with one motif per "
                                  "map; the map clamps at 1.0, which no score in this panel "
                                  "approaches)."),
    ("motif_score_background", "The same quantity for the background exon set. Identical to the "
                               "v1 column bg_mean_count."),
    ("motif_score_ratio", "motif_score_changed / motif_score_background. NA when the background "
                          "row sum is zero, flagged as background_zero_hits in motif_score_flag. "
                          "Identical to the v1 column count_ratio wherever both are defined."),
    ("motif_score_flag", "Empty when the ratio is defined; background_zero_hits when the "
                         "background carries no motif hit at that window, which makes the ratio "
                         "undefined by definition rather than by failure."),
    ("n_changed / n_background", "Exon counts of the two groups; the same values as the v1 "
                                 "columns n_fg_exons / n_bg_exons."),
    ("argmin_window", "Global window slot, 0-1199, over the eight sub-regions in order. Identical "
                      "to the v1 column native_argmin_position, which despite its name holds the "
                      "global slot, not the within-region position."),
    ("argmin_window_position", "Position of that window WITHIN its sub-region, the coordinate the "
                               "RNA map x-axis uses. Restarts at 0 in every sub-region."),
    ("authors_map_png", "Full path of the per-motif map image the released run itself wrote for "
                        "this motif and arm, verified to exist on disk. One image per motif "
                        "covering both directions, so the up and dn rows of a motif carry the "
                        "same path."),
    ("Columns ending _pooled", "The same value at the pooled region of that row, taken from the "
                               "member sub-region named in pooled_source_sub_region."),
    ("Species / assembly", "Human / GRCh38 (hg38) / GENCODE v49; SE events only."),
    ("Caveat", "A score measured at the argmin window is measured at a selected extreme. It "
               "describes that window and is not an estimate of the score across the sub-region; "
               "the per-position curve is in positions_long.tsv and in the released map image."),
    ("Provenance", "Script tools/motif_scores_ranksum.py of the rMAPS3 lab fork, seed 149, inputs "
                   "the released_counts archives, the v1 summary tables and the released run "
                   "directory; versions_motif_scores.txt records interpreter and library "
                   "versions. The released engine was neither run nor modified."),
]


def write_workbook(path: Path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    for name, (columns, rows) in sheets.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([excel_value(row.get(c)) for c in columns])
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def excel_value(value):
    if value is None:
        return "NA"
    if isinstance(value, float) and not math.isfinite(value):
        return "NA"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        if value in ("NA", "TRUE", "FALSE", "") or "\\" in value:
            return value
        try:
            return int(value)
        except ValueError:
            pass
        try:
            return float(value)
        except ValueError:
            return value
    return value


def retire_section(readout: Path, heading: str, marker: str, archive: Path = None) -> str:
    """Remove a section this dispatch appended earlier, restoring the prefix byte-exactly."""
    text = readout.read_text(encoding="utf-8")
    if heading not in text:
        return "absent"
    if text.count(heading) != 1:
        raise SystemExit("more than one {} section in {}".format(heading.strip(), readout))
    index = text.rindex(heading)
    removed = text[index:]
    if marker not in removed:
        raise SystemExit("unexpected content under {} in {}".format(heading.strip(), readout))
    if archive is not None:
        archive.write_text(removed, encoding="utf-8")
    readout.write_text(text[:index], encoding="utf-8")
    return "retired {} characters{}".format(
        len(removed), "" if archive is None else " to " + archive.name)


def readout_paragraph(arm, score_rows, verify_rows):
    flagged = sum(1 for r in score_rows if r["motif_score_flag"])
    exact = sum(1 for r in verify_rows if r["score_equals_sum_over_n"])
    return [
        "",
        "## Motif scores",
        "",
        "Every p in this supplement is a tail probability: it says how surprising a signal is and "
        "nothing about how large it is. The companion tables `motif_scores_per_motif.tsv`, "
        "`per_motif_regions_v2.tsv` and `condensed_per_rbp_v2.tsv` add the size in the only "
        "currency the released tool itself reports, the motif score: the countDist row sum of a "
        "group at a window divided by the number of exons in that group, which is motif hits per "
        "exon. They carry "
        "`motif_score_changed`, `motif_score_background` and their ratio at the same "
        "regional-minimum window the reported p comes from, the exon counts of both groups, the "
        "argmin window, and the path of the map image the released run wrote for that motif. No "
        "statistic outside the released tool was added.",
        "",
        "For this arm all {} of {} verification rows reproduce the score as the row sum divided by "
        "the exon count exactly, and {} of {} score rows are flagged where the background carries "
        "no motif hit at that window and the ratio is undefined by definition. One caveat travels "
        "with them: a score read at the argmin window is read at a selected extreme, so it "
        "describes that window rather than the whole sub-region. The per-position curve is in "
        "`positions_long.tsv` and in the released map image, whose path each row carries.".format(
            exact, len(verify_rows), flagged, len(score_rows)),
        "",
        "motif_score = mean per-exon hit count per window = the curve drawn by the released "
        "tool's per-motif map (verified in released motifMapSE_MP.py, commit b9a9dce: ccc at "
        "lines 1020-1036 sums the per-exon counts and plotMotifs at lines 1088-1105 divides by "
        "the group exon count, with one motif per map; the map clamps at 1.0, which no score in "
        "this panel approaches).",
    ]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    parser.add_argument("--arms", nargs="+", required=True)
    parser.add_argument("--counts-root", required=True,
                        help="<ARM>/<motif>.counts.npz from tools/countdist_to_npz.py")
    parser.add_argument("--summary-root", required=True,
                        help="v1 row-unit summary root; <ARM>/ receives the v2 tables")
    parser.add_argument("--released-root", required=True,
                        help="released run root holding <ARM>/maps/SE.*.png")
    parser.add_argument("--report-root", required=True,
                        help="cross-arm summary and md5 guard of the protected v1 files")
    args = parser.parse_args(argv)

    random.seed(SEED)
    np.random.seed(SEED)

    summary_root = Path(args.summary_root)
    counts_root = Path(args.counts_root)
    released_root = Path(args.released_root)
    protected = ("per_motif_regions.tsv", "condensed_per_rbp.tsv", "positions_long.tsv",
                 "refinement_report.json", "refinement_tasks.tsv", "input_md5s.tsv",
                 "versions.txt")
    before = {}
    for arm in args.arms:
        for name in protected + (arm + "_calibrated_ranksum.xlsx",):
            path = summary_root / arm / name
            if path.exists():
                before[str(path)] = md5(path)

    started = time.perf_counter()
    summary_lines = []
    for arm in args.arms:
        arm_started = time.perf_counter()
        out_dir = summary_root / arm
        log = open(out_dir / "command.log", "a", encoding="utf-8")

        def emit(message, log=log):
            print(message, flush=True)
            log.write(message + "\n")
            log.flush()

        emit("start={} cmd={}".format(time.strftime("%Y-%m-%dT%H:%M:%S"),
                                      " ".join([sys.executable] + sys.argv)))
        (score_rows, verify_rows, per_motif_v2, condensed_v2,
         v2_per_motif_columns, v2_condensed_columns) = run_arm(
            arm, counts_root, summary_root, released_root, emit)

        not_exact = [r for r in verify_rows if not r["score_equals_sum_over_n"]]
        if not_exact:
            raise SystemExit("motif score is not the row sum over the exon count in {} rows of {}"
                             .format(len(not_exact), arm))
        flagged = sum(1 for r in score_rows if r["motif_score_flag"])
        na_ratio = sum(1 for r in score_rows if not math.isfinite(r["motif_score_ratio"]))
        if flagged != na_ratio:
            raise SystemExit("flag and NA ratio disagree in " + arm)

        write_workbook(out_dir / (arm + "_calibrated_ranksum_v2.xlsx"), {
            "README": (["field", "description"],
                       [{"field": k, "description": v} for k, v in README_ROWS]),
            "condensed_per_rbp": (v2_condensed_columns, condensed_v2),
            "per_motif_regions": (v2_per_motif_columns, per_motif_v2),
            "motif_scores": (MOTIF_SCORE_COLUMNS, score_rows),
        })
        readout = out_dir / "readout.md"
        emit("  {}: readout.md effect-sizes section {}".format(arm, retire_section(
            readout, RETIRED_HEADING, "auc_cles",
            out_dir / "retired_readout_section_effect_sizes.md")))
        emit("  {}: readout.md motif-scores section {}".format(arm, retire_section(
            readout, "\n## Motif scores\n", "motif_score_changed")))
        with open(readout, "a", encoding="utf-8") as handle:
            handle.write("\n".join(readout_paragraph(arm, score_rows, verify_rows)) + "\n")
        (out_dir / "versions_motif_scores.txt").write_text("\n".join([
            "python\t" + platform.python_version(),
            "numpy\t" + np.__version__,
            "scipy\t" + scipy.__version__,
            "platform\t" + platform.platform(),
            "script\t" + str(Path(__file__).resolve()),
            "script_md5\t" + md5(Path(__file__).resolve()),
            "module_md5\t" + md5(Path(io.__file__).resolve()),
            "seed\t" + str(SEED),
            "generated\t" + time.strftime("%Y-%m-%dT%H:%M:%S"),
            "OMP_NUM_THREADS\t" + os.environ.get("OMP_NUM_THREADS", "unset"),
            "OPENBLAS_NUM_THREADS\t" + os.environ.get("OPENBLAS_NUM_THREADS", "unset"),
            "MKL_NUM_THREADS\t" + os.environ.get("MKL_NUM_THREADS", "unset"),
        ]) + "\n", encoding="utf-8")
        elapsed = time.perf_counter() - arm_started
        emit("  {}: {} score rows, {} exact score checks, {} flagged NA ratios, {:.1f}s".format(
            arm, len(score_rows), len(verify_rows) - len(not_exact), flagged, elapsed))
        emit("exit=0")
        log.close()
        summary_lines.append({
            "arm": arm, "score_rows": len(score_rows), "verify_rows": len(verify_rows),
            "per_motif_v2_rows": len(per_motif_v2), "condensed_v2_rows": len(condensed_v2),
            "score_checks_exact": len(verify_rows) - len(not_exact),
            "na_ratio_rows_flagged": flagged,
            "map_images_matched": len({r["authors_map_png"] for r in score_rows}),
            "wall_seconds": elapsed,
        })

    changed = [path for path, digest in before.items() if md5(Path(path)) != digest]
    if changed:
        raise SystemExit("protected v1 files changed: " + ", ".join(changed))
    report_root = Path(args.report_root)
    report_root.mkdir(parents=True, exist_ok=True)
    write_tsv(report_root / "motif_scores_summary.tsv", summary_lines, list(summary_lines[0]))
    with open(report_root / "motif_scores_md5_guard.tsv", "w", newline="\n",
              encoding="utf-8") as handle:
        handle.write("path\tmd5_before\tmd5_after\n")
        for path, digest in sorted(before.items()):
            handle.write("{}\t{}\t{}\n".format(path, digest, md5(Path(path))))
    print("total wall {:.1f}s, {} arms, {} protected v1 files unchanged".format(
        time.perf_counter() - started, len(args.arms), len(before)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

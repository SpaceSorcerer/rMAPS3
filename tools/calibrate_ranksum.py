"""Westfall-Young min-P calibration of the released rMAPS3 rank-sum statistic: v1, the ROW-UNIT
version, kept for SENSITIVITY ONLY (superseded 2026-09-22 by tools/calibrate_ranksum_v2.py).

Duplicate target exons (rMATS rows sharing chr/strand/exonStart/exonEnd with different flanking
exons) make the rMATS row an invalid permutation unit, so the reportable p and q come from
calibrate_ranksum_v2.py (target-exon cluster unit, RBP-level min-P, unique-k-mer family).
This module is still imported by v2 for MotifModel and its helpers; run it with
--permutation-unit row to produce the *_rowunit sensitivity columns v2 carries.

The main figure layer is the authors' released engine run with
--stat-method mannwhitney: a one-sided Mann-Whitney U on per-exon motif hit
counts, changed > background, evaluated by the tie-corrected normal
approximation. That approximation is badly anti-conservative in this sparse,
extremely unbalanced regime, so this supplement calibrates THAT SAME statistic
by permuting the changed/background labels. No statistic is redefined and no
released table is rewritten.

Key vectorization: the pooled multiset of counts at a window is fixed under any
permutation of the labels, so the midrank of a count value, the null mean and
the tie-corrected variance are all permutation-invariant. A permutation only
changes which rank values are summed, so the whole calibration is a sparse
selection over one precomputed per-exon rank-increment matrix.

Because p = Phi_bar(z) is a strictly decreasing transform applied identically to
every window, the minimum p over the windows of a region equals the maximum z
over those windows. The comparison is therefore done in z so that extreme
windows do not collapse to an underflowed zero.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import random
import sys
import time
from pathlib import Path

import numpy as np
import scipy
from scipy import sparse
from scipy.stats import false_discovery_control

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rmaps_countdist_io as io  # noqa: E402

STATISTICS = list(io.REGIONS) + list(io.POOL_TO_REGIONS)


def load_alias(path: Path):
    alias = {}
    with open(path, "r", encoding="utf-8") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        if header[:2] != ["table_name", "hgnc_symbol"]:
            raise ValueError("unexpected alias header in " + str(path))
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            alias[fields[0]] = fields[1]
    return alias


PERMUTATION_UNITS = ("row", "exon-dedupe", "cluster")
UNIT_NOTES = {
    "row": "Labels permuted over rMATS event ROWS. A target exon carried by several rows can be "
           "split between the changed and background draw, which the data can never do, so the "
           "null is narrower than the truth and the p is anti-conservative.",
    "cluster": "Labels permuted over TARGET-EXON CLUSTERS at the observed foreground's exact "
               "cluster-size composition, so every row of a target exon moves together and the "
               "drawn changed set has exactly the observed number of rows.",
    "exon-dedupe": "Duplicate target-exon rows collapsed to their first occurrence in both groups "
                   "before ranking, then rows permuted. This changes the observed statistic, so "
                   "the released root tables are NOT reproduced and are not cross-checked.",
}


def target_exon_key(exon_ids: np.ndarray):
    """The 4-field target-exon identity of each engine row: chr:strand:exonStart:exonEnd.

    The engine's fasta headers are chr:strand:exonStart:exonEnd:firstES:firstEE:secondES:secondEE,
    so rows sharing a target exon with different flanking exons collapse to one key here.
    Returns (keys, parsed): an identifier that is not in the engine's eight-field form is kept
    whole and reported as unparsed, so a non-engine archive can still be calibrated row-wise but
    can never be silently clustered.
    """
    out = np.empty(exon_ids.size, dtype=object)
    parsed = True
    for i, value in enumerate(exon_ids):
        fields = str(value).split(":")
        if len(fields) != 8:
            parsed = False
            out[i] = str(value)
        else:
            out[i] = ":".join(fields[:4])
    return out, parsed


def first_occurrence_mask(keys: np.ndarray) -> np.ndarray:
    seen, mask = set(), np.zeros(keys.size, dtype=bool)
    for i, key in enumerate(keys):
        if key not in seen:
            seen.add(key)
            mask[i] = True
    return mask


class MotifModel:
    """Permutation-invariant rank scaffolding for one motif and direction."""

    def __init__(self, npz_path: Path, direction: str, unit: str = "row"):
        if unit not in PERMUTATION_UNITS:
            raise ValueError("unknown permutation unit: " + unit)
        self.unit = unit
        with np.load(npz_path, allow_pickle=False) as data:
            fg = io.group_csr(data, direction)
            bg = io.group_csr(data, "bg")
            fg_keys, fg_parsed = target_exon_key(np.asarray(data[direction + "_exon_id"]))
            bg_keys, bg_parsed = target_exon_key(np.asarray(data["bg_exon_id"]))
            self.region_index = np.asarray(data["region_of_position"])
            self.position = np.asarray(data["position"])
        self.exon_ids_parsed = bool(fg_parsed and bg_parsed)
        if unit != "row" and not self.exon_ids_parsed:
            raise ValueError(
                "permutation unit {!r} needs the engine's eight-field exon identifiers; this archive "
                "does not carry them".format(unit))
        self.n_rows_input = (int(fg.shape[0]), int(bg.shape[0]))
        if unit == "exon-dedupe":
            fg_keep, bg_keep = first_occurrence_mask(fg_keys), first_occurrence_mask(bg_keys)
            fg, bg = fg[fg_keep], bg[bg_keep]
            fg_keys, bg_keys = fg_keys[fg_keep], bg_keys[bg_keep]
        self.n_distinct_exons = (len(set(fg_keys)), len(set(bg_keys)))
        shared = set(fg_keys) & set(bg_keys)
        if shared and unit != "row":
            raise ValueError(
                "a target exon appears in both the changed and the background set ({} of them, e.g. {}); "
                "cluster permutation is undefined".format(len(shared), sorted(shared)[0]))
        self.cluster_codes = np.unique(np.concatenate([fg_keys, bg_keys]), return_inverse=True)[1]
        self.n1 = int(fg.shape[0])
        self.n0 = int(bg.shape[0])
        self.n_total = self.n1 + self.n0
        self.n_windows = int(fg.shape[1])
        pooled = sparse.vstack([fg, bg], format="csr")
        kmax = int(pooled.data.max(initial=0))
        tied = io.csr_histograms(pooled, kmax).astype(np.float64)
        cumulative_below = np.cumsum(tied, axis=0) - tied
        midrank = cumulative_below + (tied + 1.0) / 2.0
        self.mu = self.n1 * self.n0 / 2.0
        tie_term = (tied ** 3 - tied).sum(axis=0)
        with np.errstate(invalid="ignore", divide="ignore"):
            variance = self.n1 * self.n0 / 12.0 * (
                (self.n_total + 1.0) - tie_term / (self.n_total * (self.n_total - 1.0)))
        self.sigma = np.sqrt(np.where(variance > 0, variance, np.nan))
        self.valid = np.isfinite(self.sigma)
        self.baseline = self.n1 * midrank[0]
        increment = midrank[pooled.data.astype(np.int64),
                            pooled.indices.astype(np.int64)] - midrank[0, pooled.indices]
        self.delta = sparse.csr_matrix(
            (increment.astype(np.float64), pooled.indices, pooled.indptr), shape=pooled.shape)
        self.observed_z = self.z(self.baseline + np.asarray(
            self.delta[: self.n1].sum(axis=0)).ravel())
        self.fg_carrying = np.asarray((fg > 0).sum(axis=0)).ravel()
        self.bg_carrying = np.asarray((bg > 0).sum(axis=0)).ravel()
        self.fg_total = np.asarray(fg.sum(axis=0)).ravel()
        self.bg_total = np.asarray(bg.sum(axis=0)).ravel()

    def z(self, rank_sum):
        u1 = rank_sum - self.n1 * (self.n1 + 1.0) / 2.0
        with np.errstate(invalid="ignore", divide="ignore"):
            return (u1 - self.mu - 0.5) / self.sigma

    def window_masks(self):
        masks = {}
        for region_id, region in enumerate(io.REGIONS):
            masks[region] = (self.region_index == region_id) & self.valid
        for pool, members in io.POOL_TO_REGIONS.items():
            masks[pool] = np.logical_or.reduce([masks[m] for m in members])
        return masks

    def permutation_z(self, selection: np.ndarray) -> np.ndarray:
        """z at every window for a batch of label assignments (rows of indices)."""
        batch, size = selection.shape
        if size != self.n1:
            raise ValueError("assignment cardinality must equal the changed-set size")
        indptr = np.arange(0, batch * size + 1, size, dtype=np.int64)
        picker = sparse.csr_matrix(
            (np.ones(batch * size, dtype=np.float64), selection.ravel().astype(np.int32), indptr),
            shape=(batch, self.n_total))
        rank_sum = np.asarray((picker @ self.delta).todense()) + self.baseline[None, :]
        return self.z(rank_sum)


def statistic_maxima(z: np.ndarray, masks: dict) -> dict:
    out = {}
    for name, mask in masks.items():
        if not mask.any():
            out[name] = np.full(z.shape[0], np.nan) if z.ndim == 2 else math.nan
            continue
        block = z[:, mask] if z.ndim == 2 else z[mask]
        out[name] = np.nanmax(block, axis=-1)
    return out


def draw_selection(rng, n_total: int, n1: int, batch: int) -> np.ndarray:
    """Uniform fixed-cardinality subsets; identical in law to shuffling labels."""
    out = np.empty((batch, n1), dtype=np.int32)
    for i in range(batch):
        out[i] = rng.choice(n_total, n1, replace=False)
    return out


class ClusterSampler:
    """Size-matched target-exon cluster draws: the changed set is resampled as whole exons.

    A target exon is changed or background in every one of its rMATS rows, so free row permutation
    splits a cluster that the data can never split and makes the null too narrow. Drawing the
    observed foreground's exact cluster-size composition from the pooled cluster pool keeps the
    drawn foreground at exactly n1 rows and conditions on that composition.
    """

    def __init__(self, codes: np.ndarray, n1: int):
        self.n1 = int(n1)
        order = np.argsort(codes, kind="stable")
        boundaries = np.flatnonzero(np.diff(codes[order])) + 1
        members = np.split(order, boundaries)
        sizes = np.array([m.size for m in members], dtype=np.int64)
        observed = np.unique(codes[:n1])
        observed_rows = int(sizes[observed].sum())
        if observed_rows != self.n1:
            raise ValueError(
                "the observed changed set splits a target-exon cluster: its {} clusters hold {} "
                "rows, not {}".format(observed.size, observed_rows, self.n1))
        observed_sizes = sizes[observed]
        self.pool, self.need = {}, {}
        for size in np.unique(sizes):
            self.pool[int(size)] = np.flatnonzero(sizes == size)
            self.need[int(size)] = int(np.sum(observed_sizes == size))
        for size, count in self.need.items():
            if count and self.pool[size].size < count:
                raise ValueError(
                    "only {} pooled clusters of size {}; the observed foreground needs {}".format(
                        self.pool[size].size, size, count))
        self.members = members
        self.n_clusters_observed = int(observed.size)
        self.n_clusters_pool = int(sizes.size)
        self.signature = hashlib.md5(np.ascontiguousarray(codes).tobytes()).hexdigest()

    def draw(self, rng, batch: int) -> np.ndarray:
        out = np.empty((batch, self.n1), dtype=np.int32)
        for i in range(batch):
            chosen = [self.members[c] for size, count in self.need.items() if count
                      for c in rng.choice(self.pool[size], count, replace=False)]
            rows = np.concatenate(chosen)
            if rows.size != self.n1:
                raise ValueError("cluster draw produced {} rows, expected {}".format(rows.size, self.n1))
            out[i] = rows
        return out


def calibrate(model: MotifModel, masks: dict, selection: np.ndarray, chunk: int):
    observed = statistic_maxima(model.observed_z, masks)
    exceed = {name: 0 for name in masks}
    usable = {name: 0 for name in masks}
    for start in range(0, selection.shape[0], chunk):
        batch = selection[start: start + chunk]
        z = model.permutation_z(batch)
        maxima = statistic_maxima(z, masks)
        for name in masks:
            values = maxima[name]
            finite = np.isfinite(values)
            usable[name] += int(finite.sum())
            if math.isfinite(observed[name]):
                exceed[name] += int(np.sum(values[finite] >= observed[name]))
    total = selection.shape[0]
    return {name: {
        "observed_max_z": float(observed[name]),
        "observed_min_p": float(io.p_from_z(np.asarray(observed[name]))) if math.isfinite(
            observed[name]) else math.nan,
        "calibrated_p": ((1 + exceed[name]) / (1 + total)) if math.isfinite(observed[name])
        else math.nan,
        "permutations": total,
        "usable_permutations": usable[name],
        "reason": "" if math.isfinite(observed[name]) else "no window with usable variance",
    } for name in masks}


def bh_adjust(values):
    values = np.asarray(values, dtype=float)
    out = np.full(values.shape, np.nan)
    finite = np.isfinite(values)
    if np.any(finite):
        out[finite] = false_discovery_control(values[finite], method="bh")
    return out


def write_tsv(path: Path, rows, columns):
    with open(path, "w", newline="\n", encoding="utf-8") as handle:
        handle.write("\t".join(columns) + "\n")
        for row in rows:
            handle.write("\t".join(format_cell(row.get(c)) for c in columns) + "\n")


def format_cell(value):
    if value is None:
        return "NA"
    if isinstance(value, float):
        return "NA" if not math.isfinite(value) else repr(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


PER_MOTIF_COLUMNS = [
    "arm", "motif_key", "RBP", "rbp_table_name", "direction", "direction_label",
    "region", "pooled_region", "plot",
    "native_ranksum_p", "native_argmin_position",
    "calib_observed_min_p", "calibrated_p", "calib_perms_used", "calib_stage", "calib_reason",
    "native_ranksum_p_pooled", "calib_observed_min_p_pooled", "calibrated_p_pooled",
    "calib_perms_used_pooled", "calib_stage_pooled", "calib_pooled_reason",
    "n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
    "fg_proportion", "bg_proportion", "enrichment_ratio",
    "fg_mean_count", "bg_mean_count", "count_ratio",
    "permutation_unit", "permutation_unit_status",
    "n_fg_distinct_target_exons", "n_bg_distinct_target_exons",
    "native_q", "calibrated_q",
]
CONDENSED_COLUMNS = [
    "arm", "RBP", "rbp_table_names", "direction", "direction_label", "pooled_region", "plot",
    "selected_motif_key", "selected_native_region",
    "native_ranksum_p", "native_q_at_selected_region", "calibrated_p", "calibrated_q",
    "calib_perms_used", "calib_stage", "native_argmin_position",
    "n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
    "fg_proportion", "bg_proportion", "enrichment_ratio",
    "fg_mean_count", "bg_mean_count", "count_ratio",
    "n_motifs_total", "n_motifs_calib_q_lt_0.05", "n_motifs_native_q_lt_0.05",
]
POSITIONS_COLUMNS = [
    "arm", "motif_key", "RBP", "direction", "direction_label", "region", "pooled_region",
    "position", "native_ranksum_p", "n_fg_carrying", "n_bg_carrying",
    "fg_total_hits", "bg_total_hits",
]


def effect_at(model: MotifModel, window):
    if window is None:
        return {k: math.nan for k in ("fg_proportion", "bg_proportion", "enrichment_ratio",
                                      "fg_mean_count", "bg_mean_count", "count_ratio",
                                      "n_fg_carrying", "n_bg_carrying")}
    fg_carrying = int(model.fg_carrying[window])
    bg_carrying = int(model.bg_carrying[window])
    fg_proportion = fg_carrying / model.n1
    bg_proportion = bg_carrying / model.n0
    fg_mean = float(model.fg_total[window]) / model.n1
    bg_mean = float(model.bg_total[window]) / model.n0
    return {
        "n_fg_carrying": fg_carrying, "n_bg_carrying": bg_carrying,
        "fg_proportion": fg_proportion, "bg_proportion": bg_proportion,
        "enrichment_ratio": (fg_proportion / bg_proportion) if bg_proportion > 0
        else (math.inf if fg_proportion > 0 else math.nan),
        "fg_mean_count": fg_mean, "bg_mean_count": bg_mean,
        "count_ratio": (fg_mean / bg_mean) if bg_mean > 0
        else (math.inf if fg_mean > 0 else math.nan),
    }


_SAMPLER_CACHE = {}


def cluster_sampler(model: "MotifModel") -> ClusterSampler:
    signature = hashlib.md5(np.ascontiguousarray(model.cluster_codes).tobytes()).hexdigest()
    key = (signature, model.n1)
    if key not in _SAMPLER_CACHE:
        _SAMPLER_CACHE[key] = ClusterSampler(model.cluster_codes, model.n1)
    return _SAMPLER_CACHE[key]


def unit_record(model: "MotifModel", sampler) -> dict:
    return {
        "permutation_unit": model.unit,
        "n_rows_changed": model.n1, "n_rows_background": model.n0,
        "n_rows_changed_before_dedupe": model.n_rows_input[0],
        "n_rows_background_before_dedupe": model.n_rows_input[1],
        "n_distinct_target_exons_changed": model.n_distinct_exons[0],
        "n_distinct_target_exons_background": model.n_distinct_exons[1],
        "n_clusters_drawn": None if sampler is None else sampler.n_clusters_observed,
        "n_clusters_pool": None if sampler is None else sampler.n_clusters_pool,
    }


def run_arm(args, emit):
    arm = args.arm
    counts_dir = Path(args.counts_root) / arm
    released_dir = Path(args.released_root) / arm
    out_dir = Path(args.out_root) / arm
    out_dir.mkdir(parents=True, exist_ok=True)

    alias = load_alias(Path(args.alias_table))
    roots = {d: io.read_root_table(released_dir / ("pVal." + d + ".vs.bg.RNAmap.txt"))
             for d in ("up", "dn")}
    motifs = sorted(p.name[: -len(".counts.npz")] for p in counts_dir.glob("*.counts.npz"))
    if not motifs:
        raise SystemExit("no archives in " + str(counts_dir))
    emit("motifs={} stage1_B={} stage2_B={} threshold={}".format(
        len(motifs), args.permutations, args.refine_perms, args.refine_threshold))

    rows = []
    positions_handle = None
    if args.positions_long:
        positions_handle = open(out_dir / "positions_long.tsv", "w", newline="\n",
                                encoding="utf-8")
        positions_handle.write("\t".join(POSITIONS_COLUMNS) + "\n")
    stage1 = {}
    models_meta = {}
    selection_cache = {}
    unit_audit = {}
    started = time.perf_counter()
    for index, motif in enumerate(motifs, 1):
        table_name = motif.split(".", 1)[0]
        rbp = alias.get(table_name, table_name)
        for direction in ("up", "dn"):
            model = MotifModel(counts_dir / (motif + ".counts.npz"), direction,
                               args.permutation_unit)
            masks = model.window_masks()
            sampler = cluster_sampler(model) if args.permutation_unit == "cluster" else None
            key = (direction, model.n_total, model.n1,
                   sampler.signature if sampler is not None else args.permutation_unit)
            if key not in selection_cache:
                rng = np.random.default_rng(np.random.SeedSequence(
                    [args.seed, 1, io.DIRECTION_CODE[direction]]))
                selection_cache[key] = (sampler.draw(rng, args.permutations) if sampler is not None
                                        else draw_selection(rng, model.n_total, model.n1,
                                                            args.permutations))
                unit_audit[key[:3]] = unit_record(model, sampler)
            result = calibrate(model, masks, selection_cache[key], args.chunk_size)
            stage1[(motif, direction)] = result
            models_meta[(motif, direction)] = {
                "n1": model.n1, "n0": model.n0,
                "argmin": {}, "effect": {}, "rbp": rbp, "table_name": table_name,
                "n_fg_distinct": model.n_distinct_exons[0],
                "n_bg_distinct": model.n_distinct_exons[1],
            }
            native_p = io.p_from_z(model.observed_z)
            for region_id, region in enumerate(io.REGIONS):
                mask = model.region_index == region_id
                window_ids = np.flatnonzero(mask)
                local = native_p[mask]
                best = int(window_ids[int(np.argmin(local))])
                models_meta[(motif, direction)]["argmin"][region] = best
                models_meta[(motif, direction)]["effect"][region] = effect_at(model, best)
                released = roots[direction][motif][region]
                recomputed = float(local.min())
                if args.permutation_unit != "exon-dedupe" and released > 0 and recomputed > 0:
                    if abs(math.log10(recomputed) - math.log10(released)) > 1e-6:
                        raise SystemExit(
                            "stage A invariant broken for {} {} {}".format(motif, direction, region))
                if positions_handle is not None:
                    prefix = "\t".join([arm, motif, rbp, direction,
                                        io.DIRECTION_LABEL[direction], region,
                                        io.REGION_TO_POOL[region]])
                    for slot in window_ids:
                        positions_handle.write(prefix + "\t" + "\t".join([
                            str(int(model.position[slot])), repr(float(native_p[slot])),
                            str(int(model.fg_carrying[slot])), str(int(model.bg_carrying[slot])),
                            str(int(model.fg_total[slot])), str(int(model.bg_total[slot])),
                        ]) + "\n")
            for pool, members in io.POOL_TO_REGIONS.items():
                best = min(members, key=lambda r: roots[direction][motif][r])
                models_meta[(motif, direction)]["argmin"][pool] = (
                    models_meta[(motif, direction)]["argmin"][best])
                models_meta[(motif, direction)]["effect"][pool] = (
                    models_meta[(motif, direction)]["effect"][best])
                models_meta[(motif, direction)]["native_region_" + pool] = best
            del model
        if index % 10 == 0 or index == len(motifs):
            emit("  stage 1: {}/{} motifs, {:.1f}s".format(index, len(motifs),
                                                           time.perf_counter() - started))
    if positions_handle is not None:
        positions_handle.close()
    stage1_seconds = time.perf_counter() - started

    qualifying = sorted(
        ((min(v["calibrated_p"] for v in stage1[k].values()
              if math.isfinite(v["calibrated_p"])), k)
         for k in stage1
         if any(math.isfinite(v["calibrated_p"]) and v["calibrated_p"] <= args.refine_threshold
                for v in stage1[k].values())),
        key=lambda x: (x[0], x[1]))
    selected = [k for _, k in qualifying[: args.refine_max_pairs]]
    deferred = [k for _, k in qualifying[args.refine_max_pairs:]]
    emit("stage 2: {} motif-direction pairs selected, {} deferred by cap".format(
        len(selected), len(deferred)))

    stage = {k: 1 for k in stage1}
    final = {k: dict(v) for k, v in stage1.items()}
    refine_audit = []
    stage2_started = time.perf_counter()
    for done, key in enumerate(selected, 1):
        motif, direction = key
        task_started = time.perf_counter()
        model = MotifModel(counts_dir / (motif + ".counts.npz"), direction,
                           args.permutation_unit)
        masks = model.window_masks()
        sampler = cluster_sampler(model) if args.permutation_unit == "cluster" else None
        rng = np.random.default_rng(np.random.SeedSequence(
            [args.seed, 2, io.DIRECTION_CODE[direction]]))
        exceed = {name: 0 for name in masks}
        observed = statistic_maxima(model.observed_z, masks)
        remaining = args.refine_perms
        while remaining > 0:
            batch = min(args.chunk_size, remaining)
            selection = (sampler.draw(rng, batch) if sampler is not None
                         else draw_selection(rng, model.n_total, model.n1, batch))
            maxima = statistic_maxima(model.permutation_z(selection), masks)
            for name in masks:
                if math.isfinite(observed[name]):
                    values = maxima[name]
                    exceed[name] += int(np.sum(values[np.isfinite(values)] >= observed[name]))
            remaining -= batch
        for name in masks:
            if not math.isfinite(observed[name]):
                continue
            if abs(observed[name] - final[key][name]["observed_max_z"]) > 1e-9:
                raise SystemExit("observed statistic changed during refinement: " + str(key))
            final[key][name] = dict(final[key][name])
            final[key][name]["calibrated_p"] = (1 + exceed[name]) / (1 + args.refine_perms)
            final[key][name]["permutations"] = args.refine_perms
        stage[key] = 2
        refine_audit.append({"motif_key": motif, "direction": direction, "status": "refined",
                             "permutations": args.refine_perms,
                             "wall_seconds": time.perf_counter() - task_started})
        del model
        if done % 10 == 0 or done == len(selected):
            emit("  stage 2: {}/{} pairs, {:.1f}s".format(
                done, len(selected), time.perf_counter() - stage2_started))
    for key in deferred:
        refine_audit.append({"motif_key": key[0], "direction": key[1],
                             "status": "deferred_by_cap",
                             "permutations": args.permutations, "wall_seconds": 0.0})
    stage2_seconds = time.perf_counter() - stage2_started

    for motif in motifs:
        for direction in ("up", "dn"):
            meta = models_meta[(motif, direction)]
            for region in io.REGIONS:
                pool = io.REGION_TO_POOL[region]
                stat = final[(motif, direction)][region]
                pooled = final[(motif, direction)][pool]
                row = {
                    "arm": arm, "motif_key": motif, "RBP": meta["rbp"],
                    "rbp_table_name": meta["table_name"], "direction": direction,
                    "direction_label": io.DIRECTION_LABEL[direction], "region": region,
                    "pooled_region": pool, "plot": pool in io.PLOT_POOLS,
                    "native_ranksum_p": roots[direction][motif][region],
                    "native_argmin_position": int(meta["argmin"][region]),
                    "calib_observed_min_p": stat["observed_min_p"],
                    "calibrated_p": stat["calibrated_p"],
                    "calib_perms_used": stat["permutations"],
                    "calib_stage": stage[(motif, direction)],
                    "calib_reason": stat["reason"],
                    "native_ranksum_p_pooled": min(roots[direction][motif][r]
                                                   for r in io.POOL_TO_REGIONS[pool]),
                    "calib_observed_min_p_pooled": pooled["observed_min_p"],
                    "calibrated_p_pooled": pooled["calibrated_p"],
                    "calib_perms_used_pooled": pooled["permutations"],
                    "calib_stage_pooled": stage[(motif, direction)],
                    "calib_pooled_reason": pooled["reason"],
                    "n_fg_exons": meta["n1"], "n_bg_exons": meta["n0"],
                    "permutation_unit": args.permutation_unit,
                    "permutation_unit_status": args.unit_status,
                    "n_fg_distinct_target_exons": meta["n_fg_distinct"],
                    "n_bg_distinct_target_exons": meta["n_bg_distinct"],
                }
                row.update(meta["effect"][region])
                rows.append(row)

    native_q = bh_adjust([r["native_ranksum_p"] for r in rows])
    for row, q in zip(rows, native_q):
        row["native_q"] = float(q)
    bh_keys = [(m, d, p) for m in motifs for d in ("up", "dn") for p in io.PLOT_POOLS]
    pooled_q = bh_adjust([final[(m, d)][p]["calibrated_p"] for m, d, p in bh_keys])
    q_lookup = {k: float(q) for k, q in zip(bh_keys, pooled_q)}
    emit("BH family for the calibrated layer: {} tests ({} motifs x 2 directions x {} pooled regions)"
         .format(len(bh_keys), len(motifs), len(io.PLOT_POOLS)))
    for row in rows:
        row["calibrated_q"] = q_lookup.get(
            (row["motif_key"], row["direction"], row["pooled_region"]), math.nan)

    by_key = {(r["motif_key"], r["direction"], r["region"]): r for r in rows}
    condensed = []
    rbps = sorted({alias.get(m.split(".", 1)[0], m.split(".", 1)[0]) for m in motifs})
    for rbp in rbps:
        rbp_motifs = [m for m in motifs
                      if alias.get(m.split(".", 1)[0], m.split(".", 1)[0]) == rbp]
        table_names = sorted({m.split(".", 1)[0] for m in rbp_motifs})
        for direction in ("up", "dn"):
            for pool in io.POOL_TO_REGIONS:
                candidates = [(final[(m, direction)][pool]["calibrated_p"],
                               min(roots[direction][m][r] for r in io.POOL_TO_REGIONS[pool]), m)
                              for m in rbp_motifs]
                finite = [c for c in candidates if math.isfinite(c[0])]
                chosen = min(finite, key=lambda c: (c[0], c[1], c[2])) if finite else (
                    math.nan, math.nan, rbp_motifs[0])
                motif = chosen[2]
                member_rows = [by_key[(motif, direction, r)] for r in io.POOL_TO_REGIONS[pool]]
                effect_row = min(member_rows, key=lambda r: r["native_ranksum_p"])
                record = {
                    "arm": arm, "RBP": rbp, "rbp_table_names": ",".join(table_names),
                    "direction": direction, "direction_label": io.DIRECTION_LABEL[direction],
                    "pooled_region": pool, "plot": pool in io.PLOT_POOLS,
                    "selected_motif_key": motif, "selected_native_region": effect_row["region"],
                    "native_ranksum_p": effect_row["native_ranksum_p_pooled"],
                    "native_q_at_selected_region": effect_row["native_q"],
                    "calibrated_p": final[(motif, direction)][pool]["calibrated_p"],
                    "calibrated_q": q_lookup.get((motif, direction, pool), math.nan),
                    "calib_perms_used": final[(motif, direction)][pool]["permutations"],
                    "calib_stage": stage[(motif, direction)],
                    "native_argmin_position": effect_row["native_argmin_position"],
                    "n_motifs_total": len(rbp_motifs),
                    "n_motifs_calib_q_lt_0.05": sum(
                        1 for m in rbp_motifs
                        if math.isfinite(q_lookup.get((m, direction, pool), math.nan))
                        and q_lookup[(m, direction, pool)] < 0.05),
                    "n_motifs_native_q_lt_0.05": sum(
                        1 for m in rbp_motifs
                        if any(math.isfinite(by_key[(m, direction, r)]["native_q"])
                               and by_key[(m, direction, r)]["native_q"] < 0.05
                               for r in io.POOL_TO_REGIONS[pool])),
                }
                for column in ("n_fg_exons", "n_bg_exons", "n_fg_carrying", "n_bg_carrying",
                               "fg_proportion", "bg_proportion", "enrichment_ratio",
                               "fg_mean_count", "bg_mean_count", "count_ratio"):
                    record[column] = effect_row[column]
                condensed.append(record)

    write_tsv(out_dir / "per_motif_regions.tsv", rows, PER_MOTIF_COLUMNS)
    write_tsv(out_dir / "condensed_per_rbp.tsv", condensed, CONDENSED_COLUMNS)
    write_tsv(out_dir / "refinement_tasks.tsv", refine_audit,
              ["motif_key", "direction", "status", "permutations", "wall_seconds"])

    report = {
        "arm": arm, "seed": args.seed, "stage1_permutations": args.permutations,
        "stage2_permutations": args.refine_perms, "refine_threshold": args.refine_threshold,
        "refine_max_pairs": args.refine_max_pairs,
        "pairs_total": len(stage1), "pairs_refined": len(selected),
        "pairs_deferred_by_cap": len(deferred),
        "stage1_wall_seconds": stage1_seconds, "stage2_wall_seconds": stage2_seconds,
        "stage2_p_floor": 1.0 / (1.0 + args.refine_perms),
        "bh_family_size": len(bh_keys),
        "smallest_calibrated_p": min((v for v in
                                      (final[(m, d)][p]["calibrated_p"] for m, d, p in bh_keys)
                                      if math.isfinite(v)), default=None),
        "smallest_calibrated_q": min((v for v in q_lookup.values() if math.isfinite(v)),
                                     default=None),
        "singleton_bh_q_floor": min(1.0, len(bh_keys) / (1.0 + args.refine_perms)),
        "permutation_unit": args.permutation_unit,
        "permutation_unit_status": args.unit_status,
        "permutation_unit_audit": [dict(zip(("direction", "n_total", "n_changed"), k), **v)
                                   for k, v in sorted(unit_audit.items())],
    }
    (out_dir / "refinement_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return out_dir, rows, condensed, report, roots, motifs


def readout_lines(arm, condensed, report):
    lines = [
        "# " + arm + " - permutation-calibrated rank-sum supplement",
        "",
        "Layer: the authors' released rMAPS3 engine with --stat-method mannwhitney. The "
        "statistic is unchanged; only its p-value is recalibrated by permuting the "
        "changed/background exon labels (Westfall-Young min-P over the windows of a region).",
        "Stage 1 used {} permutations at seed {}; {} of {} motif-direction pairs were promoted "
        "to {} permutations because at least one of their twelve statistics reached "
        "stage-1 p <= {}.".format(report["stage1_permutations"], report["seed"],
                                  report["pairs_refined"], report["pairs_total"],
                                  report["stage2_permutations"], report["refine_threshold"]),
        "Adaptive stage selection does not establish super-uniform final p-values or a strict "
        "BH false-discovery-rate guarantee; the calibrated q is a ranking aid with a floor of "
        "{:.3g}.".format(report["singleton_bh_q_floor"]),
        "",
    ]
    for pool in io.PLOT_POOLS:
        for direction in ("up", "dn"):
            subset = [r for r in condensed if r["pooled_region"] == pool
                      and r["direction"] == direction
                      and math.isfinite(r["calibrated_q"])]
            # The calibrated p cannot resolve below its permutation floor, so RBPs
            # tied there are ordered by the native statistic they are calibrating.
            subset.sort(key=lambda r: (r["calibrated_q"], r["calibrated_p"],
                                       r["native_ranksum_p"], r["RBP"]))
            leaders = ", ".join("{} (q={:.3g})".format(r["RBP"], r["calibrated_q"])
                                for r in subset[:3]) or "none available"
            lines.append("- {}, {}: {}; n={} RBPs tested.".format(
                pool, io.DIRECTION_LABEL[direction], leaders, len(subset)))
    lines += [
        "",
        "Both the native released rank-sum p and the calibrated p are reported for every motif, "
        "region and direction. The native p is the number the published figure ranks on; the "
        "calibrated p is the supplement that says how much of that signal survives a permutation "
        "null of the same statistic.",
    ]
    return lines


def write_workbook(path: Path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    for name, (columns, rows) in sheets.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([excel_value(row.get(c)) for c in columns])
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def excel_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return "NA"
    if value is None:
        return "NA"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value


README_ROWS = [
    ("Layer", "Released rMAPS3 engine, --stat-method mannwhitney: one-sided Mann-Whitney U on "
              "per-exon motif hit counts, changed > background, tie-corrected normal "
              "approximation with continuity correction."),
    ("What is calibrated", "That exact statistic. The released per-window p-values and root "
                           "tables are read, never recomputed or rewritten."),
    ("native_ranksum_p", "The released engine value: the smallest p over the windows of the "
                         "sub-region. Anti-conservative in this sparse regime; use for ORDER."),
    ("calibrated_p", "Westfall-Young min-P permutation p for the same regional minimum, from "
                     "permuting changed/background labels among the union of exons."),
    ("calibrated_q", "Benjamini-Hochberg over 126 motifs x 2 directions x 3 pooled regions = "
                     "756 tests."),
    ("native_q", "Benjamini-Hochberg over 126 motifs x 2 directions x 8 sub-regions = 2016 "
                 "released values."),
    ("Pooled regions", "Upstream Intron = UpstreamExonIntron + UpstreamIntron; Exon Body = "
                       "TargetExon_5prime + TargetExon-3prime; Downstream Intron = "
                       "DownstreamIntron + DownstreamExonIntron; Flanking Exon kept as its own "
                       "column, not plotted."),
    ("direction", "up = exons more INCLUDED in the changed set; dn = exons more SKIPPED."),
    ("Effect columns", "Reported at the native argmin window of the sub-region with the smallest "
                       "released p inside the pooled region. fg/bg_proportion are carrying-exon "
                       "fractions; fg/bg_mean_count are hits per exon."),
    ("Stages", "Stage 1 = 2,000 permutations for every pair. Stage 2 = 100,000 permutations for "
               "pairs whose stage-1 p reached the refinement threshold. Adaptive stage selection "
               "voids a strict FDR guarantee."),
    ("Species / assembly", "Human / GRCh38 (hg38) / GENCODE v49; SE (skipped exon) events only."),
    ("Caveat", "A calibrated p at the permutation floor means only that no permutation matched "
               "the observed statistic, not that the true p is that small."),
]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--arm", required=True)
    parser.add_argument("--counts-root", required=True,
                        help="root holding <arm>/*.counts.npz from countdist_to_npz.py")
    parser.add_argument("--released-root", required=True,
                        help="root holding <arm>/pVal.{up,dn}.vs.bg.RNAmap.txt")
    parser.add_argument("--out-root", required=True)
    parser.add_argument("--alias-table", required=True,
                        help="HGNC alias table: table_name<TAB>hgnc_symbol")
    parser.add_argument("--permutations", type=int, default=2000)
    parser.add_argument("--refine-perms", type=int, default=100000)
    parser.add_argument("--refine-threshold", type=float, default=0.005)
    parser.add_argument("--refine-max-pairs", type=int, default=1000)
    parser.add_argument("--chunk-size", type=int, default=500)
    parser.add_argument("--seed", type=int, default=149)
    parser.add_argument("--permutation-unit", choices=list(PERMUTATION_UNITS), default="row",
                        help="row: permute rMATS event rows (historical default). cluster: permute "
                             "target-exon clusters at the observed size composition. exon-dedupe: "
                             "collapse duplicate target exons first, then permute rows.")
    parser.add_argument("--unit-status",
                        default="row-unit sensitivity only; the reportable p is calibrate_ranksum_v2.py",
                        help="free text recorded beside every calibrated p; set to a decided status "
                             "only once the unit question is settled")
    parser.add_argument("--positions-long", action="store_true", default=True)
    parser.add_argument("--no-positions-long", dest="positions_long", action="store_false")
    args = parser.parse_args(argv)

    random.seed(args.seed)
    np.random.seed(args.seed)

    out_dir = Path(args.out_root) / args.arm
    out_dir.mkdir(parents=True, exist_ok=True)
    log = open(out_dir / "command.log", "a", encoding="utf-8")

    def emit(message):
        print(message, flush=True)
        log.write(message + "\n")
        log.flush()

    started = time.perf_counter()
    emit("start={} cmd={}".format(time.strftime("%Y-%m-%dT%H:%M:%S"),
                                  " ".join([sys.executable] + sys.argv)))
    out_dir, rows, condensed, report, roots, motifs = run_arm(args, emit)

    (out_dir / "readout.md").write_text(
        "\n".join(readout_lines(args.arm, condensed, report)) + "\n", encoding="utf-8")
    write_workbook(out_dir / (args.arm + "_calibrated_ranksum.xlsx"), {
        "README": (["field", "description"],
                   [{"field": k, "description": v} for k, v in README_ROWS] + [
                       {"field": "permutation_unit",
                        "description": UNIT_NOTES[args.permutation_unit]},
                       {"field": "permutation_unit_status",
                        "description": args.unit_status}]),
        "condensed_per_rbp": (CONDENSED_COLUMNS, condensed),
        "per_motif_regions": (PER_MOTIF_COLUMNS, rows),
    })

    counts_dir = Path(args.counts_root) / args.arm
    released_dir = Path(args.released_root) / args.arm
    with open(out_dir / "input_md5s.tsv", "w", newline="\n", encoding="utf-8") as handle:
        handle.write("path\tmd5\n")
        for path in [Path(args.alias_table),
                     released_dir / "pVal.up.vs.bg.RNAmap.txt",
                     released_dir / "pVal.dn.vs.bg.RNAmap.txt",
                     counts_dir / "conversion_manifest.tsv",
                     counts_dir / "VERIFY.md"]:
            if path.exists():
                handle.write("{}\t{}\n".format(path, io.md5(path)))
        for motif in motifs:
            path = counts_dir / (motif + ".counts.npz")
            handle.write("{}\t{}\n".format(path, io.md5(path)))

    with open(out_dir / "versions.txt", "w", encoding="utf-8") as handle:
        handle.write("python\t{}\n".format(sys.version.split()[0]))
        handle.write("numpy\t{}\n".format(np.__version__))
        handle.write("scipy\t{}\n".format(scipy.__version__))
        handle.write("openpyxl\t{}\n".format(importlib.metadata.version("openpyxl")))
        handle.write("platform\t{}\n".format(platform.platform()))
        handle.write("script\t{}\n".format(Path(__file__).resolve()))
        handle.write("script_md5\t{}\n".format(io.md5(Path(__file__).resolve())))
        handle.write("module_md5\t{}\n".format(
            io.md5(Path(__file__).resolve().parent / "rmaps_countdist_io.py")))
        handle.write("seed\t{}\n".format(args.seed))
        handle.write("generated\t{}\n".format(time.strftime("%Y-%m-%dT%H:%M:%S")))
        for var in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
            handle.write("{}\t{}\n".format(var, os.environ.get(var, "unset")))

    emit("wall_seconds={:.1f} stage1={:.1f} stage2={:.1f}".format(
        time.perf_counter() - started, report["stage1_wall_seconds"],
        report["stage2_wall_seconds"]))
    emit("exit=0")
    log.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

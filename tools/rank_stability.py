from __future__ import annotations

import os
import sys

sys.dont_write_bytecode = True
for _thread_variable in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS'):
    os.environ[_thread_variable] = '3'

import argparse
import csv
import hashlib
import importlib.metadata
import json
import platform
import random
import re
import subprocess
import time
from pathlib import Path

import numpy as np
from scipy.stats import rankdata

SUMMARY_PATH = Path(r'E:\rmaps_summ\scripts\summarize_rmaps_regions_v5.py')
ENGINE_ROOT = Path(r'E:\Claude\rMAPS3_fix')
ALIASES_PATH = Path(r'F:\rMAPS\scripts\rbp_alias_hgnc_2026-09-17.tsv')
FIGURE_SKILL = Path(r'C:\Users\ambur\.claude\skills\rnaseq-figure-style\SKILL.md')
XLSX_SKILL = Path(r'C:\Users\ambur\.codex\plugins\cache\anthropic-agent-skills\document-skills\local\skills\xlsx\SKILL.md')
sys.path.insert(0, str(SUMMARY_PATH.parent))
import summarize_rmaps_regions_v5 as native

POOLS = {k: native.POOL_TO_REGIONS[k] for k in ('Upstream Intron', 'Exon Body', 'Downstream Intron')}
DIRECTIONS = {'up': 0, 'dn': 1}
COLUMNS = ['arm', 'direction', 'direction_label', 'pooled_region', 'RBP', 'observed_raw_p',
           'observed_rank', 'median_rank', 'rank_p05', 'rank_p95', 'top3_frequency',
           'top5_frequency', 'top10_frequency', 'n_motifs']


def bootstrap_weights(n, boot, rng):
    if n < 1 or boot < 1:
        raise ValueError('n and boot must be positive')
    return rng.multinomial(n, np.full(n, 1.0 / n), size=boot).astype(np.float64)


def weighted_counts(W, H, E):
    W, H, E = np.asarray(W), np.asarray(H), np.asarray(E)
    if H.ndim != 2 or H.shape != E.shape or W.ndim != 2 or W.shape[1] != H.shape[0]:
        raise ValueError('Expected W[resample,event] and matching H/E[event,window]')
    if np.any(~np.isfinite(W)) or np.any(W < 0) or np.any(W != np.rint(W)):
        raise ValueError('Weights must be finite nonnegative integers')
    if not np.all(np.isin(H, [0, 1])) or not np.all(np.isin(E, [0, 1])) or np.any(H > E):
        raise ValueError('Hits and eligibility must be binary; hits require eligibility')
    joined = np.concatenate((H, E), axis=1).astype(np.float64)
    counts = np.rint(W.astype(np.float64, copy=False) @ joined).astype(np.int64)
    return counts[:, :H.shape[1]], counts[:, H.shape[1]:]


def fixed_bg_pvalues(a, n, c, m):
    a, n, c, m = np.broadcast_arrays(*[np.asarray(x, dtype=np.int64) for x in (a, n, c, m)])
    valid = (n > 0) & (m > 0) & (a >= 0) & (c >= 0) & (a <= n) & (c <= m)
    out = np.full(a.shape, np.nan)
    if np.any(valid):
        tuples = np.column_stack([x[valid] for x in (a, n, c, m)])
        unique, inverse = np.unique(tuples, axis=0, return_inverse=True)
        values = native.greater_pvalues(*unique.T)
        out[valid] = values[inverse]
    return out


def rank_matrix(p):
    p = np.asarray(p, dtype=float)
    if p.ndim < 1 or p.shape[-1] == 0 or np.any(~np.isfinite(p)) or np.any((p < 0) | (p > 1)):
        raise ValueError('Every ranked RBP must have a finite raw p in [0,1]; no missing-rank imputation')
    return rankdata(p, axis=-1, method='average')


def read_aliases(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.DictReader(handle, delimiter='\t'))
    if not rows or not {'table_name', 'hgnc_symbol'}.issubset(rows[0]):
        raise ValueError('Alias table requires table_name and hgnc_symbol')
    result = {}
    for row in rows:
        key, value = row['table_name'].strip(), row['hgnc_symbol'].strip()
        if not key or not value or (key in result and result[key] != value):
            raise ValueError(f'Invalid or ambiguous alias row: {key!r}')
        result[key] = value
    return result


def display_name(motif, aliases):
    raw = motif.split('.', 1)[0]
    return 'ESRP-like' if re.fullmatch(r'motif_\d+', raw) else aliases.get(raw, raw)


def sha256(path):
    result = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            result.update(block)
    return result.hexdigest()


def minimum_valid(p):
    p = np.asarray(p)
    valid = np.any(np.isfinite(p), axis=-1)
    output = np.full(p.shape[:-1], np.nan)
    output[valid] = np.nanmin(p[valid], axis=-1)
    return output


def write_table(path, rows, columns):
    with Path(path).open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter='\t')
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path, rows, definitions):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    readme = wb.active
    readme.title = 'README'
    readme.append(['Field', 'Definition'])
    for key, value in definitions.items():
        readme.append([key, str(value)])
    readme.column_dimensions['A'].width = 27
    readme.column_dimensions['B'].width = 115
    for row in readme.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        readme.row_dimensions[row[0].row].height = 42
    ws = wb.create_sheet('Rank stability')
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
    for col, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = max(16, len(name) + 2)
    ws.auto_filter.ref = ws.dimensions
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Arial', size=10)
        for cell in sheet[1]:
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0072B2')
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = '0.000E+00'
        for cell in row[6:10]:
            cell.number_format = '0.00'
        for cell in row[10:13]:
            cell.number_format = '0.0%'
    wb.save(path)


def write_figure(out, arm, rows, boot):
    os.environ['MPLCONFIGDIR'] = str(out / '.matplotlib')
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt, font_manager
    font_path = font_manager.findfont('Arial', fallback_to_default=False)
    plt.rcParams.update({'font.family': 'Arial', 'svg.fonttype': 'none', 'font.size': 9,
                         'axes.titlesize': 10, 'axes.labelsize': 9, 'axes.spines.top': False,
                         'axes.spines.right': False, 'axes.linewidth': 0.75})
    fig, axes = plt.subplots(2, 3, figsize=(13.5, 8.5))
    fig.subplots_adjust(left=.09, right=.98, bottom=.10, top=.78, wspace=.53, hspace=.50)
    fig.suptitle(f'{arm}: RBP rank stability', y=.97, fontsize=15)
    fig.text(.5, .905, 'Bar = fraction of foreground-exon bootstrap resamples with average RBP rank ≤ 10.',
             ha='center', fontsize=11)
    fig.text(.5, .025, f'Shown: observed top 10 per panel; {boot} resamples; fixed background; stability of order, not significance.',
             ha='center', fontsize=9)
    for ri, direction in enumerate(DIRECTIONS):
        for ci, pool in enumerate(POOLS):
            ax = axes[ri, ci]
            subset = sorted([r for r in rows if r['direction'] == direction and r['pooled_region'] == pool],
                            key=lambda r: (r['observed_rank'], r['RBP']))[:10]
            ax.barh(np.arange(len(subset)), [r['top10_frequency'] for r in subset],
                    color='#D55E00' if direction == 'up' else '#0072B2', height=.65)
            ax.set_yticks(np.arange(len(subset)), [r['RBP'] for r in subset])
            ax.invert_yaxis()
            ax.set_xlim(0, 1.05)
            ax.set_xticks([0, .25, .5, .75, 1])
            ax.set_xlabel('Top-10 frequency')
            ax.set_title(f'{pool} · {native.DIRECTION_LABEL[direction].lower()}', pad=10)
    stem = out / f'{arm}_rank_stability'
    fig.savefig(stem.with_suffix('.png'), dpi=600, facecolor='white')
    fig.savefig(stem.with_suffix('.svg'), facecolor='white')
    plt.close(fig)
    return font_path


def run_analysis(args):
    started = time.perf_counter()
    run, out, aliases_path = args.run.resolve(), args.out.resolve(), args.aliases.resolve()
    if out == run or run in out.parents or Path(r'E:\rmaps_summ') == out or Path(r'E:\rmaps_summ') in out.parents:
        raise ValueError('Output must be outside read-only input and production summarizer directories')
    out.mkdir(parents=True, exist_ok=True)
    prefix = out / f'{args.arm}_rank_stability.tsv'
    if prefix.exists() or (out / 'run_summary.json').exists():
        raise FileExistsError(f'Preserving existing analysis outputs in {out}; choose a new --out')
    command = subprocess.list2cmdline([sys.executable, str(Path(__file__).resolve()), '--run', str(run),
                                     '--arm', args.arm, '--out', str(out), '--boot', str(args.boot),
                                     '--seed', str(args.seed), '--aliases', str(aliases_path)])
    (out / 'command.log').write_text(command + '\n', encoding='utf-8')
    versions = {'Python': platform.python_version(), **{p: importlib.metadata.version(p)
                for p in ('numpy', 'scipy', 'openpyxl', 'matplotlib')}}
    versions.update({x: os.environ[x] for x in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS')})
    (out / 'versions.txt').write_text(''.join(f'{k}\t{v}\n' for k, v in versions.items()), encoding='utf-8')
    random.seed(args.seed)
    np.random.seed(args.seed)
    manifest_path = run / 'run_manifest.json'
    manifest = native.read_manifest(manifest_path)
    parameters = manifest['parameters']
    if parameters.get('genome') != 'hg38' or parameters.get('fisher_alternative') != 'greater':
        raise ValueError('Manifest must specify genome=hg38 and fisher_alternative=greater')
    if manifest.get('statistical_method') != 'fisher':
        raise ValueError('Manifest must specify statistical_method=fisher')
    sources = [Path(__file__).resolve(), SUMMARY_PATH, ENGINE_ROOT / 'rmaps_core' / 'positional_io.py',
               aliases_path, manifest_path, FIGURE_SKILL, XLSX_SKILL]
    roots = {}
    for direction in DIRECTIONS:
        path = run / f'pVal.{direction}.vs.bg.RNAmap.txt'
        sources.append(path)
        roots[direction] = native.read_root_table(path)
    sources_before = {str(p): sha256(p) for p in sources}
    if not roots['up'] or set(roots['up']) != set(roots['dn']):
        raise ValueError('Root motif tables must have identical, nonempty motif key sets')
    motifs = sorted(roots['up'])
    aliases = read_aliases(aliases_path)
    groups = {}
    for motif in motifs:
        groups.setdefault(display_name(motif, aliases), []).append(motif)
    rbps = sorted(groups)
    rbp_index = {rbp: i for i, rbp in enumerate(rbps)}
    expected = {p['path'].replace('\\', '/'): p for p in manifest['outputs']}
    regions = [region for members in POOLS.values() for region in members]
    paths = { (motif, region): run / 'positional' / f'{motif}.{region}.hits.npz'
             for motif in motifs for region in regions }
    for path in paths.values():
        relative = path.relative_to(run).as_posix()
        if not path.is_file() or relative not in expected or not expected[relative].get('sha256'):
            raise FileNotFoundError(f'Required manifested hit archive missing: {path}')
    first_hits, first_eligible, labels = native.load_sparse_hits(paths[motifs[0], regions[0]], ENGINE_ROOT)
    labels = labels.copy()
    del first_hits, first_eligible
    fg_indices = {d: np.flatnonzero(labels == code) for d, code in DIRECTIONS.items()}
    bg_indices = np.flatnonzero(labels == 2)
    if not len(bg_indices) or any(not len(x) for x in fg_indices.values()):
        raise ValueError('Both foreground directions and background must contain events')
    weights = {d: np.vstack([np.ones((1, len(ix))), bootstrap_weights(
                    len(ix), args.boot, np.random.default_rng(np.random.SeedSequence([args.seed, DIRECTIONS[d]])))])
               for d, ix in fg_indices.items()}
    np.savez_compressed(out / 'bootstrap_weights.npz', up=weights['up'], dn=weights['dn'], set_labels=labels)
    pooled = {(d, pool): np.full((args.boot + 1, len(rbps)), np.nan)
              for d in DIRECTIONS for pool in POOLS}
    archive_records, motif_records = [], []
    phase = {'preflight_seconds': time.perf_counter() - started, 'load_hash_seconds': 0.,
             'counts_seconds': 0., 'fisher_seconds': 0.}
    count_done = 0
    eligibility_hashes = {}
    for pool, members in POOLS.items():
        for motif in motifs:
            minima = {d: np.full(args.boot + 1, np.nan) for d in DIRECTIONS}
            for region in members:
                tick = time.perf_counter()
                path = paths[motif, region]
                digest = sha256(path)
                record = expected[path.relative_to(run).as_posix()]
                if digest != record['sha256'] or path.stat().st_size != record['size_bytes']:
                    raise ValueError(f'Archive differs from manifest: {path}')
                hits, eligible, these_labels = native.load_sparse_hits(path, ENGINE_ROOT)
                if not np.array_equal(these_labels, labels):
                    raise ValueError(f'Event label/order mismatch: {path}')
                if hits.shape[1] != native.expected_windows(manifest, region):
                    raise ValueError(f'Unexpected window count: {path}')
                eh = hashlib.sha256(eligible.tobytes()).hexdigest()
                if region in eligibility_hashes and eligibility_hashes[region] != eh:
                    raise ValueError(f'Eligibility mismatch across motifs: {path}')
                eligibility_hashes[region] = eh
                archive_records.append({'path': str(path), 'sha256': digest, 'size_bytes': path.stat().st_size})
                phase['load_hash_seconds'] += time.perf_counter() - tick
                c = hits[bg_indices].sum(axis=0, dtype=np.int64)
                m = eligible[bg_indices].sum(axis=0, dtype=np.int64)
                for d, ix in fg_indices.items():
                    tick = time.perf_counter()
                    a, n = weighted_counts(weights[d], hits[ix], eligible[ix])
                    phase['counts_seconds'] += time.perf_counter() - tick
                    tick = time.perf_counter()
                    region_min = minimum_valid(fixed_bg_pvalues(a, n, c, m))
                    phase['fisher_seconds'] += time.perf_counter() - tick
                    minima[d] = np.fmin(minima[d], region_min)
                del hits, eligible
            index = rbp_index[display_name(motif, aliases)]
            for d in DIRECTIONS:
                pooled[d, pool][:, index] = np.fmin(pooled[d, pool][:, index], minima[d])
                motif_records.append({'direction': d, 'pooled_region': pool, 'motif_key': motif,
                                      'RBP': display_name(motif, aliases), 'observed_raw_p': float(minima[d][0])})
            count_done += 1
            if count_done % 20 == 0 or count_done == len(motifs) * len(POOLS):
                print(f'Completed {count_done}/{len(motifs) * len(POOLS)} motif pools; '
                      f'elapsed {time.perf_counter() - started:.1f} s', flush=True)
    compute_end = time.perf_counter()
    rows, audit_arrays, stable_counts = [], {}, {}
    for d in DIRECTIONS:
        for pi, pool in enumerate(POOLS):
            raw = pooled[d, pool]
            if np.any(~np.isfinite(raw)):
                bad = np.argwhere(~np.isfinite(raw))[0]
                raise ValueError(f'No valid window for RBP {rbps[bad[1]]}, direction {d}, pool {pool}, row {bad[0]}')
            ranks = rank_matrix(raw)
            audit_arrays[f'{d}_pool{pi}_raw_p'] = raw
            audit_arrays[f'{d}_pool{pi}_ranks'] = ranks
            boot_ranks = ranks[1:]
            quantiles = np.percentile(boot_ranks, [5, 50, 95], axis=0, method='linear')
            for i, rbp in enumerate(rbps):
                rows.append({'arm': args.arm, 'direction': d, 'direction_label': native.DIRECTION_LABEL[d],
                             'pooled_region': pool, 'RBP': rbp, 'observed_raw_p': float(raw[0, i]),
                             'observed_rank': float(ranks[0, i]), 'median_rank': float(quantiles[1, i]),
                             'rank_p05': float(quantiles[0, i]), 'rank_p95': float(quantiles[2, i]),
                             'top3_frequency': float(np.mean(boot_ranks[:, i] <= 3)),
                             'top5_frequency': float(np.mean(boot_ranks[:, i] <= 5)),
                             'top10_frequency': float(np.mean(boot_ranks[:, i] <= 10)),
                             'n_motifs': len(groups[rbp])})
    rows.sort(key=lambda r: (list(DIRECTIONS).index(r['direction']), list(POOLS).index(r['pooled_region']),
                             r['observed_rank'], r['RBP']))
    for d in DIRECTIONS:
        for pool in POOLS:
            selected = [r for r in rows if r['direction'] == d and r['pooled_region'] == pool][:10]
            stable_counts[f'{pool} | {d}'] = sum(r['top10_frequency'] >= .8 for r in selected)
    np.savez_compressed(out / 'rank_bootstrap.npz', rbps=np.array(rbps), pools=np.array(list(POOLS)), **audit_arrays)
    definitions = {
        'Purpose': 'Measures stability of ORDER under resampling of the changed exons, not significance.',
        'Sampling': f'{args.boot} nonparametric bootstrap resamples per direction, foreground events sampled with replacement; background fixed. Seed {args.seed}.',
        'Shared weights': 'One multinomial weight matrix per direction, reused across every motif and native region. SeedSequence([seed,direction code]); up=0, dn=1. Audit row 0 is all ones (observed).',
        'Direction': 'up = included (set label 0); dn = skipped (set label 1); background = label 2.',
        'Statistic': 'Exact one-sided greater Fisher p per eligible window; minimum across windows, the two native regions in each pool, then all motifs assigned to the RBP.',
        'Pools': '; '.join(f'{p}: {" + ".join(r)}' for p, r in POOLS.items()),
        'RBP grouping': f'Key text before first dot, mapped through {aliases_path}; motif_N = ESRP-like. n_motifs counts keys after alias grouping.',
        'Ranks': 'Ascending raw pooled p; ties receive average rank. Ranks are recalculated across every RBP within each panel and resample.',
        'observed_rank': 'Rank from original foreground events (all weights equal 1), with no resampling.',
        'rank intervals': 'Median and 5th/95th percentiles across bootstrap ranks; NumPy linear interpolation. These describe bootstrap rank variability.',
        'topK_frequency': 'Fraction of all bootstrap resamples with average rank <= K. Ties at a threshold are handled by average rank, not arbitrary tie breaking.',
        'Observed top 10': 'First 10 RBPs sorted by observed average rank then alphabetically by RBP (deterministic display tie breaking only).',
        'Missing windows': 'Windows with no eligible foreground or background are excluded from minima. Abort if any RBP has no valid pooled value in any resample; no imputation.',
        'Scope': 'Assigned scope: Homo sapiens / GRCh38 (hg38) / GENCODE v49. Input manifest confirms hg38; annotation release is not independently recorded in this manifest.',
        'Reference provenance': f"Engine motif references: {parameters.get('knownMotifs', 'missing')}; {parameters.get('motif', 'missing')}; FASTA root {parameters.get('fastaRoot', 'missing')}. No new motif scan or annotation loading.",
        'Run': str(run), 'Software': str(Path(__file__).resolve()),
        'Audit arrays': 'bootstrap_weights.npz has observed row 0 + resamples and label vector; rank_bootstrap.npz has raw pooled p and ranks with RBP and pool axes.',
        'Projection': 'Measured total wall x 88631 / input total events. Heuristic only: fixed setup/figure costs and Fisher-tail dependence on foreground counts mean actual runtime need not scale linearly.'}
    write_table(prefix, rows, COLUMNS)
    write_table(out / 'observed_per_motif.tsv', motif_records,
                ['direction', 'pooled_region', 'motif_key', 'RBP', 'observed_raw_p'])
    write_workbook(out / f'{args.arm}_rank_stability.xlsx', rows, definitions)
    (out / 'README.md').write_text('\n'.join(f'- **{k}:** {v}' for k, v in definitions.items()) + '\n', encoding='utf-8')
    font_path = write_figure(out, args.arm, rows, args.boot)
    sources_after = {str(p): sha256(p) for p in sources}
    if sources_before != sources_after:
        raise RuntimeError('Source changed during run; outputs are not approved for use')
    provenance = {'source_sha256_before': sources_before, 'source_sha256_after': sources_after,
                  'archive_manifest_validation': 'all six required regions of every motif verified before loading',
                  'archives': archive_records, 'engine_git_revision_from_manifest': manifest.get('git_revision'),
                  'reference_paths_from_manifest': {k: parameters.get(k, 'missing') for k in ('knownMotifs', 'motif', 'fastaRoot', 'genome')},
                  'font_path': font_path, 'definitions': definitions}
    (out / 'provenance.json').write_text(json.dumps(provenance, indent=2) + '\n', encoding='utf-8')
    wall = time.perf_counter() - started
    phase.update({'compute_seconds': compute_end - started, 'export_seconds': time.perf_counter() - compute_end})
    qki = {f"{r['pooled_region']} | {r['direction']}": r['top3_frequency'] for r in rows if r['RBP'] == 'QKI'}
    summary = {'arm': args.arm, 'boot': args.boot, 'seed': args.seed, 'wall_seconds': wall, 'total_events': len(labels),
               'foreground_counts': {d: len(ix) for d, ix in fg_indices.items()}, 'background_count': len(bg_indices),
               'n_motifs': len(motifs), 'n_rbps': len(rbps), 'n_rows': len(rows), 'qki_top3_frequency': qki,
               'stable_top10_counts': stable_counts, 'stable_top10_range': [min(stable_counts.values()), max(stable_counts.values())],
               'projected_seconds': wall * 88631 / len(labels), 'projection_target_events': 88631,
               'projection_caveat': definitions['Projection'], 'phase_seconds': phase, 'source_hashes_unchanged': True}
    (out / 'run_summary.json').write_text(json.dumps(summary, indent=2) + '\n', encoding='utf-8')
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description='Foreground-bootstrap stability of RBP order; fixed background, raw Fisher minima.')
    parser.add_argument('--run', required=True, type=Path)
    parser.add_argument('--arm', required=True)
    parser.add_argument('--out', required=True, type=Path)
    parser.add_argument('--boot', type=int, default=300)
    parser.add_argument('--seed', type=int, default=149)
    parser.add_argument('--aliases', type=Path, default=ALIASES_PATH)
    args = parser.parse_args(argv)
    if args.boot < 1 or args.seed < 0:
        parser.error('--boot must be positive and --seed nonnegative')
    if not re.fullmatch(r'[A-Za-z0-9_-]+', args.arm):
        parser.error('--arm must contain only letters, digits, underscore or hyphen')
    print(json.dumps(run_analysis(args), indent=2), flush=True)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

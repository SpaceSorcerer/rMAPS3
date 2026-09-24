"""Argument, filter, summary and permutation-unit logic of tools/rmaps3_skill_run.py."""
import json
import os
import subprocess
import sys
from pathlib import Path

import numpy as np
import pytest

from test_engine_synthetic import synthetic, input_paths  # noqa: F401

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import calibrate_ranksum as calib  # noqa: E402
import rmaps_countdist_io as io  # noqa: E402
import rmaps3_skill_run as skill  # noqa: E402


def parse(*argv):
    return skill.build_parser().parse_args(list(argv))


BASE = ("--mode", "quick", "--arm", "A1", "--out", "out", "--no-figures")


# ------------------------------------------------------------------ argument logic
def test_input_modes_are_exclusive_and_complete(tmp_path):
    with pytest.raises(ValueError, match="either --up/--dn/--bg or --rmats-se"):
        skill.validate(parse(*BASE))
    with pytest.raises(ValueError, match="either --up/--dn/--bg or --rmats-se"):
        skill.validate(parse(*BASE, "--up", "u", "--dn", "d", "--bg", "b",
                             "--rmats-se", "se.txt", "--filter", "f.json"))
    with pytest.raises(ValueError, match="must be given together"):
        skill.validate(parse(*BASE, "--up", "u", "--dn", "d"))
    with pytest.raises(ValueError, match="needs --filter"):
        skill.validate(parse(*BASE, "--rmats-se", "se.txt"))
    skill.validate(parse(*BASE, "--up", "u", "--dn", "d", "--bg", "b"))
    skill.validate(parse(*BASE, "--rmats-se", "se.txt", "--filter", "f.json"))


def test_arm_must_be_filename_safe():
    for bad in ("..", "a/b", "../x", ""):
        with pytest.raises(ValueError, match="filename-safe"):
            skill.validate(parse("--mode", "quick", "--arm", bad, "--out", "o",
                                 "--up", "u", "--dn", "d", "--bg", "b"))


def test_geometry_defaults_are_the_locked_values():
    args = parse(*BASE, "--up", "u", "--dn", "d", "--bg", "b")
    assert (args.window, args.step, args.intron, args.exon) == (50, 1, 250, 50)
    assert args.stat_method == "mannwhitney"
    assert args.engine == "released"
    assert args.workers == 1
    assert args.calib_unit == "cluster"
    assert args.seed == 149
    assert (args.permutations, args.refine_perms) == (2000, 100000)


def test_nonpositive_geometry_is_refused():
    with pytest.raises(ValueError, match="--window must be positive"):
        skill.validate(parse(*BASE, "--up", "u", "--dn", "d", "--bg", "b", "--window", "0"))


def test_mouse_scope_is_announced_loudly(capsys):
    skill.validate(parse(*BASE, "--up", "u", "--dn", "d", "--bg", "b",
                         "--genome", "mm10", "--species", "Mus musculus"))
    assert "MOUSE REFERENCE IN PLAY" in capsys.readouterr().err


def test_output_directory_must_be_absent_or_empty(tmp_path):
    (tmp_path / "already").mkdir()
    (tmp_path / "already" / "x.txt").write_text("x")
    with pytest.raises(ValueError, match="Refusing non-empty"):
        skill.fresh_directory(tmp_path / "already")
    skill.fresh_directory(tmp_path / "fresh")


# ------------------------------------------------------------------ filter spec
def write_spec(tmp_path, spec, name="filter.json"):
    path = tmp_path / name
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


def translate(tmp_path, spec, rmats="se.txt"):
    """Run the spec through build_event_sets up to the point it writes the builder config."""
    calls = []

    def fake_step(name, command, log, env, cwd, out):
        calls.append(command)
        return 0.0

    original = skill.run_step
    skill.run_step = fake_step
    try:
        args = parse(*BASE[:4], "--out", str(tmp_path), "--rmats-se", rmats,
                     "--filter", str(write_spec(tmp_path, spec)))
        skill.build_event_sets(args, tmp_path, lambda m: None, {})
    finally:
        skill.run_step = original
    return json.loads((tmp_path / "event_set_config.json").read_text()), calls


def test_filter_keys_are_translated_to_builder_gates(tmp_path):
    config, calls = translate(tmp_path, {"fdr": 0.05, "min_abs_dpsi": 0.10,
                                         "min_jc_per_sample": 10, "bg_fdr_min": 0.5})
    assert config["gates"] == {"fdr": 0.05, "abs_dpsi": 0.10,
                               "min_sample_coverage": 10, "background_fdr": 0.5}
    assert "build_event_sets.py" in " ".join(calls[0])


def test_expression_floor_requires_an_expression_table(tmp_path):
    with pytest.raises(ValueError, match="base_mean_floor needs expr_table"):
        translate(tmp_path, {"fdr": 0.05, "base_mean_floor": 50})
    config, _ = translate(tmp_path, {"fdr": 0.05, "base_mean_floor": 50,
                                     "expr_table": str(tmp_path / "deseq2.csv")})
    assert config["gates"]["base_mean_floor"] == 50
    assert config["inputs"]["deseq2"].endswith("deseq2.csv")


def test_unknown_filter_keys_are_refused_not_ignored(tmp_path):
    with pytest.raises(ValueError, match="Unknown filter keys"):
        translate(tmp_path, {"fdr": 0.05, "pvalue": 0.01})


def test_yaml_spec_extension_is_accepted_or_refused_clearly(tmp_path):
    path = tmp_path / "f.txt"
    path.write_text("{}")
    with pytest.raises(ValueError, match="must be .json"):
        skill.load_spec(path)


# ------------------------------------------------------------------ summary tables
def synthetic_engine_run(tmp_path, arm="A1", n_windows=None):
    """Root tables plus matching count archives for two motifs, in the engine's own layout."""
    per_region = 3
    n_windows = n_windows or per_region * len(io.REGIONS)
    engine = tmp_path / "engine" / arm
    (engine / "temp").mkdir(parents=True)
    counts_root = tmp_path / "counts"
    (counts_root / arm).mkdir(parents=True)
    rng = np.random.default_rng(149)
    motifs = ["QKI.ACTAAC[ACG]", "OTHER.TTTT"]
    roots = {"up": {}, "dn": {}}
    region_index = np.repeat(np.arange(len(io.REGIONS), dtype=np.int8), per_region)
    position = np.tile(np.arange(per_region, dtype=np.int32), len(io.REGIONS))
    for index, motif in enumerate(motifs):
        store = {"schema_version": np.asarray(io.SCHEMA_VERSION), "arm": np.asarray(arm),
                 "motif": np.asarray(motif), "region_name": np.asarray(io.REGIONS, dtype="U"),
                 "region_of_position": region_index, "position": position}
        matrices = {}
        for group, n_rows, rate in (("up", 12, 0.9 if index == 0 else 0.1),
                                    ("dn", 11, 0.1), ("bg", 60, 0.1)):
            matrix = rng.poisson(rate, size=(n_rows, n_windows)).astype(np.int16)
            matrices[group] = matrix
            store[group + "_exon_id"] = np.asarray(
                ["chr1:+:{}:{}:1:2:3:4".format(1000 + i, 1100 + i) for i in range(n_rows)],
                dtype="U")
            io.pack_group(store, group, matrix)
        np.savez_compressed(counts_root / arm / (motif + ".counts.npz"), **store)
        for direction in ("up", "dn"):
            # Deterministic, distinct p per sub-region so the pooled minimum is checkable.
            # QKI (index 0) is the most significant everywhere, so the panel order is checkable.
            roots[direction][motif] = {
                region: 10.0 ** -(40 - index * 20 - r - (0 if direction == "up" else 1))
                for r, region in enumerate(io.REGIONS)}
    for direction in ("up", "dn"):
        path = engine / f"pVal.{direction}.vs.bg.RNAmap.txt"
        with open(path, "w", newline="\n", encoding="utf-8") as handle:
            handle.write("\t".join(["RBP"] + [io.ROOT_COLUMNS[r] for r in io.REGIONS]) + "\n")
            for motif in motifs:
                handle.write("\t".join([motif] + [repr(roots[direction][motif][r])
                                                  for r in io.REGIONS]) + "\n")
    return engine, counts_root, roots, motifs


def fake_counts(n_up=12, n_dn=11, n_bg=60):
    return {key: {"n_events": n, "renamed": {}, "output": {"path": f"{key}.coord.txt"}}
            for key, n in (("up", n_up), ("dn", n_dn), ("bg", n_bg))}


def test_quick_tables_take_pooled_minima_from_the_root_tables(tmp_path):
    engine, counts_root, roots, motifs = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b")
    per_motif, rbp_rows, seen = skill.quick_tables(args, engine, counts_root, fake_counts(), {})
    assert sorted(seen) == sorted(motifs)
    row = next(r for r in per_motif["up"] if r["motif_key"] == "QKI.ACTAAC[ACG]")
    for pool, members in io.POOL_TO_REGIONS.items():
        expected = min(roots["up"]["QKI.ACTAAC[ACG]"][m] for m in members)
        assert row[f"{pool} | min_p"] == expected
        assert row[f"{pool} | sub_region"] in members
    assert row["n_changed_events"] == 12 and row["n_background_events"] == 60


def test_quick_tables_carry_the_engines_own_motif_scores(tmp_path):
    engine, counts_root, _, _ = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b")
    per_motif, _, _ = skill.quick_tables(args, engine, counts_root, fake_counts(), {})
    row = next(r for r in per_motif["up"] if r["motif_key"] == "QKI.ACTAAC[ACG]")
    ratio = row["Upstream Intron | count_ratio"]
    assert ratio is not None and ratio > 1.0, "planted enrichment must raise the count ratio"
    assert row["Upstream Intron | fg_mean_count"] > row["Upstream Intron | bg_mean_count"]


def test_motif_scores_are_na_without_archives(tmp_path):
    engine, _, _, _ = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b")
    per_motif, _, _ = skill.quick_tables(args, engine, None, fake_counts(), {})
    assert per_motif["up"][0]["Upstream Intron | count_ratio"] is None


def test_rbp_sheet_ranks_by_best_motif_and_applies_the_alias(tmp_path):
    engine, counts_root, _, _ = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b")
    _, rbp_rows, _ = skill.quick_tables(args, engine, counts_root, fake_counts(),
                                        {"OTHER": "RENAMED"})
    panel = sorted((r for r in rbp_rows
                    if r["direction_label"] == "INCLUDED" and r["pooled_region"] == "Upstream Intron"),
                   key=lambda r: r["rank"])
    assert [r["RBP"] for r in panel] == ["QKI", "RENAMED"]
    assert panel[0]["min_p"] <= panel[1]["min_p"]
    assert panel[0]["selected_motif_key"] == "QKI.ACTAAC[ACG]"


def test_positive_control_passes_only_at_rank_one(tmp_path):
    engine, counts_root, _, _ = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b")
    _, rbp_rows, _ = skill.quick_tables(args, engine, counts_root, fake_counts(), {})
    panels = [("INCLUDED", "Upstream Intron")]
    assert all(r["pass"] for r in skill.positive_control(rbp_rows, "QKI", panels))
    other = skill.positive_control(rbp_rows, "OTHER", panels)
    assert not other[0]["pass"] and other[0]["symbol_rank"] == 2
    absent = skill.positive_control(rbp_rows, "NOT_A_MOTIF", panels)
    assert not absent[0]["pass"] and absent[0]["symbol_rank"] is None


def test_readme_names_the_statistic_caveat_and_the_counted_unit(tmp_path):
    engine, _, _, motifs = synthetic_engine_run(tmp_path)
    args = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b",
                 "--genome-root", "genomes")
    rows = skill.readme_rows(args, engine, ROOT, fake_counts(), {"reason": "x"}, motifs,
                             Path("known.txt"), Path("esrp.txt"))
    text = json.dumps(rows)
    assert "anti-conservative" in text
    assert "one rMATS SE event, not one distinct target exon" in text
    fisher = parse(*BASE[:4], "--out", str(tmp_path), "--up", "u", "--dn", "d", "--bg", "b",
                   "--genome-root", "genomes", "--stat-method", "fisher")
    fisher_text = json.dumps(skill.readme_rows(fisher, engine, ROOT, fake_counts(), {}, motifs,
                                               Path("k"), Path("e")))
    assert "counts MOTIF HITS, not exons" in fisher_text


def test_workbook_round_trips_every_sheet(tmp_path):
    import openpyxl
    path = tmp_path / "book.xlsx"
    skill.write_workbook(path, {"README": (["field", "description"],
                                           [{"field": "a", "description": "b"}]),
                                "INCLUDED": (["x", "y"], [{"x": 1, "y": None},
                                                          {"x": 2, "y": float("inf")}])})
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["README", "INCLUDED"]
    assert [c.value for c in book["INCLUDED"][2]] == [1, "NA"]
    assert [c.value for c in book["INCLUDED"][3]] == [2, "NA"]
    assert book["INCLUDED"].auto_filter.ref == "A1:B3"


# ------------------------------------------------------------------ permutation units
def test_target_exon_key_collapses_flanking_exon_variants():
    ids = np.asarray(["chr1:+:100:200:1:2:3:4", "chr1:+:100:200:9:9:9:9",
                      "chr1:+:300:400:1:2:3:4"], dtype="U")
    keys, parsed = calib.target_exon_key(ids)
    assert parsed is True
    assert list(keys) == ["chr1:+:100:200", "chr1:+:100:200", "chr1:+:300:400"]
    opaque, parsed = calib.target_exon_key(np.asarray(["exon0", "exon1"], dtype="U"))
    assert parsed is False and list(opaque) == ["exon0", "exon1"]


def clustered_npz(tmp_path, fg_rows, bg_rows, n_windows=24, name="T.AAA"):
    rng = np.random.default_rng(149)
    per_region = n_windows // len(io.REGIONS)
    store = {"schema_version": np.asarray(io.SCHEMA_VERSION), "arm": np.asarray("S"),
             "motif": np.asarray(name), "region_name": np.asarray(io.REGIONS, dtype="U"),
             "region_of_position": np.repeat(np.arange(len(io.REGIONS), dtype=np.int8), per_region),
             "position": np.tile(np.arange(per_region, dtype=np.int32), len(io.REGIONS))}
    for group, ids in (("up", fg_rows), ("dn", fg_rows), ("bg", bg_rows)):
        store[group + "_exon_id"] = np.asarray(ids, dtype="U")
        io.pack_group(store, group,
                      rng.poisson(0.4, size=(len(ids), n_windows)).astype(np.int16))
    path = tmp_path / (name + ".counts.npz")
    np.savez_compressed(path, **store)
    return path


def ids_for(exons, copies):
    out = []
    for exon, n in zip(exons, copies):
        for k in range(n):
            out.append(f"chr1:+:{exon}:{exon + 100}:{k}:{k}:{k}:{k}")
    return out


def test_cluster_draw_never_splits_a_target_exon_and_keeps_the_row_count(tmp_path):
    fg = ids_for([1000, 2000, 3000], [2, 1, 1])          # 4 rows over 3 exons
    bg = ids_for([10000 + 1000 * i for i in range(12)], [2, 2, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1])
    model = calib.MotifModel(clustered_npz(tmp_path, fg, bg), "up", "cluster")
    sampler = calib.ClusterSampler(model.cluster_codes, model.n1)
    draws = sampler.draw(np.random.default_rng(149), 40)
    assert draws.shape == (40, model.n1)
    codes = model.cluster_codes
    for row in draws:
        assert len(set(row.tolist())) == model.n1, "a row was drawn twice"
        for code in set(codes[row]):
            assert int((codes[row] == code).sum()) == int((codes == code).sum()), \
                "a target-exon cluster was split between the drawn and undrawn set"


def test_cluster_draw_matches_the_observed_size_composition(tmp_path):
    fg = ids_for([1000, 2000], [3, 1])
    bg = ids_for([10000 + 1000 * i for i in range(8)], [3, 3, 1, 1, 1, 1, 1, 1])
    model = calib.MotifModel(clustered_npz(tmp_path, fg, bg), "up", "cluster")
    sampler = calib.ClusterSampler(model.cluster_codes, model.n1)
    assert sampler.need == {1: 1, 3: 1}
    for row in sampler.draw(np.random.default_rng(3), 25):
        sizes = sorted(int((model.cluster_codes[row] == c).sum())
                       for c in set(model.cluster_codes[row]))
        assert sizes == [1, 3]


def test_cluster_sampler_refuses_a_changed_set_that_splits_a_cluster():
    # Rows 0..3 are the changed set, but cluster 1 straddles the boundary at row 3.
    codes = np.array([0, 0, 1, 1, 1, 2, 3, 4], dtype=np.int64)
    with pytest.raises(ValueError, match="splits a target-exon cluster"):
        calib.ClusterSampler(codes, 4)
    calib.ClusterSampler(codes, 5)  # whole clusters: fine


def test_exon_dedupe_drops_duplicate_target_exons(tmp_path):
    fg = ids_for([1000, 2000], [3, 1])
    bg = ids_for([10000 + 1000 * i for i in range(6)], [2, 1, 1, 1, 1, 1])
    path = clustered_npz(tmp_path, fg, bg)
    rows = calib.MotifModel(path, "up", "row")
    dedupe = calib.MotifModel(path, "up", "exon-dedupe")
    assert (rows.n1, rows.n0) == (4, 7)
    assert (dedupe.n1, dedupe.n0) == (2, 6)
    assert dedupe.n_distinct_exons == (2, 6)


def test_cluster_unit_refuses_archives_without_engine_exon_identifiers(tmp_path):
    path = clustered_npz(tmp_path, ["a", "b", "c"], ["d", "e", "f", "g"])
    calib.MotifModel(path, "up", "row")  # row permutation never needs the identity
    with pytest.raises(ValueError, match="eight-field exon identifiers"):
        calib.MotifModel(path, "up", "cluster")


# ------------------------------------------------------------------ end to end, synthetic genome
def run_skill(args, tmp_path):
    return subprocess.run([sys.executable, str(ROOT / "tools" / "rmaps3_skill_run.py")] + args,
                          cwd=ROOT, capture_output=True, text=True, timeout=900,
                          env=dict(os.environ, RMAPS_FORCE_MOTIF_FALLBACK="1"))


def test_quick_mode_runs_end_to_end_on_the_synthetic_genome(synthetic, tmp_path):  # noqa: F811
    _, genome_root = synthetic
    paths = input_paths(tmp_path)
    motifs = tmp_path / "motifs.tsv"
    motifs.write_text("Protein_name\tregularExpression\nTEST\tAA\nOTHER\tATAT\n")
    alias = tmp_path / "alias.tsv"
    alias.write_text("table_name\thgnc_symbol\nTEST\tTESTHGNC\n")
    out = tmp_path / "run"
    result = run_skill(["--mode", "quick", "--arm", "SYNTH", "--out", str(out),
                        "--up", str(paths["up"]), "--dn", str(paths["dn"]),
                        "--bg", str(paths["bg"]),
                        "--genome-root", str(genome_root), "--genome", "synthetic",
                        "--engine", "audited", "--engine-root", str(ROOT),
                        "--stat-method", "fisher", "--known-motifs", str(motifs),
                        "--additional-motifs", "NA", "--alias-table", str(alias),
                        "--window", "4", "--step", "2", "--intron", "20", "--exon", "10",
                        "--workers", "1", "--blas-threads", "1"], tmp_path)
    assert result.returncode == 0, result.stdout + result.stderr
    manifest = json.loads((out / "run_manifest.json").read_text())
    assert manifest["status"] == "complete"      # audited Fisher: its positional archives are the required set
    assert manifest["event_counts"] == {"up": 1, "dn": 1, "bg": 1}
    for name in ("quick_summary.xlsx", "index.html", "command.log", "versions.txt", "md5.txt"):
        assert (out / name).is_file() and (out / name).stat().st_size > 0, name
    import openpyxl
    book = openpyxl.load_workbook(out / "quick_summary.xlsx")
    assert {"README", "INCLUDED", "SKIPPED", "RBP_best_motif"} <= set(book.sheetnames)
    rbp = book["RBP_best_motif"]
    header = [c.value for c in rbp[1]]
    names = {row[header.index("RBP")] for row in rbp.iter_rows(min_row=2, values_only=True)}
    assert "TESTHGNC" in names, "the alias table must rename the motif table's RBP"
    index = (out / "index.html").read_text(encoding="utf-8")
    assert "pVal.up.vs.bg.RNAmap.txt" in index and "quick_summary.xlsx" in index
    assert "counts MOTIF HITS" in index, "the index must carry the statistic caveat"


def test_full_mode_refuses_a_fisher_run_before_doing_any_work(tmp_path):
    out = tmp_path / "full"
    result = run_skill(["--mode", "full", "--arm", "SYNTH", "--out", str(out),
                        "--up", "u", "--dn", "d", "--bg", "b",
                        "--stat-method", "fisher"], tmp_path)
    assert result.returncode == 2
    assert "needs --stat-method mannwhitney" in result.stderr
    assert not out.exists(), "a refused run must not create an output tree"


def test_full_mode_refuses_the_audited_engine(tmp_path):
    result = run_skill(["--mode", "full", "--arm", "SYNTH", "--out", str(tmp_path / "f2"),
                        "--up", "u", "--dn", "d", "--bg", "b",
                        "--engine", "audited"], tmp_path)
    assert result.returncode == 2
    assert "--engine released" in result.stderr


# ------------------------------------------------------------------ quick-mode figures
FIG_ARGS = ("--mode", "quick", "--arm", "QKI_KO_T", "--out", "o", "--up", "u", "--dn", "d",
            "--bg", "b")


def test_figures_need_an_arm_rule_suffix_and_a_gate_record():
    with pytest.raises(ValueError, match="<NAME>_<RULE>"):
        skill.validate(parse("--mode", "quick", "--arm", "NOSUFFIX", "--out", "o",
                             "--up", "u", "--dn", "d", "--bg", "b", "--gate-counts", "c.json"))
    with pytest.raises(ValueError, match="--gate-counts"):
        skill.validate(parse(*FIG_ARGS))
    with pytest.raises(ValueError, match="none is stated"):     # suffix T is not a rule
        skill.validate(parse(*FIG_ARGS, "--gate-counts", "c.json"))
    skill.validate(parse(*FIG_ARGS, "--gate-counts", "c.json", "--gate-rule", "A"))
    skill.validate(parse(*FIG_ARGS, "--no-figures"))
    skill.validate(parse("--mode", "quick", "--arm", "NOSUFFIX", "--out", "o", "--up", "u",
                         "--dn", "d", "--bg", "b", "--stat-method", "fisher"))
    with pytest.raises(ValueError, match="for pre-split inputs"):
        skill.validate(parse("--mode", "quick", "--arm", "QKI_KO_T", "--out", "o",
                             "--rmats-se", "se.txt", "--filter", "f.json",
                             "--gate-counts", "c.json"))


def test_figures_are_skipped_with_a_reason_off_the_released_rank_sum_layer():
    args = parse(*FIG_ARGS, "--stat-method", "fisher")
    assert "--stat-method fisher" in skill.figure_skip_reason(args, {"verified": False})
    args = parse(*FIG_ARGS)
    assert "verified count archives" in skill.figure_skip_reason(args, {"verified": False,
                                                                        "reason": "x"})
    assert skill.figure_skip_reason(args, {"verified": True}) == ""


def test_gate_record_must_match_the_input_event_counts(synthetic, tmp_path):  # noqa: F811
    _, genome_root = synthetic
    paths = input_paths(tmp_path)
    gate = tmp_path / "counts.json"
    gate.write_text(json.dumps({"n_up": 1, "n_dn": 1, "n_bg": 1, "n_expr_unknown_in_fg": 0,
                                "n_expr_unknown_in_bg": 3}))
    args = parse("--mode", "quick", "--arm", "S_A", "--out", str(tmp_path / "o"),
                 "--up", str(paths["up"]), "--dn", str(paths["dn"]), "--bg", str(paths["bg"]),
                 "--genome-root", str(genome_root), "--genome", "synthetic",
                 "--gate-counts", str(gate))
    (tmp_path / "o").mkdir()
    skill.prepare_inputs(args, tmp_path / "o", lambda m: None, {})
    record = json.loads((tmp_path / "o" / "event_counts.json").read_text())
    assert record["n_expr_unknown_in_bg"] == 3
    assert record["gate_counts_source"] == str(gate.resolve())
    gate.write_text(json.dumps({"n_up": 2, "n_dn": 1, "n_bg": 1, "n_expr_unknown_in_fg": 0,
                                "n_expr_unknown_in_bg": 3}))
    (tmp_path / "o2").mkdir()
    with pytest.raises(ValueError, match="n_up=2 disagrees"):
        skill.prepare_inputs(args, tmp_path / "o2", lambda m: None, {})


def fork_motif_keys():
    """Every motif key the released engine emits for the fork's own data/ motif tables."""
    import csv
    keys = []
    with open(ROOT / "data" / "knownMotifs.human.mouse.txt", encoding="utf-8") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            keys.append(f"{row['Protein_name']}.{row['regularExpression']}")
    with open(ROOT / "data" / "ESRP.like.motif.txt", encoding="utf-8") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            keys.append(f"{row['name']}.{row['motif']}")
    return keys


def fork_data_engine_run(tmp_path, arm="QKI_KO_T"):
    """A released-layout engine run over the fork's data/ motif tables with QKI planted first
    in INCLUDED x Upstream Intron and SKIPPED x Downstream Intron."""
    out = tmp_path / "run"
    engine = out / "engine" / arm
    engine.mkdir(parents=True)
    counts_root = out / "counts"
    (counts_root / arm).mkdir(parents=True)
    motifs = sorted(fork_motif_keys())
    rng = np.random.default_rng(149)
    per_region = 4
    n_windows = per_region * len(io.REGIONS)
    region_index = np.repeat(np.arange(len(io.REGIONS), dtype=np.int8), per_region)
    position = np.tile(np.arange(per_region, dtype=np.int32), len(io.REGIONS))
    roots = {"up": {}, "dn": {}}
    for motif in motifs:
        store = {"schema_version": np.asarray(io.SCHEMA_VERSION), "arm": np.asarray(arm),
                 "motif": np.asarray(motif), "region_name": np.asarray(io.REGIONS, dtype="U"),
                 "region_of_position": region_index, "position": position}
        qki = motif.startswith("QKI.")
        for group, n_rows in (("up", 20), ("dn", 18), ("bg", 200)):
            rate = 0.8 if qki and group != "bg" else 0.2
            matrix = rng.poisson(rate, size=(n_rows, n_windows)).astype(np.int16)
            store[group + "_exon_id"] = np.asarray(
                ["chr1:+:{}:{}:1:2:3:4".format(1000 + i, 1100 + i) for i in range(n_rows)],
                dtype="U")
            io.pack_group(store, group, matrix)
        np.savez_compressed(counts_root / arm / (motif + ".counts.npz"), **store)
        for direction in ("up", "dn"):
            roots[direction][motif] = {r: float(10.0 ** -rng.uniform(0.0, 4.0)) for r in io.REGIONS}
        if qki:
            roots["up"][motif]["UpstreamIntron"] = 1e-12
            roots["dn"][motif]["DownstreamIntron"] = 1e-11
    for direction in ("up", "dn"):
        with open(engine / f"pVal.{direction}.vs.bg.RNAmap.txt", "w", newline="\n",
                  encoding="utf-8") as handle:
            handle.write("\t".join(["RBP"] + [io.ROOT_COLUMNS[r] for r in io.REGIONS]) + "\n")
            for motif in motifs:
                handle.write("\t".join([motif] + [repr(roots[direction][motif][r])
                                                  for r in io.REGIONS]) + "\n")
    (out / "engine" / f"{arm}_command.log").write_text(
        f"set={arm} engine=released ({ROOT} @ {skill.git_revision(ROOT)}) stat=mannwhitney\n"
        "exit=0 wall_s=1\n", encoding="utf-8")
    (out / "event_counts.json").write_text(json.dumps(
        {"n_up": 20, "n_dn": 18, "n_bg": 200, "n_expr_unknown_in_fg": 0,
         "n_expr_unknown_in_bg": 2}), encoding="utf-8")
    alias = skill.load_alias(ROOT / "data" / "rbp_alias_hgnc_2026-09-17.tsv")
    names = {m.split(".", 1)[0] for m in motifs if not m.startswith("motif_")}
    symbols = sorted({alias.get(n, n) for n in names} | set(alias.values()))
    gtf = tmp_path / "genes.gtf"
    gtf.write_text("".join(
        f'chr1\tT\tgene\t1\t2\t.\t+\t.\tgene_id "ENSG{i:011d}.1"; gene_name "{s}";\n'
        for i, s in enumerate(symbols)), encoding="utf-8")
    (tmp_path / "splice.txt").write_text("SRSF1\nSNRPA\n", encoding="utf-8")
    (tmp_path / "broad.txt").write_text("PTBP1\nHNRNPC\n", encoding="utf-8")
    return out, engine, counts_root, roots, motifs, alias, gtf


def build_fork_figures(tmp_path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    out, engine, counts_root, roots, motifs, alias, gtf = fork_data_engine_run(tmp_path)
    args = parse("--mode", "quick", "--arm", "QKI_KO_T", "--out", str(out),
                 "--up", "u", "--dn", "d", "--bg", "b", "--gate-counts", "c.json",
                 "--engine-root", str(ROOT), "--gtf", str(gtf),
                 "--spliceosome-list", str(tmp_path / "splice.txt"),
                 "--broad-binders-list", str(tmp_path / "broad.txt"),
                 "--positive-control", "QKI")
    scores = {m: skill.motif_scores(counts_root, args.arm, m) for m in motifs}
    result = skill.build_quick_figures(args, out, engine, roots, motifs, scores, alias,
                                       ROOT / "data" / "knownMotifs.human.mouse.txt",
                                       ROOT / "data" / "ESRP.like.motif.txt", lambda m: None)
    return args, out, engine, result


def test_quick_mode_draws_the_four_main_layer_figures_on_the_fork_data_tables(tmp_path):
    args, out, engine, result = build_fork_figures(tmp_path)
    figures = out / "figures"
    expected = [f"QKI_KO_T_SE_{kind}_released_ranksum_rawP{variant}"
                for kind in ("byRBP", "byMotif") for variant in ("", "_noSpliceosome_noBroad")]
    for stem in expected:
        for suffix in (".png", ".svg"):
            path = figures / (stem + suffix)
            assert path.is_file() and path.stat().st_size > 0, path
    assert sorted(p.name for p in result["figures"]) == sorted(
        s + x for s in expected for x in (".png", ".svg"))
    assert not list(figures.rglob("*calibrated*")), "quick mode must never draw the calibrated layer"
    svg = (figures / (expected[0] + ".svg")).read_text(encoding="utf-8")
    assert "use for RBP ORDER only" in svg
    assert "n events (rMATS SE rows) included=20 / skipped=18 / background=200" in svg
    assert "0 foreground / 2 background events" in svg, "the footer must print the gate record"
    assert "Dot size = motif-score ratio" in svg
    controls = result["controls"]
    assert len(controls) == 8 and all(r["pass"] for r in controls)
    assert (figures / "positive_control_audit.tsv").is_file()
    header = (figures / "motif_scores" / "QKI_KO_T" / "per_motif_regions.tsv").read_text(
        encoding="utf-8").splitlines()[0]
    assert "calibrated" not in header and "count_ratio" in header
    index = skill.write_index(args, out, engine, fake_counts(), [], {"verified": True}, [], result)
    html = index.read_text(encoding="utf-8")
    for stem in expected:
        assert f"figures/{stem}.png" in html
    assert "use them for the RBP ORDER only" in html


def test_quick_figures_are_byte_identical_across_two_builds(tmp_path):
    import hashlib
    _, _, _, first = build_fork_figures(tmp_path / "a")
    _, _, _, second = build_fork_figures(tmp_path / "b")

    def digest(path):
        return hashlib.md5(path.read_bytes()).hexdigest()

    assert [p.name for p in first["figures"]] == [p.name for p in second["figures"]]
    assert [digest(p) for p in first["figures"]] == [digest(p) for p in second["figures"]]


def test_site_paths_have_no_default_and_are_refused_when_needed():
    presplit = ("--up", "u", "--dn", "d", "--bg", "b")
    with pytest.raises(ValueError, match="--genome-root"):
        skill.require_site_paths(parse(*BASE, *presplit, "--engine-root", "e"))
    with pytest.raises(ValueError, match="--engine-root"):
        skill.require_site_paths(parse(*BASE, *presplit, "--genome-root", "g"))
    skill.require_site_paths(parse(*BASE, *presplit, "--genome-root", "g", "--engine-root", "e"))
    skill.require_site_paths(parse(*BASE, *presplit, "--genome-root", "g", "--engine", "audited"))
    figures = ("--mode", "quick", "--arm", "QKI_KO_T", "--out", "o", *presplit,
               "--genome-root", "g", "--engine-root", "e")
    with pytest.raises(ValueError, match="--gtf, --spliceosome-list, --broad-binders-list"):
        skill.require_site_paths(parse(*figures))
    skill.require_site_paths(parse(*figures, "--gtf", "x", "--spliceosome-list", "s",
                                   "--broad-binders-list", "b"))
    with pytest.raises(ValueError, match="released engine"):
        skill.engine_root_of(parse(*BASE, *presplit))
    assert skill.engine_root_of(parse(*BASE, *presplit, "--engine", "audited")) == ROOT
